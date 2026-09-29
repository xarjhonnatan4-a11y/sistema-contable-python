import streamlit as st
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, Border, Side
import io
import os
import json
from datetime import datetime
import calendar
from utils.encoding import normalizar_columnas_dataframe

CONFIG_FILE = "empresas.json"

# -------- FUNCIONES AUXILIARES --------
def cargar_empresas():
    if os.path.exists(CONFIG_FILE):
        with open(CONFIG_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    return []

def guardar_empresas(empresas):
    with open(CONFIG_FILE, "w", encoding="utf-8") as f:
        json.dump(empresas, f, indent=4, ensure_ascii=False)

def cargar_partidas_seguro(path):
    if os.path.exists(path):
        with open(path, "r", encoding="utf-8") as f:
            try:
                data = json.load(f)
                return data if isinstance(data, list) else []
            except json.JSONDecodeError:
                return []
    return []

def convertir_json_serializable(obj):
    if isinstance(obj, dict):
        return {k: convertir_json_serializable(v) for k, v in obj.items()}
    elif isinstance(obj, list):
        return [convertir_json_serializable(i) for i in obj]
    elif isinstance(obj, np.integer):
        return int(obj)
    elif isinstance(obj, np.floating):
        return float(obj)
    elif isinstance(obj, pd.Timestamp):
        return obj.strftime("%Y-%m-%d")
    elif pd.isna(obj):
        return None
    return obj

# -------- INTERFAZ PRINCIPAL --------
st.title("📘 Generador de Libro Contable - Ventas SAT")

menu = st.sidebar.radio("Menú", ["Registrar empresa", "Generar libro"],index=1)

empresas = cargar_empresas()

# -------- REGISTRAR / EDITAR / ELIMINAR EMPRESA --------
if menu == "Registrar empresa":
    st.header("Gestión de empresas")

    # Registrar nueva empresa
    with st.expander("➕ Registrar nueva empresa", expanded=True):
        with st.form("registro_empresa"):
            nombre = st.text_input("Nombre de la empresa")
            nit = st.text_input("NIT de la empresa")
            propietario = st.text_input("Propietario")
            regimen = st.selectbox("Régimen", ["Mensual", "Trimestral"])
            es_gasolinera = st.checkbox("¿Es gasolinera?")
            submitted = st.form_submit_button("Guardar empresa")

        if submitted:
            if not nombre.strip():
                st.warning("⚠️ El nombre de la empresa es obligatorio.")
            else:
                nueva = {
                    "nombre": nombre,
                    "nit": nit,
                    "propietario": propietario,
                    "regimen": regimen,
                    "es_gasolinera": es_gasolinera
                }
                empresas.append(nueva)
                guardar_empresas(empresas)
                st.success(f"✅ Empresa '{nombre}' registrada con éxito")

    # Editar empresa
    if empresas:
        with st.expander("✏️ Editar empresa"):
            empresa_sel = st.selectbox(
                "Seleccione la empresa a editar",
                [e["nombre"] for e in empresas]
            )
            empresa_data = next(e for e in empresas if e["nombre"] == empresa_sel)

            with st.form("editar_empresa"):
                nuevo_nombre = st.text_input("Nombre de la empresa", value=empresa_data["nombre"])
                nuevo_nit = st.text_input("NIT de la empresa", value=empresa_data.get("nit", ""))
                nuevo_propietario = st.text_input("Propietario", value=empresa_data["propietario"])
                nuevo_regimen = st.selectbox(
                    "Régimen",
                    ["Mensual", "Trimestral"],
                    index=["Mensual", "Trimestral"].index(empresa_data["regimen"])
                )
                nuevo_es_gasolinera = st.checkbox(
                    "¿Es gasolinera?",
                    value=empresa_data["es_gasolinera"]
                )
                actualizar = st.form_submit_button("Actualizar empresa")

            if actualizar:
                empresa_data["nombre"] = nuevo_nombre
                empresa_data["nit"] = nuevo_nit
                empresa_data["propietario"] = nuevo_propietario
                empresa_data["regimen"] = nuevo_regimen
                empresa_data["es_gasolinera"] = nuevo_es_gasolinera
                guardar_empresas(empresas)
                st.success(f"✅ Empresa '{nuevo_nombre}' actualizada con éxito")

    # Eliminar empresa
    if empresas:
        with st.expander("🗑️ Eliminar empresa"):
            empresa_sel_del = st.selectbox(
                "Seleccione la empresa a eliminar",
                [e["nombre"] for e in empresas],
                key="del"
            )
            if st.button("Eliminar empresa seleccionada"):
                empresas = [e for e in empresas if e["nombre"] != empresa_sel_del]
                guardar_empresas(empresas)
                st.success(f"✅ Empresa '{empresa_sel_del}' eliminada con éxito")

# -------- GENERAR LIBRO --------
elif menu == "Generar libro":
    if not empresas:
        st.warning("⚠️ No hay empresas registradas. Primero registre una en el menú.")
    else:
        st.header("Generar libro contable")
        empresa_sel = st.selectbox("Seleccione la empresa", [e["nombre"] for e in empresas])
        empresa_data = next(e for e in empresas if e["nombre"] == empresa_sel)

        # Año y mes
        anio_actual = datetime.now().year
        anio = st.number_input("Año a trabajar", value=anio_actual)
        mes_actual = datetime.now().month
        mes_defecto = 12 if mes_actual == 1 else mes_actual - 1
        meses_nombre = [
            "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio",
            "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"
        ]
        mes_num = st.number_input("Mes a trabajar", min_value=1, max_value=12, value=mes_defecto)
        mes = meses_nombre[int(mes_num) - 1]

        ultimo_folio = st.number_input("Último folio utilizado", min_value=0, value=0)
        archivo_excel = st.file_uploader(
            "Sube el archivo Excel descargado del SAT",
            type=["xls", "xlsx"]
        )

        st.info(f"""
        **NIT:** {empresa_data.get('nit', '')}  
        **Propietario:** {empresa_data['propietario']}  
        **Régimen:** {empresa_data['regimen']}  
        **Gasolinera:** {"Sí" if empresa_data['es_gasolinera'] else "No"}  
        """)

        if st.button("Generar libro contable") and archivo_excel:
            try:
                extension = os.path.splitext(archivo_excel.name)[1].lower()
                if extension == '.xls':
                    df = pd.read_excel(archivo_excel, engine='xlrd')
                else:
                    df = pd.read_excel(archivo_excel)
                df = normalizar_columnas_dataframe(df)

                if "Código de establecimiento" not in df.columns:
                    st.error("❌ El archivo no contiene la columna 'Código de establecimiento'")
                else:
                    codigos_empresas = df["Código de establecimiento"].dropna().unique().tolist()

                    wb = Workbook()
                    if "Sheet" in wb.sheetnames:
                        del wb["Sheet"]

                    resumen_establecimientos = []
                    total_folios_global = 0

                    for codigo in codigos_empresas:
                        # FILTRAR DATOS DE CADA EMPRESA
                        df_empresa = df[df["Código de establecimiento"] == codigo].copy()

                        columnas_requeridas = [
                            'Fecha de emisión',
                            'Tipo de DTE (nombre)',
                            'Número del DTE',
                            'ID del receptor',
                            'Nombre completo del receptor',
                            'Gran Total (Moneda Original)',
                            'Estado'
                        ]
                        if empresa_data['es_gasolinera']:
                            columnas_requeridas.append('Petróleo (monto de este impuesto)')

                        df_empresa = df_empresa[columnas_requeridas]

                        # Formateo y orden
                        df_empresa['Fecha de emisión'] = pd.to_datetime(
                            df_empresa['Fecha de emisión'].astype(str).str.split('T').str[0],
                            errors='coerce'
                        ).dt.day

                        df_empresa = df_empresa.sort_values(by="Fecha de emisión")

                        if empresa_data['es_gasolinera']:
                            df_empresa["IDP"] = pd.to_numeric(
                                df_empresa["Petróleo (monto de este impuesto)"],
                                errors="coerce"
                            ).fillna(0.00)

                            df_empresa["Total"] = pd.to_numeric(
                                df_empresa["Gran Total (Moneda Original)"],
                                errors="coerce"
                            ).fillna(0) - df_empresa["IDP"]
                        else:
                            df_empresa["Total"] = pd.to_numeric(
                                df_empresa["Gran Total (Moneda Original)"],
                                errors="coerce"
                            ).fillna(0)

                        df_empresa["Venta Neta"] = (df_empresa["Total"] / 1.12).round(2)
                        df_empresa["IVA Débito"] = (df_empresa["Venta Neta"] * 0.12).round(2)

                        anulados = df_empresa["Estado"].astype(str).str.lower() == "anulado"
                        df_empresa.loc[anulados, ["Total", "Venta Neta", "IVA Débito"]] = 0.00
                        df_empresa.loc[anulados, "Nombre completo del receptor"] = "ANULADO"

                        if empresa_data['es_gasolinera']:
                            df_empresa.loc[anulados, "IDP"] = 0.00

                        df_empresa = df_empresa.rename(columns={
                            'Fecha de emisión': 'DIA',
                            'Tipo de DTE (nombre)': 'Tipo DTE',
                            'Número del DTE': 'Numero',
                            'ID del receptor': 'Nit',
                            'Nombre completo del receptor': 'Cliente'
                        })

                        df_empresa = df_empresa.drop(columns=["Estado", "Gran Total (Moneda Original)"])
                        if empresa_data['es_gasolinera']:
                            df_empresa = df_empresa.drop(columns=["Petróleo (monto de este impuesto)"])

                        # CREAR HOJA PARA CADA EMPRESA
                        nombre_hoja = f"Empresa_{codigo}"
                        ws = wb.create_sheet(title=nombre_hoja)

                        if empresa_data['es_gasolinera']:
                            columnas_exportar = [
                                "DIA", "Tipo DTE", "Numero", "Nit", "Cliente",
                                "IDP", "Total", "Venta Neta", "IVA Débito"
                            ]
                        else:
                            columnas_exportar = [
                                "DIA", "Tipo DTE", "Numero", "Nit", "Cliente",
                                "Total", "Venta Neta", "IVA Débito"
                            ]

                        total_lineas = int(len(df_empresa))
                        lineas_por_folio = 50
                        total_folios = (
                            (total_lineas + lineas_por_folio - 1) // lineas_por_folio
                            if total_lineas > 0 else 1
                        )
                        total_folios_global += int(total_folios)

                        acumulado_total = 0.0
                        acumulado_venta_neta = 0.0
                        acumulado_iva = 0.0
                        acumulado_idp = 0.0
                        fila_actual = 1

                        borde_fino = Border(
                            left=Side(style='thin'),
                            right=Side(style='thin'),
                            top=Side(style='thin'),
                            bottom=Side(style='thin')
                        )

                        estilo_anulado = Font(color="FF0000", bold=True)

                        for folio in range(total_folios):
                            folio_actual = int(ultimo_folio) + folio + 1

                            ws.merge_cells(
                                start_row=fila_actual,
                                start_column=1,
                                end_row=fila_actual,
                                end_column=len(columnas_exportar)
                            )
                            cell = ws.cell(row=fila_actual, column=1)
                            cell.value = (
                                f"{empresa_data['nombre']} | NIT: {empresa_data.get('nit', '')} | "
                                f"Propietario: {empresa_data['propietario']} | Año: {int(anio)} | "
                                f"Mes: {mes} | Folio: {folio_actual}"
                            )
                            cell.font = Font(bold=True)
                            cell.alignment = Alignment(horizontal='center')
                            fila_actual += 1

                            ws.merge_cells(
                                start_row=fila_actual,
                                start_column=1,
                                end_row=fila_actual,
                                end_column=len(columnas_exportar)
                            )
                            cell = ws.cell(row=fila_actual, column=1)
                            cell.value = "CIFRAS EXPRESADAS EN QUETZALES - LIBRO DE VENTAS"
                            cell.font = Font(bold=True, italic=True)
                            cell.alignment = Alignment(horizontal='center')
                            fila_actual += 1

                            for col_num, col_name in enumerate(columnas_exportar, 1):
                                cell = ws.cell(row=fila_actual, column=col_num, value=col_name)
                                cell.font = Font(bold=True)
                                cell.alignment = Alignment(horizontal='center')
                                cell.border = borde_fino
                            fila_actual += 1

                            # VIENEN solo si no es el primer folio
                            if folio > 0:
                                inicio_col = len(columnas_exportar) - 4 if empresa_data['es_gasolinera'] else len(columnas_exportar) - 3
                                ws.cell(row=fila_actual, column=inicio_col, value="VIENEN:").font = Font(italic=True, bold=True)

                                if empresa_data['es_gasolinera']:
                                    ws.cell(row=fila_actual, column=inicio_col + 1, value=float(round(acumulado_idp, 2))).font = Font(italic=True, bold=True)
                                    ws.cell(row=fila_actual, column=inicio_col + 2, value=float(round(acumulado_total, 2))).font = Font(italic=True, bold=True)
                                    ws.cell(row=fila_actual, column=inicio_col + 3, value=float(round(acumulado_venta_neta, 2))).font = Font(italic=True, bold=True)
                                    ws.cell(row=fila_actual, column=inicio_col + 4, value=float(round(acumulado_iva, 2))).font = Font(italic=True, bold=True)
                                else:
                                    ws.cell(row=fila_actual, column=inicio_col + 1, value=float(round(acumulado_total, 2))).font = Font(italic=True, bold=True)
                                    ws.cell(row=fila_actual, column=inicio_col + 2, value=float(round(acumulado_venta_neta, 2))).font = Font(italic=True, bold=True)
                                    ws.cell(row=fila_actual, column=inicio_col + 3, value=float(round(acumulado_iva, 2))).font = Font(italic=True, bold=True)
                                fila_actual += 1

                            inicio = folio * lineas_por_folio
                            fin = min(inicio + lineas_por_folio, total_lineas)
                            datos_folio = df_empresa.iloc[inicio:fin]

                            for fila_dato in datos_folio.itertuples(index=False):
                                for col_num, valor in enumerate(fila_dato, 1):
                                    valor_limpio = convertir_json_serializable(valor)
                                    cell = ws.cell(row=fila_actual, column=col_num, value=valor_limpio)
                                    cell.border = borde_fino
                                    col_name = columnas_exportar[col_num - 1]

                                    if col_name == "Cliente" and valor_limpio == "ANULADO":
                                        cell.font = estilo_anulado

                                    if col_name in ["Total", "Venta Neta", "IVA Débito", "IDP"]:
                                        cell.number_format = '"Q"#,##0.00'
                                fila_actual += 1

                            for _ in range(lineas_por_folio - len(datos_folio)):
                                for col in range(1, len(columnas_exportar) + 1):
                                    cell = ws.cell(row=fila_actual, column=col, value=None)
                                    cell.border = borde_fino
                                fila_actual += 1

                            suma_total = float(datos_folio["Total"].sum())
                            suma_venta_neta = float(datos_folio["Venta Neta"].sum())
                            suma_iva = float(datos_folio["IVA Débito"].sum())
                            suma_idp = float(datos_folio["IDP"].sum()) if empresa_data['es_gasolinera'] else 0.0

                            acumulado_total += round(suma_total, 2)
                            acumulado_venta_neta += round(suma_venta_neta, 2)
                            acumulado_iva += round(suma_iva, 2)
                            acumulado_idp += round(suma_idp, 2)

                            inicio_col = len(columnas_exportar) - 4 if empresa_data['es_gasolinera'] else len(columnas_exportar) - 3
                            if folio < total_folios - 1:
                                ws.cell(row=fila_actual, column=inicio_col, value="VAN:").font = Font(bold=True)
                            else:
                                ws.cell(row=fila_actual, column=inicio_col, value="TOTAL:").font = Font(bold=True)

                            if empresa_data['es_gasolinera']:
                                ws.cell(row=fila_actual, column=inicio_col + 1, value=float(round(acumulado_idp, 2))).number_format = '"Q"#,##0.00'
                                ws.cell(row=fila_actual, column=inicio_col + 2, value=float(round(acumulado_total, 2))).number_format = '"Q"#,##0.00'
                                ws.cell(row=fila_actual, column=inicio_col + 3, value=float(round(acumulado_venta_neta, 2))).number_format = '"Q"#,##0.00'
                                ws.cell(row=fila_actual, column=inicio_col + 4, value=float(round(acumulado_iva, 2))).number_format = '"Q"#,##0.00'
                            else:
                                ws.cell(row=fila_actual, column=inicio_col + 1, value=float(round(acumulado_total, 2))).number_format = '"Q"#,##0.00'
                                ws.cell(row=fila_actual, column=inicio_col + 2, value=float(round(acumulado_venta_neta, 2))).number_format = '"Q"#,##0.00'
                                ws.cell(row=fila_actual, column=inicio_col + 3, value=float(round(acumulado_iva, 2))).number_format = '"Q"#,##0.00'

                            fila_actual += 2

                        # --- Cálculo de Ventas Lubricantes y Venta Exhibición Shell ---
                        if empresa_data['es_gasolinera']:
                            df_empresa["Nit"] = df_empresa["Nit"].astype(str)
                            venta_lubricantes = float(
                                df_empresa[
                                    (df_empresa["IDP"] == 0) & (df_empresa["Nit"] != '321052')
                                ]["Venta Neta"].sum()
                            )
                            venta_exhibicion_shell = float(
                                df_empresa[
                                    (df_empresa["IDP"] == 0) & (df_empresa["Nit"] == '321052')
                                ]["Venta Neta"].sum()
                            )
                        else:
                            venta_lubricantes = 0.0
                            venta_exhibicion_shell = 0.0

                        # Resumen final
                        total_facturas = int(total_lineas)

                        ws.cell(row=fila_actual, column=5, value="Número de Facturas:").font = Font(bold=True)
                        ws.cell(row=fila_actual, column=6, value=total_facturas)
                        fila_actual += 1

                        if empresa_data['regimen'] == "Mensual":
                            if acumulado_venta_neta > 30000:
                                isr_mensual = float((acumulado_venta_neta - 30000) * 0.07 + 1500)
                            else:
                                isr_mensual = float(acumulado_venta_neta * 0.05)

                            ws.cell(row=fila_actual, column=5, value="ISR Mensual:").font = Font(bold=True)
                            ws.cell(row=fila_actual, column=6, value=float(round(isr_mensual, 2))).number_format = '"Q"#,##0.00'
                            fila_actual += 1

                            ws.cell(row=fila_actual, column=5, value="Retenciones:").font = Font(bold=True)
                            fila_actual += 1

                            ws.cell(row=fila_actual, column=5, value="ISR por pagar:").font = Font(bold=True)
                            ws.cell(row=fila_actual, column=6, value=f"=F{fila_actual-2}-F{fila_actual-1}").number_format = '"Q"#,##0.00'
                            fila_actual += 1

                        resumen_establecimientos.append({
                            "codigo_establecimiento": str(codigo),
                            "total": float(round(acumulado_total, 2)),
                            "venta_neta": float(round(acumulado_venta_neta, 2)),
                            "iva_debito": float(round(acumulado_iva, 2)),
                            "idp": float(round(acumulado_idp, 2)),
                            "facturas": int(total_facturas),
                            "venta_lubricantes": float(round(venta_lubricantes, 2)),
                            "venta_exhibicion_shell": float(round(venta_exhibicion_shell, 2))
                        })

                    # -------- HOJA: PARTIDA CONTABLE --------
                    ws_partida = wb.create_sheet(title="Partida Contable")
                    partida_numero = int(ultimo_folio) + int(total_folios_global)
                    ultimo_dia = calendar.monthrange(int(anio), int(mes_num))[1]
                    fecha_partida = f"{ultimo_dia:02d}/{int(mes_num):02d}/{int(anio)}"
                    fila = 1

                    # Encabezado
                    ws_partida.cell(row=fila, column=1, value=f"PDA {partida_numero}")
                    ws_partida.cell(row=fila, column=2, value=fecha_partida)
                    fila += 1

                    total_general = float(round(sum(x["total"] for x in resumen_establecimientos), 2))
                    venta_neta_general = float(round(sum(x["venta_neta"] for x in resumen_establecimientos), 2))
                    iva_general = float(round(sum(x["iva_debito"] for x in resumen_establecimientos), 2))

                    # Partida para GASOLINERA
                    if empresa_data['es_gasolinera']:
                        datos_debe = [
                            ("Bancos", 0.0),
                            ("Retencion IVA", 0.0),
                            ("Comisiones Pagadas", 0.0),
                            ("IVA Credito", 0.0),
                            ("Clientes", 0.0),
                            ("Evaporación", 0.0),
                            ("Cupones", 0.0)
                        ]

                        for cuenta, monto in datos_debe:
                            ws_partida.cell(row=fila, column=2, value=cuenta)
                            ws_partida.cell(row=fila, column=3, value=float(monto))
                            ws_partida.cell(row=fila, column=3).number_format = '"Q"#,##0.00'
                            fila += 1

                        for item in resumen_establecimientos:
                            ws_partida.cell(row=fila, column=2, value=f"    A: Ventas Est. {item['codigo_establecimiento']}")
                            ws_partida.cell(row=fila, column=4, value=0.0)
                            ws_partida.cell(row=fila, column=4).number_format = '"Q"#,##0.00'
                            fila += 1

                            otras_cuentas = [
                                ("        Ventas Regular", 0.0),
                                ("        Ventas Diesel", 0.0),
                                ("        Ventas Lubricantes", float(round(item["venta_lubricantes"], 2))),
                                ("        Venta Exhibición Shell", float(round(item["venta_exhibicion_shell"], 2)))
                            ]

                            for cuenta, monto in otras_cuentas:
                                ws_partida.cell(row=fila, column=2, value=cuenta)
                                ws_partida.cell(row=fila, column=4, value=float(monto))
                                ws_partida.cell(row=fila, column=4).number_format = '"Q"#,##0.00'
                                fila += 1

                        ws_partida.cell(row=fila, column=2, value="        IVA Debito")
                        ws_partida.cell(row=fila, column=4, value=float(round(iva_general, 2)))
                        ws_partida.cell(row=fila, column=4).number_format = '"Q"#,##0.00'
                        fila += 1

                        otras_cuentas = [
                            ("        IDP Super", 0.0),
                            ("        IDP Regular", 0.0),
                            ("        IDP Diesel", 0.0),
                            ("        Clientes", 0.0)
                        ]

                        for cuenta, monto in otras_cuentas:
                            ws_partida.cell(row=fila, column=2, value=cuenta)
                            ws_partida.cell(row=fila, column=4, value=float(monto))
                            ws_partida.cell(row=fila, column=4).number_format = '"Q"#,##0.00'
                            fila += 1

                        ws_partida.cell(row=fila, column=2, value="Registro de los ingresos obtenidos en el presente mes")
                        ws_partida.cell(row=fila, column=3, value=f"=SUM(C2:C{fila-1})")
                        ws_partida.cell(row=fila, column=3).number_format = '"Q"#,##0.00'
                        ws_partida.cell(row=fila, column=4, value=f"=SUM(D2:D{fila-1})")
                        ws_partida.cell(row=fila, column=4).number_format = '"Q"#,##0.00'

                    else:
                        # No es gasolinera – partida sencilla con ventas separadas por establecimiento
                        ws_partida.cell(row=fila, column=2, value="Bancos")
                        ws_partida.cell(row=fila, column=3, value=float(round(total_general, 2)))
                        ws_partida.cell(row=fila, column=3).number_format = '"Q"#,##0.00'
                        fila += 1

                        ws_partida.cell(row=fila, column=2, value="Retencion IVA")
                        ws_partida.cell(row=fila, column=3, value=0.0)
                        ws_partida.cell(row=fila, column=3).number_format = '"Q"#,##0.00'
                        fila += 1

                        ws_partida.cell(row=fila, column=2, value="Retencion ISR")
                        ws_partida.cell(row=fila, column=3, value=0.0)
                        ws_partida.cell(row=fila, column=3).number_format = '"Q"#,##0.00'
                        fila += 1

                        for item in resumen_establecimientos:
                            ws_partida.cell(row=fila, column=2, value=f"    A: Ventas Est. {item['codigo_establecimiento']}")
                            ws_partida.cell(row=fila, column=4, value=float(round(item["venta_neta"], 2)))
                            ws_partida.cell(row=fila, column=4).number_format = '"Q"#,##0.00'
                            fila += 1

                        ws_partida.cell(row=fila, column=2, value="        IVA Débito")
                        ws_partida.cell(row=fila, column=4, value=float(round(iva_general, 2)))
                        ws_partida.cell(row=fila, column=4).number_format = '"Q"#,##0.00'
                        fila += 1

                        ws_partida.cell(row=fila, column=2, value="Registro de los ingresos obtenidos en el presente mes")
                        ws_partida.cell(row=fila, column=3, value=f"=SUM(C2:C{fila-1})")
                        ws_partida.cell(row=fila, column=3).number_format = '"Q"#,##0.00'
                        ws_partida.cell(row=fila, column=4, value=f"=SUM(D2:D{fila-1})")
                        ws_partida.cell(row=fila, column=4).number_format = '"Q"#,##0.00'

                    # Alinear texto
                    for i in range(2, fila + 1):
                        ws_partida.cell(row=i, column=2).alignment = Alignment(horizontal='left')

                    # -------- GUARDAR PARTIDA EN JSON --------
                    PARTIDAS_FILE = "partidas.json"

                    partida_json = {
                        "empresa": empresa_data["nombre"],
                        "empresa_nit": str(empresa_data.get("nit", "")).strip(),
                        "anio": int(anio),
                        "mes": mes,
                        "fecha": fecha_partida,
                        "cuentas": [],
                        "glosa": "Registro de los ingresos obtenidos en el presente mes"
                    }

                    if empresa_data['es_gasolinera']:
                        partida_json["cuentas"].append({"cuenta": "Bancos", "debe": 0.0, "haber": 0.0})
                        partida_json["cuentas"].append({"cuenta": "Retencion IVA", "debe": 0.0, "haber": 0.0})
                        partida_json["cuentas"].append({"cuenta": "Comisiones Pagadas", "debe": 0.0, "haber": 0.0})
                        partida_json["cuentas"].append({"cuenta": "IVA Credito", "debe": 0.0, "haber": 0.0})
                        partida_json["cuentas"].append({"cuenta": "Clientes", "debe": 0.0, "haber": 0.0})
                        partida_json["cuentas"].append({"cuenta": "Evaporación", "debe": 0.0, "haber": 0.0})
                        partida_json["cuentas"].append({"cuenta": "Cupones", "debe": 0.0, "haber": 0.0})

                        for item in resumen_establecimientos:
                            partida_json["cuentas"].append({
                                "cuenta": f"A: Ventas Est. {item['codigo_establecimiento']}",
                                "debe": 0.0,
                                "haber": 0.0
                            })
                            partida_json["cuentas"].append({"cuenta": "Ventas Regular", "debe": 0.0, "haber": 0.0})
                            partida_json["cuentas"].append({"cuenta": "Ventas Diesel", "debe": 0.0, "haber": 0.0})
                            partida_json["cuentas"].append({
                                "cuenta": "Ventas Lubricantes",
                                "debe": 0.0,
                                "haber": float(round(item["venta_lubricantes"], 2))
                            })
                            partida_json["cuentas"].append({
                                "cuenta": "Venta Exhibición Shell",
                                "debe": 0.0,
                                "haber": float(round(item["venta_exhibicion_shell"], 2))
                            })

                        partida_json["cuentas"].append({
                            "cuenta": "IVA Débito",
                            "debe": 0.0,
                            "haber": float(round(iva_general, 2))
                        })
                        partida_json["cuentas"].append({"cuenta": "IDP Super", "debe": 0.0, "haber": 0.0})
                        partida_json["cuentas"].append({"cuenta": "IDP Regular", "debe": 0.0, "haber": 0.0})
                        partida_json["cuentas"].append({"cuenta": "IDP Diesel", "debe": 0.0, "haber": 0.0})
                        partida_json["cuentas"].append({"cuenta": "Clientes", "debe": 0.0, "haber": 0.0})
                    else:
                        partida_json["cuentas"].append({
                            "cuenta": "Bancos",
                            "debe": float(round(total_general, 2)),
                            "haber": 0.0
                        })
                        partida_json["cuentas"].append({"cuenta": "Retencion IVA", "debe": 0.0, "haber": 0.0})
                        partida_json["cuentas"].append({"cuenta": "Retencion ISR", "debe": 0.0, "haber": 0.0})

                        for item in resumen_establecimientos:
                            partida_json["cuentas"].append({
                                "cuenta": f"A: Ventas Est. {item['codigo_establecimiento']}",
                                "debe": 0.0,
                                "haber": float(round(item["venta_neta"], 2))
                            })

                        partida_json["cuentas"].append({
                            "cuenta": "IVA Débito",
                            "debe": 0.0,
                            "haber": float(round(iva_general, 2))
                        })

                    partidas = cargar_partidas_seguro(PARTIDAS_FILE)

                    nit_actual = str(empresa_data.get("nit", "")).strip()
                    nombre_actual = str(empresa_data.get("nombre", "")).strip()

                    existe = any(
                        (
                            str(p.get("empresa_nit", "")).strip() == nit_actual
                            or str(p.get("empresa", "")).strip() == nombre_actual
                        ) and
                        int(p.get("anio", 0) or 0) == int(anio) and
                        str(p.get("mes")) == str(mes)
                        for p in partidas
                    )

                    if not existe:
                        partidas.append(partida_json)
                        partidas_serializables = convertir_json_serializable(partidas)

                        with open(PARTIDAS_FILE, "w", encoding="utf-8") as f:
                            json.dump(partidas_serializables, f, indent=4, ensure_ascii=False)

                        st.success("✅ Partida guardada correctamente.")
                    else:
                        st.info("ℹ️ Ya existe una partida para este mes y empresa. No se guardó nuevamente.")

                    # Guardar y ofrecer descarga
                    output = io.BytesIO()
                    wb.save(output)
                    output.seek(0)

                    nombre_archivo = f"Libro_Ventas_{mes}_{int(anio)}.xlsx"
                    st.success("✅ Libro generado con éxito.")

                    st.download_button(
                        label="📥 Descargar libro contable",
                        data=output,
                        file_name=nombre_archivo,
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )

            except Exception as e:
                st.error(f"Ocurrió un error: {e}")
