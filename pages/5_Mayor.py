import streamlit as st
import pandas as pd
import json
import os
import base64
import re
from io import BytesIO
from datetime import datetime
import html as html_lib
import streamlit.components.v1 as components
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from utils.cuentas import cargar_cuentas

VENTAS_FILE = "partidas.json"
COMPRAS_FILE = "partidascompras.json"
PARTIDAS_GENERADAS_FILE = "partidas_generadas.json"
EMPRESAS_FILE = "empresas.json"
CIERRES_PARCIALES_FILE = "cierres_parciales.json"

MESES = {
    1: "Enero", 2: "Febrero", 3: "Marzo", 4: "Abril",
    5: "Mayo", 6: "Junio", 7: "Julio", 8: "Agosto",
    9: "Septiembre", 10: "Octubre", 11: "Noviembre", 12: "Diciembre",
}


# =========================
# UTILIDADES GENERALES
# =========================
def cargar_json_lista(ruta):
    if os.path.exists(ruta):
        with open(ruta, "r", encoding="utf-8") as f:
            try:
                data = json.load(f)
                return data if isinstance(data, list) else []
            except json.JSONDecodeError:
                return []
    return []


def normalizar_nit(nit):
    return str(nit or "").strip().replace("-", "").replace(" ", "").upper()


def normalizar_txt(txt):
    return " ".join(str(txt or "").strip().lower().split())


def normalizar_nombre_cuenta(nombre):
    txt = str(nombre or "").strip()
    if not txt:
        return ""

    txt = txt.replace("A :", "A:")
    txt = txt.replace("a :", "A:")

    if txt[:2].lower() == "a:":
        txt = txt[2:]
    elif txt[:3].lower() == "a :":
        txt = txt[3:]

    txt = " ".join(txt.strip().split())
    return txt


def clave_cuenta(nombre):
    return normalizar_txt(normalizar_nombre_cuenta(nombre))


def cargar_empresas():
    return cargar_json_lista(EMPRESAS_FILE)


def cargar_cierres_parciales():
    return cargar_json_lista(CIERRES_PARCIALES_FILE)


def guardar_cierres_parciales(data):
    with open(CIERRES_PARCIALES_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=4)


def es_mes_cierre_trimestral(mes):
    return int(mes or 0) in {3, 6, 9, 12}


def trimestre_desde_mes(mes):
    mes = int(mes or 0)
    if mes < 1:
        return 0
    return ((mes - 1) // 3) + 1


def rango_meses_hasta_trimestre(mes_corte):
    return list(range(1, int(mes_corte or 0) + 1))


def obtener_regimen_empresa(empresa_id, empresas):
    emp = buscar_empresa_por_nit(empresa_id, empresas)
    if not emp:
        return "Mensual"
    return str(emp.get("regimen", "Mensual")).strip() or "Mensual"


def dataframe_a_registros_json(df):
    if df is None or df.empty:
        return []

    df_limpio = df.copy()
    df_limpio = df_limpio.where(pd.notnull(df_limpio), None)

    registros = []
    for row in df_limpio.to_dict(orient="records"):
        limpio = {}
        for k, v in row.items():
            if isinstance(v, (int, float)) and pd.isna(v):
                limpio[k] = None
            elif isinstance(v, (int, float)):
                limpio[k] = round(float(v), 2)
            else:
                limpio[k] = v
        registros.append(limpio)
    return registros


def buscar_cierre_parcial(data, empresa_id, anio, trimestre):
    empresa_id = normalizar_nit(empresa_id)
    for item in data or []:
        if (
            normalizar_nit(item.get("empresa_nit")) == empresa_id
            and int(item.get("anio", 0) or 0) == int(anio)
            and int(item.get("trimestre", 0) or 0) == int(trimestre)
        ):
            return item
    return None


def guardar_o_actualizar_cierre_parcial(data, registro):
    empresa_id = normalizar_nit(registro.get("empresa_nit"))
    anio = int(registro.get("anio", 0) or 0)
    trimestre = int(registro.get("trimestre", 0) or 0)

    actualizado = False
    for i, item in enumerate(data or []):
        if (
            normalizar_nit(item.get("empresa_nit")) == empresa_id
            and int(item.get("anio", 0) or 0) == anio
            and int(item.get("trimestre", 0) or 0) == trimestre
        ):
            data[i] = registro
            actualizado = True
            break

    if not actualizado:
        data.append(registro)

    data.sort(key=lambda x: (
        str(x.get("empresa", "")).lower(),
        int(x.get("anio", 0) or 0),
        int(x.get("trimestre", 0) or 0)
    ))

    guardar_cierres_parciales(data)
    return registro


def construir_registro_cierre_parcial(empresa, empresa_id, regimen, anio, mes, df_resumen, estado_cuadre, df_pendientes=None):
    trimestre = trimestre_desde_mes(mes)
    meses_acumulados = rango_meses_hasta_trimestre(mes)

    return {
        "empresa": str(empresa).strip(),
        "empresa_nit": normalizar_nit(empresa_id),
        "regimen": str(regimen or "").strip(),
        "anio": int(anio),
        "trimestre": int(trimestre),
        "mes_corte": int(mes),
        "mes_corte_nombre": MESES.get(int(mes), str(mes)),
        "meses_acumulados": meses_acumulados,
        "fecha_guardado": str(pd.Timestamp.now()),
        "estado_cuadre": {
            "cuadra_inicial": bool((estado_cuadre or {}).get("cuadra_inicial", True)),
            "cuadra_movimientos": bool((estado_cuadre or {}).get("cuadra_movimientos", True)),
            "cuadra_final": bool((estado_cuadre or {}).get("cuadra_final", True)),
            "totales": (estado_cuadre or {}).get("totales", {}) or {},
        },
        "resumen_mayor": dataframe_a_registros_json(df_resumen),
        "saldos_pendientes": dataframe_a_registros_json(df_pendientes if df_pendientes is not None else pd.DataFrame()),
    }

def construir_apertura_desde_saldo_final(df_resumen_anterior, anio, mes_anterior):
    if df_resumen_anterior is None or df_resumen_anterior.empty:
        return None

    cuentas = []

    for _, row in df_resumen_anterior.iterrows():
        cuenta = str(row.get("cuenta", "")).strip()
        if not cuenta or cuenta.upper() == "TOTAL":
            continue

        saldo_debe = round(float(row.get("saldo_final_debe", 0) or 0), 2)
        saldo_haber = round(float(row.get("saldo_final_haber", 0) or 0), 2)

        if saldo_debe == 0 and saldo_haber == 0:
            continue

        nombre_cuenta = cuenta
        if saldo_haber > 0 and not nombre_cuenta.startswith("A: "):
            nombre_cuenta = f"A: {nombre_cuenta}"

        cuentas.append({
            "cuenta": nombre_cuenta,
            "debe": saldo_debe,
            "haber": saldo_haber
        })

    if not cuentas:
        return None

    return {
        "codigo": f"SALDO INICIAL {mes_anterior}",
        "fecha": f"{anio}-{int(mes_anterior):02d}-31",
        "glosa": f"Saldo inicial proveniente del saldo final de {MESES[int(mes_anterior)]}",
        "tipo_partida": "saldo_inicial",
        "cuentas": cuentas
    }


def buscar_empresa_por_nit(nit, empresas):
    nit_norm = normalizar_nit(nit)
    for e in empresas:
        if normalizar_nit(e.get("nit")) == nit_norm:
            return e
    return None


def buscar_empresa_por_nombre(nombre, empresas):
    nombre_norm = str(nombre or "").strip()
    for e in empresas:
        if str(e.get("nombre", "")).strip() == nombre_norm:
            return e
    return None


def obtener_id_empresa_desde_partida(partida, empresas):
    nit_partida = normalizar_nit(partida.get("empresa_nit"))
    nombre_partida = str(partida.get("empresa", "")).strip()

    if nit_partida:
        emp = buscar_empresa_por_nit(nit_partida, empresas)
        if emp:
            return normalizar_nit(emp.get("nit"))

    if nombre_partida:
        emp = buscar_empresa_por_nombre(nombre_partida, empresas)
        if emp:
            return normalizar_nit(emp.get("nit"))

    return ""


def buscar_partidas_generadas(data, empresa_id, anio, mes):
    empresa_id = normalizar_nit(empresa_id)

    for item in data:
        if (
            normalizar_nit(item.get("empresa_nit")) == empresa_id
            and int(item.get("anio", 0) or 0) == int(anio)
            and int(item.get("mes", 0) or 0) == int(mes)
        ):
            return item
    return None


def mes_a_numero(valor):
    if isinstance(valor, int):
        return valor
    if isinstance(valor, float):
        return int(valor)

    txt = str(valor).strip().lower()
    if txt.isdigit():
        return int(txt)

    meses_inv = {v.lower(): k for k, v in MESES.items()}
    return meses_inv.get(txt, 0)


def buscar_partida_apertura(data, empresa_id, anio):
    empresa_id = normalizar_nit(empresa_id)

    for p in data:
        nit_p = normalizar_nit(p.get("empresa_nit"))
        anio_p = int(p.get("anio", 0) or 0)
        mes_p = mes_a_numero(p.get("mes"))
        tipo_partida = str(p.get("tipo_partida", "")).strip().lower()
        codigo = str(p.get("codigo", "")).strip().upper()

        if (
            nit_p == empresa_id
            and anio_p == int(anio)
            and (
                mes_p == 0
                or tipo_partida == "apertura"
                or codigo == "PDA 0"
            )
        ):
            return p

    return None


def normalizar_apertura_contrapartidas(apertura):
    if not apertura:
        return None

    nueva_apertura = apertura.copy()
    cuentas_nuevas = []

    for mov in apertura.get("cuentas", []):
        cuenta = str(mov.get("cuenta", "")).strip()
        debe = float(mov.get("debe", 0) or 0)
        haber = float(mov.get("haber", 0) or 0)

        if haber > 0 and not cuenta.startswith("A: "):
            cuenta = f"A: {cuenta}"

        cuentas_nuevas.append({
            "cuenta": cuenta,
            "debe": debe,
            "haber": haber
        })

    nueva_apertura["cuentas"] = cuentas_nuevas
    return nueva_apertura


# =========================
# CATÁLOGO / NATURALEZA
# =========================
def obtener_catalogo_map():
    cuentas = cargar_cuentas() or []
    catalogo = {}

    for c in cuentas:
        nombre = str(c.get("nombre", "")).strip()
        if not nombre:
            continue

        catalogo[normalizar_txt(nombre)] = {
            "codigo": str(c.get("codigo", "")).strip(),
            "nombre": nombre,
            "tipo": normalizar_txt(c.get("tipo", "")),
            "naturaleza": normalizar_txt(c.get("naturaleza", "")),
            "activa": bool(c.get("activa", True))
        }

    return catalogo


def obtener_info_cuenta_catalogo(nombre_cuenta, catalogo_map=None):
    if catalogo_map is None:
        catalogo_map = obtener_catalogo_map()

    base = clave_cuenta(nombre_cuenta)
    return catalogo_map.get(base, {
        "codigo": "",
        "nombre": normalizar_nombre_cuenta(nombre_cuenta),
        "tipo": "",
        "naturaleza": "",
        "activa": True
    })


def obtener_naturaleza_real(nombre_cuenta, catalogo_map=None):
    info = obtener_info_cuenta_catalogo(nombre_cuenta, catalogo_map=catalogo_map)
    naturaleza = normalizar_txt(info.get("naturaleza"))
    tipo = normalizar_txt(info.get("tipo"))
    nombre = clave_cuenta(nombre_cuenta)

    if naturaleza in {"deudora", "acreedora"}:
        return naturaleza

    if tipo in {"activo", "gasto"}:
        return "deudora"

    if tipo in {"pasivo", "patrimonio", "ingreso"}:
        return "acreedora"

    if "venta" in nombre or "ventas" in nombre:
        return "acreedora"

    if (
        "por pagar" in nombre
        or "proveedor" in nombre
        or "acreedor" in nombre
        or "prestamo" in nombre
        or "préstamo" in nombre
        or "capital" in nombre
        or "venta" in nombre
        or "ventas" in nombre
        or "ingreso" in nombre
    ):
        return "acreedora"

    return "deudora"


def calcular_saldo_cuenta(total_debe, total_haber, naturaleza_real):
    total_debe = round(float(total_debe or 0), 2)
    total_haber = round(float(total_haber or 0), 2)

    if naturaleza_real == "deudora":
        return round(total_debe - total_haber, 2)

    return round(total_haber - total_debe, 2)


def desglosar_saldo(saldo, naturaleza_real):
    saldo = round(float(saldo or 0), 2)
    saldo_debe = 0.0
    saldo_haber = 0.0

    if saldo > 0:
        if naturaleza_real == "deudora":
            saldo_debe = saldo
        else:
            saldo_haber = saldo
    elif saldo < 0:
        if naturaleza_real == "deudora":
            saldo_haber = abs(saldo)
        else:
            saldo_debe = abs(saldo)

    return round(saldo_debe, 2), round(saldo_haber, 2)


def es_cuenta_pendiente_para_pago(nombre_cuenta, catalogo_map=None):
    info = obtener_info_cuenta_catalogo(nombre_cuenta, catalogo_map=catalogo_map)
    tipo = normalizar_txt(info.get("tipo"))
    base = clave_cuenta(nombre_cuenta)

    if tipo == "pasivo":
        return True

    return (
        "por pagar" in base
        or "proveedor" in base
        or "acreedor" in base
        or base.startswith("cuentas por pagar")
        or base.startswith("documentos por pagar")
    )


# =========================
# LÓGICA DEL MAYOR
# =========================
def es_partida_apertura_guardada(partida):
    glosa = str(partida.get("glosa", "")).strip().lower()
    codigo = str(partida.get("codigo", "")).strip().upper()
    tipo_partida = str(partida.get("tipo_partida", "")).strip().lower()

    if tipo_partida == "apertura":
        return True

    if glosa == "partida de apertura":
        return True

    if codigo in {"PDA 0", "PDA 1"}:
        cuentas = partida.get("cuentas", [])
        if any(str(c.get("cuenta", "")).startswith("A: ") for c in cuentas):
            return True

    return False


def obtener_apertura_y_movimientos_mes(
    ventas_data,
    compras_data,
    partidas_generadas_data,
    empresa_id,
    anio,
    mes
):
    empresa_id = normalizar_nit(empresa_id)
    anio = int(anio)
    mes = int(mes)

    # ENERO = usa partida de apertura normal
    if mes == 1:
        apertura = buscar_partida_apertura(ventas_data, empresa_id, anio)
        if not apertura:
            apertura = buscar_partida_apertura(compras_data, empresa_id, anio)

        apertura = normalizar_apertura_contrapartidas(apertura)

        registro_mes = buscar_partidas_generadas(
            partidas_generadas_data or [],
            empresa_id,
            anio,
            mes
        )

        partidas_mes = []
        if registro_mes and isinstance(registro_mes.get("partidas"), list):
            partidas_mes = registro_mes.get("partidas", [])

        if apertura:
            partidas_mes = [p for p in partidas_mes if not es_partida_apertura_guardada(p)]

        return apertura, partidas_mes

    # MESES > 1 = saldo inicial viene del saldo final del mes anterior
    mes_anterior = mes - 1

    apertura_anterior, partidas_mes_anterior = obtener_apertura_y_movimientos_mes(
        ventas_data=ventas_data,
        compras_data=compras_data,
        partidas_generadas_data=partidas_generadas_data,
        empresa_id=empresa_id,
        anio=anio,
        mes=mes_anterior
    )

    df_resumen_anterior = construir_resumen_mayor_mes(apertura_anterior, partidas_mes_anterior)
    apertura = construir_apertura_desde_saldo_final(df_resumen_anterior, anio, mes_anterior)

    registro_mes = buscar_partidas_generadas(
        partidas_generadas_data or [],
        empresa_id,
        anio,
        mes
    )

    partidas_mes = []
    if registro_mes and isinstance(registro_mes.get("partidas"), list):
        partidas_mes = registro_mes.get("partidas", [])

    return apertura, partidas_mes


def obtener_cuentas_historicas_hasta_mes(ventas_data, compras_data, partidas_generadas_data, empresa_id, anio, mes):
    empresa_id = normalizar_nit(empresa_id)
    anio = int(anio)
    mes = int(mes)

    cuentas = []
    vistos = set()

    def registrar_desde_partida(partida):
        if not isinstance(partida, dict):
            return
        for mov in partida.get("cuentas", []) or []:
            cuenta = normalizar_nombre_cuenta(mov.get("cuenta", ""))
            if not cuenta:
                continue
            key = clave_cuenta(cuenta)
            if key and key not in vistos:
                vistos.add(key)
                cuentas.append(cuenta)

    apertura = buscar_partida_apertura(ventas_data, empresa_id, anio)
    if not apertura:
        apertura = buscar_partida_apertura(compras_data, empresa_id, anio)
    apertura = normalizar_apertura_contrapartidas(apertura)
    registrar_desde_partida(apertura)

    for mes_iter in range(1, mes + 1):
        registro_mes = buscar_partidas_generadas(partidas_generadas_data or [], empresa_id, anio, mes_iter)
        if not registro_mes:
            continue
        for partida in registro_mes.get("partidas", []) or []:
            if mes_iter == 1 and es_partida_apertura_guardada(partida):
                continue
            registrar_desde_partida(partida)

    return cuentas


def construir_mayor_de_partidas(partidas):
    catalogo_map = obtener_catalogo_map()
    mayor = {}
    orden = []

    for partida in partidas or []:
        if not isinstance(partida, dict):
            continue

        for mov in partida.get("cuentas", []):
            cuenta_original = normalizar_nombre_cuenta(mov.get("cuenta", ""))
            cuenta_key = clave_cuenta(cuenta_original)
            debe = round(float(mov.get("debe", 0) or 0), 2)
            haber = round(float(mov.get("haber", 0) or 0), 2)

            if not cuenta_key:
                continue

            if cuenta_key not in mayor:
                info_catalogo = obtener_info_cuenta_catalogo(cuenta_original, catalogo_map=catalogo_map)
                cuenta_visible = str(info_catalogo.get("nombre", "")).strip() or cuenta_original
                naturaleza_real = obtener_naturaleza_real(cuenta_visible, catalogo_map=catalogo_map)

                mayor[cuenta_key] = {
                    "codigo": info_catalogo.get("codigo", ""),
                    "cuenta": cuenta_visible,
                    "tipo": info_catalogo.get("tipo", "") or "no_clasificada",
                    "naturaleza_catalogo": info_catalogo.get("naturaleza", "") or "",
                    "naturaleza_real": naturaleza_real,
                    "debe": 0.0,
                    "haber": 0.0,
                }
                orden.append(cuenta_key)

            mayor[cuenta_key]["debe"] += debe
            mayor[cuenta_key]["haber"] += haber

    salida = []
    for cuenta_key in orden:
        total_debe = round(mayor[cuenta_key]["debe"], 2)
        total_haber = round(mayor[cuenta_key]["haber"], 2)
        naturaleza_real = mayor[cuenta_key]["naturaleza_real"]
        saldo = calcular_saldo_cuenta(total_debe, total_haber, naturaleza_real)
        saldo_debe, saldo_haber = desglosar_saldo(saldo, naturaleza_real)

        salida.append({
            "codigo": mayor[cuenta_key]["codigo"],
            "cuenta": mayor[cuenta_key]["cuenta"],
            "tipo": mayor[cuenta_key]["tipo"],
            "naturaleza_catalogo": mayor[cuenta_key]["naturaleza_catalogo"],
            "naturaleza_real": naturaleza_real,
            "debe": total_debe,
            "haber": total_haber,
            "saldo": saldo,
            "saldo_debe": saldo_debe,
            "saldo_haber": saldo_haber,
        })

    return salida


def construir_resumen_mayor_mes(apertura, partidas_mes, cuentas_historicas=None):
    catalogo_map = obtener_catalogo_map()

    mayor_inicial = construir_mayor_de_partidas([apertura] if apertura else [])
    mayor_mov = construir_mayor_de_partidas(partidas_mes or [])
    mayor_final = construir_mayor_de_partidas(([apertura] if apertura else []) + (partidas_mes or []))

    ini_map = {clave_cuenta(x["cuenta"]): x for x in mayor_inicial}
    mov_map = {clave_cuenta(x["cuenta"]): x for x in mayor_mov}
    fin_map = {clave_cuenta(x["cuenta"]): x for x in mayor_final}

    todas = []
    for cuenta in cuentas_historicas or []:
        k = clave_cuenta(cuenta)
        if k and k not in todas:
            todas.append(k)

    for k in list(ini_map.keys()) + list(mov_map.keys()) + list(fin_map.keys()):
        if k and k not in todas:
            todas.append(k)

    cuentas_historicas_set = {clave_cuenta(x) for x in (cuentas_historicas or []) if str(x).strip()}

    filas = []
    for key in todas:
        ini = ini_map.get(key, {})
        mov = mov_map.get(key, {})
        fin = fin_map.get(key, {})

        cuenta = (
            fin.get("cuenta")
            or mov.get("cuenta")
            or ini.get("cuenta")
            or key
        )

        info_catalogo = obtener_info_cuenta_catalogo(cuenta, catalogo_map=catalogo_map)
        naturaleza_real = obtener_naturaleza_real(cuenta, catalogo_map=catalogo_map)

        ini_debe = round(float(ini.get("saldo_debe", 0) or 0), 2)
        ini_haber = round(float(ini.get("saldo_haber", 0) or 0), 2)

        mov_debe = round(float(mov.get("debe", 0) or 0), 2)
        mov_haber = round(float(mov.get("haber", 0) or 0), 2)

        fin_debe = round(float(fin.get("saldo_debe", 0) or 0), 2)
        fin_haber = round(float(fin.get("saldo_haber", 0) or 0), 2)

        incluir_fila = any([
            ini_debe, ini_haber, mov_debe, mov_haber, fin_debe, fin_haber
        ])

        if not incluir_fila and cuentas_historicas_set:
            incluir_fila = key in cuentas_historicas_set

        if incluir_fila:
            filas.append({
                "codigo": info_catalogo.get("codigo", ""),
                "cuenta": cuenta,
                "tipo": info_catalogo.get("tipo", "") or "no_clasificada",
                "naturaleza_catalogo": info_catalogo.get("naturaleza", "") or "",
                "naturaleza_real": naturaleza_real,
                "saldo_inicial_debe": ini_debe,
                "saldo_inicial_haber": ini_haber,
                "movimientos_debe": mov_debe,
                "movimientos_haber": mov_haber,
                "saldo_final_debe": fin_debe,
                "saldo_final_haber": fin_haber,
            })

    df = pd.DataFrame(filas)

    if not df.empty:
        fila_total = {
            "codigo": "",
            "cuenta": "TOTAL",
            "tipo": "",
            "naturaleza_catalogo": "",
            "naturaleza_real": "",
            "saldo_inicial_debe": round(pd.to_numeric(df["saldo_inicial_debe"], errors="coerce").fillna(0).sum(), 2),
            "saldo_inicial_haber": round(pd.to_numeric(df["saldo_inicial_haber"], errors="coerce").fillna(0).sum(), 2),
            "movimientos_debe": round(pd.to_numeric(df["movimientos_debe"], errors="coerce").fillna(0).sum(), 2),
            "movimientos_haber": round(pd.to_numeric(df["movimientos_haber"], errors="coerce").fillna(0).sum(), 2),
            "saldo_final_debe": round(pd.to_numeric(df["saldo_final_debe"], errors="coerce").fillna(0).sum(), 2),
            "saldo_final_haber": round(pd.to_numeric(df["saldo_final_haber"], errors="coerce").fillna(0).sum(), 2),
        }
        df = pd.concat([df, pd.DataFrame([fila_total])], ignore_index=True)

    return df


def evaluar_cuadre_resumen(df):
    if df is None or df.empty:
        return {
            "cuadra_inicial": True,
            "cuadra_movimientos": True,
            "cuadra_final": True,
            "totales": {}
        }

    fila_total = df.iloc[-1].to_dict() if str(df.iloc[-1].get("cuenta", "")).strip().upper() == "TOTAL" else None

    if fila_total:
        ini_debe = round(float(fila_total.get("saldo_inicial_debe", 0) or 0), 2)
        ini_haber = round(float(fila_total.get("saldo_inicial_haber", 0) or 0), 2)
        mov_debe = round(float(fila_total.get("movimientos_debe", 0) or 0), 2)
        mov_haber = round(float(fila_total.get("movimientos_haber", 0) or 0), 2)
        fin_debe = round(float(fila_total.get("saldo_final_debe", 0) or 0), 2)
        fin_haber = round(float(fila_total.get("saldo_final_haber", 0) or 0), 2)
    else:
        ini_debe = round(pd.to_numeric(df["saldo_inicial_debe"], errors="coerce").fillna(0).sum(), 2)
        ini_haber = round(pd.to_numeric(df["saldo_inicial_haber"], errors="coerce").fillna(0).sum(), 2)
        mov_debe = round(pd.to_numeric(df["movimientos_debe"], errors="coerce").fillna(0).sum(), 2)
        mov_haber = round(pd.to_numeric(df["movimientos_haber"], errors="coerce").fillna(0).sum(), 2)
        fin_debe = round(pd.to_numeric(df["saldo_final_debe"], errors="coerce").fillna(0).sum(), 2)
        fin_haber = round(pd.to_numeric(df["saldo_final_haber"], errors="coerce").fillna(0).sum(), 2)

    return {
        "cuadra_inicial": abs(ini_debe - ini_haber) <= 0.01,
        "cuadra_movimientos": abs(mov_debe - mov_haber) <= 0.01,
        "cuadra_final": abs(fin_debe - fin_haber) <= 0.01,
        "totales": {
            "ini_debe": ini_debe,
            "ini_haber": ini_haber,
            "mov_debe": mov_debe,
            "mov_haber": mov_haber,
            "fin_debe": fin_debe,
            "fin_haber": fin_haber,
        }
    }


def clasificar_cuenta_para_seguimiento(row, catalogo_map=None):
    cuenta = str(row.get("cuenta", "")).strip()
    if not cuenta or cuenta.upper() == "TOTAL":
        return None

    if catalogo_map is None:
        catalogo_map = obtener_catalogo_map()

    base = clave_cuenta(cuenta)
    tipo = normalizar_txt(row.get("tipo", ""))
    naturaleza_real = normalizar_txt(row.get("naturaleza_real", "")) or obtener_naturaleza_real(cuenta, catalogo_map=catalogo_map)

    fin_debe = round(float(row.get("saldo_final_debe", 0) or 0), 2)
    fin_haber = round(float(row.get("saldo_final_haber", 0) or 0), 2)

    if es_cuenta_pendiente_para_pago(cuenta, catalogo_map=catalogo_map):
        categoria = "Pago pendiente"
        monto = fin_haber if fin_haber > 0 else max(fin_debe, fin_haber)
        lado = "Haber" if fin_haber >= fin_debe else "Debe"
    elif (
        base.startswith("clientes")
        or "cliente" in base
        or "cuentas por cobrar" in base
        or "documentos por cobrar" in base
        or "deudores" in base
    ):
        categoria = "Cobro a clientes"
        monto = fin_debe if fin_debe > 0 else max(fin_debe, fin_haber)
        lado = "Debe" if fin_debe >= fin_haber else "Haber"
    elif "reintegro" in base:
        categoria = "Reintegro pendiente"
        if fin_debe == fin_haber:
            monto = fin_debe
            lado = "Debe" if naturaleza_real == "deudora" else "Haber"
        else:
            monto = max(fin_debe, fin_haber)
            lado = "Debe" if fin_debe >= fin_haber else "Haber"
    elif "anticipo" in base or "anticipos" in base:
        categoria = "Anticipo pendiente"
        monto = max(fin_debe, fin_haber)
        lado = "Debe" if fin_debe >= fin_haber else "Haber"
    elif "iva por cobrar" in base:
        categoria = "IVA por cobrar"
        monto = fin_debe if fin_debe > 0 else max(fin_debe, fin_haber)
        lado = "Debe" if fin_debe >= fin_haber else "Haber"
    elif (
        "credito de retencion iva" in base
        or "crédito de retención iva" in base
        or "credito retencion iva" in base
        or "crédito retención iva" in base
        or "retencion iva por acreditar" in base
        or "retención iva por acreditar" in base
    ):
        categoria = "Crédito de retención IVA"
        monto = fin_debe if fin_debe > 0 else max(fin_debe, fin_haber)
        lado = "Debe" if fin_debe >= fin_haber else "Haber"
    elif (
        tipo == "no_clasificada"
        or "regularizar" in base
        or "ajuste" in base
        or "diferencia" in base
        or "cuadre" in base
        or "reclas" in base
    ):
        categoria = "Revisión / ajustar"
        monto = max(fin_debe, fin_haber)
        lado = "Debe" if fin_debe >= fin_haber else "Haber"
    else:
        return None

    return {
        "categoria": categoria,
        "monto": round(float(monto or 0), 2),
        "lado": lado,
        "saldo_final_debe": fin_debe,
        "saldo_final_haber": fin_haber,
        "tipo": row.get("tipo", ""),
    }


def obtener_saldos_pendientes_desde_resumen(df, incluir_saldadas=False, categorias=None):
    if df is None or df.empty:
        return pd.DataFrame()

    catalogo_map = obtener_catalogo_map()
    categorias = set(categorias or [])
    filas = []

    for _, row in df.iterrows():
        cuenta = str(row.get("cuenta", "")).strip()
        if not cuenta or cuenta.upper() == "TOTAL":
            continue

        fin_debe = round(float(row.get("saldo_final_debe", 0) or 0), 2)
        fin_haber = round(float(row.get("saldo_final_haber", 0) or 0), 2)

        tuvo_movimiento = any([
            round(float(row.get("saldo_inicial_debe", 0) or 0), 2),
            round(float(row.get("saldo_inicial_haber", 0) or 0), 2),
            round(float(row.get("movimientos_debe", 0) or 0), 2),
            round(float(row.get("movimientos_haber", 0) or 0), 2),
            fin_debe,
            fin_haber,
        ])

        clasificacion = clasificar_cuenta_para_seguimiento(row, catalogo_map=catalogo_map)
        if not clasificacion:
            continue

        if categorias and clasificacion["categoria"] not in categorias:
            continue

        monto = round(float(clasificacion.get("monto", 0) or 0), 2)
        estado = "Pendiente" if monto > 0 else "Saldada"

        if monto > 0 or (incluir_saldadas and tuvo_movimiento):
            filas.append({
                "seleccionar": False,
                "codigo": row.get("codigo", ""),
                "cuenta": cuenta,
                "categoria": clasificacion.get("categoria", ""),
                "tipo": clasificacion.get("tipo", ""),
                "lado_saldo": clasificacion.get("lado", ""),
                "saldo_final_debe": clasificacion.get("saldo_final_debe", 0.0),
                "saldo_final_haber": clasificacion.get("saldo_final_haber", 0.0),
                "monto": monto,
                "estado": estado,
            })

    if not filas:
        return pd.DataFrame()

    df_out = pd.DataFrame(filas)
    df_out = df_out.sort_values(by=["categoria", "cuenta"], ascending=[True, True]).reset_index(drop=True)
    return df_out


def formatear_moneda(valor):
    return f"Q {float(valor or 0):,.2f}"


def obtener_lado_monto_saldo_visual(saldo_debe, saldo_haber):
    saldo_debe = round(float(saldo_debe or 0), 2)
    saldo_haber = round(float(saldo_haber or 0), 2)

    if abs(saldo_debe) <= 0.009 and abs(saldo_haber) <= 0.009:
        return "Sin saldo", 0.0

    if saldo_debe > 0 and saldo_haber <= 0:
        return "Debe", saldo_debe

    if saldo_haber > 0 and saldo_debe <= 0:
        return "Haber", saldo_haber

    if saldo_debe >= saldo_haber:
        return "Debe", saldo_debe

    return "Haber", saldo_haber


def obtener_alerta_saldo_contrario(row):
    cuenta = str(row.get("cuenta", "") or "").strip()
    if not cuenta or cuenta.upper() == "TOTAL":
        return None

    saldo_debe = round(float(row.get("saldo_final_debe", 0) or 0), 2)
    saldo_haber = round(float(row.get("saldo_final_haber", 0) or 0), 2)
    lado_actual, monto_actual = obtener_lado_monto_saldo_visual(saldo_debe, saldo_haber)

    if lado_actual == "Sin saldo" or monto_actual <= 0:
        return None

    tipo = normalizar_txt(row.get("tipo", ""))
    naturaleza_real = normalizar_txt(row.get("naturaleza_real", ""))

    if naturaleza_real == "deudora" and saldo_haber > 0:
        esperado = "Debe"
        actual = "Haber"
    elif naturaleza_real == "acreedora" and saldo_debe > 0:
        esperado = "Haber"
        actual = "Debe"
    else:
        return None

    tipo_legible = tipo if tipo else "cuenta"
    return {
        "cuenta": cuenta,
        "tipo": tipo_legible,
        "esperado": esperado,
        "actual": actual,
        "mensaje": f"Saldo contrario al tipo/naturaleza: {tipo_legible} normalmente cierra en {esperado} y aquí aparece en {actual}."
    }


def obtener_dataframe_alertas_saldo_contrario(df):
    if df is None or df.empty:
        return pd.DataFrame()

    filas = []
    for _, row in df.iterrows():
        alerta = obtener_alerta_saldo_contrario(row)
        if alerta:
            filas.append({
                "Cuenta": alerta["cuenta"],
                "Tipo": alerta["tipo"].title(),
                "Esperado": alerta["esperado"],
                "Actual": alerta["actual"],
                "Alerta": alerta["mensaje"],
            })

    if not filas:
        return pd.DataFrame()

    return pd.DataFrame(filas)


def render_tarjetas_cuadre(estado_cuadre):
    totales = (estado_cuadre or {}).get("totales", {}) or {}
    items = [
        {
            "titulo": "Saldo inicial",
            "cuadra": bool((estado_cuadre or {}).get("cuadra_inicial", True)),
            "debe": float(totales.get("ini_debe", 0) or 0),
            "haber": float(totales.get("ini_haber", 0) or 0),
        },
        {
            "titulo": "Movimientos del mes",
            "cuadra": bool((estado_cuadre or {}).get("cuadra_movimientos", True)),
            "debe": float(totales.get("mov_debe", 0) or 0),
            "haber": float(totales.get("mov_haber", 0) or 0),
        },
        {
            "titulo": "Saldo final",
            "cuadra": bool((estado_cuadre or {}).get("cuadra_final", True)),
            "debe": float(totales.get("fin_debe", 0) or 0),
            "haber": float(totales.get("fin_haber", 0) or 0),
        },
    ]

    cols = st.columns(3)
    for col, item in zip(cols, items):
        diferencia = abs(round(item["debe"] - item["haber"], 2))
        clase = "ok" if item["cuadra"] else "warn"
        estado = "Cuadra" if item["cuadra"] else "Revisar"

        with col:
            st.markdown(
                f"""
                <div class="cuadre-card {clase}">
                    <div class="cuadre-card-top">
                        <span class="cuadre-card-title">{item['titulo']}</span>
                        <span class="cuadre-card-pill">{estado}</span>
                    </div>
                    <div class="cuadre-card-values">
                        <div><small>Debe</small><strong>Q {item['debe']:,.2f}</strong></div>
                        <div><small>Haber</small><strong>Q {item['haber']:,.2f}</strong></div>
                    </div>
                    <div class="cuadre-card-diff">Diferencia: Q {diferencia:,.2f}</div>
                </div>
                """,
                unsafe_allow_html=True
            )


def obtener_lectura_cliente_cuenta(row):
    cuenta = str(row.get("cuenta", "") or "").strip()
    categoria = str(row.get("categoria", "") or "").strip()
    base = clave_cuenta(cuenta)
    monto = round(float(row.get("monto", 0) or 0), 2)
    estado = str(row.get("estado", "") or "").strip()

    resumen = ""
    significado = (
        "Esta cuenta presenta un saldo al cierre del período y conviene revisarla para confirmar "
        "si ya quedó resuelta o si todavía requiere movimiento."
    )
    accion = "Revisar documentos de respaldo y confirmar con el cliente cómo debe regularizarse."
    categoria_cliente = categoria or "Seguimiento"
    prioridad = "Media"

    if categoria == "Cobro a clientes" or base.startswith("clientes") or "cliente" in base:
        categoria_cliente = "Cobro / depósito pendiente"
        resumen = ""
        significado = (
            "Normalmente indica que hubo ventas o cobros registrados, pero el dinero no entró completo "
            "a banco o caja, o quedó un saldo pendiente de aplicar correctamente."
        )
        accion = (
            "Revisar ventas, recibos y boletas de depósito para confirmar cuánto falta ingresar al banco "
        )
        prioridad = "Alta" if monto > 0 else "Informativa"
    elif categoria == "Reintegro pendiente" or "reintegro" in base:
        categoria_cliente = "Reintegro pendiente"
        resumen = "Salida de dinero que todavía no quedó respaldada o reintegrada"
        significado = (
            "Indica que salió dinero de banco o caja y aún falta el reintegro, el respaldo documental "
            "o la aplicación correcta dentro de la contabilidad."
        )
        accion = (
            "Reintegrar al Banco, ya que fué un egreso demás no respaldado"
        )
        prioridad = "Alta" if monto > 0 else "Informativa"
    elif categoria == "Anticipo pendiente" or "anticipo" in base or "anticipos" in base:
        categoria_cliente = "Anticipo por aplicar"
        resumen = "Monto adelantado que aún no se ha aplicado por completo"
        significado = (
            "Corresponde a dinero entregado o recibido por anticipado que todavía debe relacionarse "
            "con una factura, compra, servicio o liquidación final."
        )
        accion = (
            "Verificar contra qué documento debe aplicarse y completar la liquidación o compensación del saldo."
        )
        prioridad = "Media" if monto > 0 else "Informativa"
    elif categoria == "Pago pendiente":
        categoria_cliente = "Pago pendiente"
        resumen = ""
        significado = (
            "Hay una obligación o cuenta por pagar que continúa pendiente al cierre del período y "
            "conviene darle seguimiento."
        )
        accion = "Validar fecha de pago, soporte y cuenta desde la cual se cancelará."
        prioridad = "Alta" if monto > 0 else "Informativa"
    elif categoria == "IVA por cobrar" or "iva por cobrar" in base:
        categoria_cliente = "IVA por cobrar"
        resumen = "Crédito fiscal pendiente de compensar o aplicar"
        significado = (
            "Esta cuenta refleja IVA crédito acumulado a favor de la empresa que todavía no se ha "
            "compensado por completo en la regularización o en períodos posteriores."
        )
        accion = "Revisar la regularización del IVA y confirmar cuánto queda pendiente de aplicar en los meses siguientes."
        prioridad = "Media" if monto > 0 else "Informativa"
    elif categoria == "Crédito de retención IVA" or "retencion iva" in base or "retención iva" in base:
        categoria_cliente = "Crédito de retención IVA"
        resumen = "Crédito de retención pendiente de aplicar"
        significado = (
            "Corresponde a retenciones de IVA acumuladas a favor que aún pueden utilizarse para "
            "rebajar pagos o regularizaciones futuras, según el control contable y fiscal."
        )
        accion = "Verificar el soporte de las retenciones y aplicar el saldo en la próxima regularización que corresponda."
        prioridad = "Media" if monto > 0 else "Informativa"
    elif categoria == "Revisión / ajustar":
        categoria_cliente = "Revisión administrativa"
        resumen = "Saldo poco usual que conviene aclarar"
        significado = (
            "Se trata de una cuenta que probablemente necesita ajuste, reclasificación o confirmación "
            "para que el cierre quede más claro y ordenado."
        )
        accion = "Revisar con administración o contabilidad qué ajuste corresponde realizar."
        prioridad = "Media" if monto > 0 else "Informativa"

    estado_cliente = "Requiere seguimiento" if (estado.lower() == "pendiente" or monto > 0) else "Regularizada"

    return {
        "categoria_cliente": categoria_cliente,
        "resumen": resumen,
        "significado": significado,
        "accion": accion,
        "estado_cliente": estado_cliente,
        "prioridad": prioridad,
    }


def obtener_accion_breve_cliente(row):
    lectura = obtener_lectura_cliente_cuenta(row)
    accion = str((row.get("que_hacer", "") if hasattr(row, "get") else "") or "").strip()
    if accion:
        return accion
    return str(lectura.get("accion", "Revisar y regularizar esta cuenta.") or "Revisar y regularizar esta cuenta.").strip()


def obtener_resumen_breve_cliente(row):
    lectura = obtener_lectura_cliente_cuenta(row)
    return str(lectura.get("resumen", "") or "").strip()


def obtener_logo_data_uri(uploaded_file):
    if uploaded_file is None:
        return ""

    try:
        contenido = uploaded_file.getvalue()
        if not contenido:
            return ""
        mime = getattr(uploaded_file, "type", "") or "image/png"
        return f"data:{mime};base64,{base64.b64encode(contenido).decode('utf-8')}"
    except Exception:
        return ""



def construir_html_reporte_pendientes(
    empresa,
    empresa_id,
    anio,
    mes,
    df_seleccionadas,
    titulo_reporte,
    observaciones="",
    logo_data_uri=""
):
    fecha_emision = datetime.now().strftime("%d/%m/%Y %H:%M")
    periodo = f"{MESES.get(int(mes), mes)} {int(anio)}"
    observaciones_html = html_lib.escape(str(observaciones or "").strip()).replace("\n", "<br>")

    total = 0.0
    filas_html = ""

    for _, row in df_seleccionadas.iterrows():
        cuenta = html_lib.escape(str(row.get("cuenta", "") or ""))
        monto = round(float(row.get("monto", 0) or 0), 2)
        total += monto
        que_hacer = html_lib.escape(obtener_accion_breve_cliente(row))
        resumen = html_lib.escape(obtener_resumen_breve_cliente(row))

        filas_html += f"""
        <tr class="fila-reporte">
            <td class="celda cuenta-celda">
                <div class="td-label">Cuenta</div>
                <div class="cuenta">{cuenta}</div>
                <div class="detalle">{resumen}</div>
            </td>
            <td class="celda monto-celda">
                <div class="td-label">Monto</div>
                <div class="monto">{formatear_moneda(monto)}</div>
            </td>
            <td class="celda accion-celda">
                <div class="td-label">¿Qué hacer?</div>
                <div class="accion">{que_hacer}</div>
            </td>
        </tr>
        """

    if not filas_html:
        filas_html = '<tr><td colspan="3" class="empty">No hay cuentas seleccionadas para el reporte.</td></tr>'

    if logo_data_uri:
        logo_html = f'<img src="{logo_data_uri}" alt="Logotipo" class="logo-img">'
    else:
        iniciales = "".join([p[:1].upper() for p in str(empresa or "Empresa").split()[:3]]) or "EMP"
        logo_html = f'<div class="logo-fallback">{html_lib.escape(iniciales)}</div>'

    total_cuentas = len(df_seleccionadas)

    observaciones_block = ""
    if observaciones_html:
        observaciones_block = f'<div class="mensaje"><strong>Nota:</strong> {observaciones_html}</div>'

    return f"""
<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="UTF-8" />
    <meta name="viewport" content="width=device-width, initial-scale=1.0" />
    <title>{html_lib.escape(str(titulo_reporte or 'Reporte de seguimiento'))}</title>
    <style>
    * {{ box-sizing: border-box; }}

    body {{
        font-family: Arial, Helvetica, sans-serif;
        margin: 0;
        background: #f5f7fb;
        color: #1f2937;
    }}

    .page {{
        max-width: 920px;
        margin: 24px auto;
        background: #ffffff;
        border-radius: 18px;
        box-shadow: 0 10px 30px rgba(15, 23, 42, 0.08);
        overflow: hidden;
    }}

    .header {{
        padding: 24px 28px 18px;
        border-bottom: 1px solid #eef2f7;
    }}

    .header-top {{
        display: flex;
        justify-content: space-between;
        align-items: center;
        gap: 18px;
    }}

    .brand {{
        display: flex;
        align-items: center;
        gap: 14px;
        min-width: 0;
    }}

    .logo-img {{
        width: 60px;
        height: 60px;
        object-fit: contain;
        border-radius: 14px;
        border: 1px solid #e5e7eb;
        background: #fff;
        padding: 6px;
        flex-shrink: 0;
    }}

    .logo-fallback {{
        width: 60px;
        height: 60px;
        border-radius: 14px;
        background: #0f172a;
        color: #fff;
        display: flex;
        align-items: center;
        justify-content: center;
        font-size: 20px;
        font-weight: 700;
        flex-shrink: 0;
    }}

    h1 {{
        margin: 0;
        font-size: 24px;
        color: #0f172a;
        line-height: 1.2;
    }}

    .sub {{
        margin: 4px 0 0 0;
        font-size: 13px;
        color: #64748b;
        line-height: 1.5;
    }}

    .meta {{
        text-align: right;
        font-size: 13px;
        color: #475569;
        line-height: 1.6;
    }}

    .content {{
        padding: 18px 28px 26px;
    }}

    .kpis {{
        display: flex;
        flex-wrap: wrap;
        gap: 10px;
        margin-bottom: 14px;
    }}

    .chip {{
        background: #f8fafc;
        border: 1px solid #e2e8f0;
        color: #0f172a;
        border-radius: 999px;
        padding: 8px 12px;
        font-size: 13px;
        font-weight: 600;
    }}

    .mensaje {{
        background: #f8fafc;
        border: 1px solid #e2e8f0;
        color: #334155;
        border-radius: 12px;
        padding: 12px 14px;
        margin-bottom: 14px;
        font-size: 13px;
        line-height: 1.5;
    }}

    .tabla-wrap {{
        border: 1px solid #e5e7eb;
        border-radius: 16px;
        overflow: hidden;
        background: #fff;
    }}

    table {{
        width: 100%;
        border-collapse: collapse;
    }}

    thead th {{
        background: #f8fafc;
        color: #334155;
        text-align: left;
        font-size: 12px;
        text-transform: uppercase;
        letter-spacing: .4px;
        padding: 14px 16px;
        border-bottom: 1px solid #e5e7eb;
    }}

    tbody td {{
        padding: 14px 16px;
        border-bottom: 1px solid #eef2f7;
        vertical-align: top;
        font-size: 14px;
        color: #1f2937;
    }}

    tbody tr:last-child td {{
        border-bottom: none;
    }}

    .td-label {{
        display: none;
    }}

    .cuenta {{
        font-weight: 700;
        color: #0f172a;
        margin-bottom: 4px;
        line-height: 1.35;
        word-break: break-word;
    }}

    .detalle {{
        font-size: 12px;
        color: #64748b;
        line-height: 1.5;
        word-break: break-word;
    }}

    .monto {{
        white-space: nowrap;
        font-weight: 800;
        color: #0f172a;
        font-size: 15px;
    }}

    .accion {{
        line-height: 1.55;
        color: #334155;
        word-break: break-word;
    }}

    .empty {{
        text-align: center;
        color: #64748b;
        padding: 22px;
    }}

    .foot {{
        padding: 10px 28px 24px;
        display: flex;
        flex-direction: column;
        align-items: center;
        gap: 10px;
    }}

    .footer-copy {{
        font-size: 12px;
        color: #64748b;
        text-align: center;
        line-height: 1.5;
    }}

    .btn-print {{
        width: 100%;
        border: none;
        background: #0f172a;
        color: #fff;
        border-radius: 12px;
        padding: 13px 16px;
        font-size: 14px;
        font-weight: 700;
        cursor: pointer;
    }}

    @media print {{
        body {{
            background: #fff;
        }}

        .page {{
            margin: 0;
            max-width: none;
            border-radius: 0;
            box-shadow: none;
        }}

        .header, .content {{
            padding-left: 18px;
            padding-right: 18px;
        }}

        .btn-print {{
            display: none;
        }}

        .foot {{
            display: block;
            padding: 12px 18px 0;
        }}

        .footer-copy {{
            text-align: center;
            color: #64748b;
        }}
    }}

    @media (max-width: 768px) {{
        body {{
            background: #eef2f7;
        }}

        .page {{
            margin: 0;
            border-radius: 0;
            box-shadow: none;
            min-height: 100vh;
        }}

        .header {{
            padding: 18px 16px 14px;
        }}

        .header-top {{
            flex-direction: column;
            align-items: flex-start;
            gap: 14px;
        }}

        .brand {{
            width: 100%;
            align-items: flex-start;
        }}

        .brand > div {{
            min-width: 0;
        }}

        h1 {{
            font-size: 20px;
        }}

        .sub {{
            font-size: 12px;
        }}

        .meta {{
            text-align: left;
            width: 100%;
            background: #f8fafc;
            border: 1px solid #e2e8f0;
            border-radius: 12px;
            padding: 10px 12px;
        }}

        .content {{
            padding: 14px 16px 22px;
        }}

        .kpis {{
            gap: 8px;
            margin-bottom: 12px;
        }}

        .chip {{
            font-size: 12px;
            padding: 7px 10px;
        }}

        .mensaje {{
            font-size: 12px;
            padding: 10px 12px;
            margin-bottom: 12px;
        }}

        .tabla-wrap {{
            border: none;
            border-radius: 0;
            overflow: visible;
            background: transparent;
        }}

        table,
        thead,
        tbody,
        th,
        td,
        tr {{
            display: block;
            width: 100%;
        }}

        thead {{
            display: none;
        }}

        tbody {{
            display: flex;
            flex-direction: column;
            gap: 12px;
        }}

        tbody tr.fila-reporte {{
            background: #ffffff;
            border: 1px solid #e5e7eb;
            border-radius: 16px;
            padding: 12px;
            box-shadow: 0 4px 14px rgba(15, 23, 42, 0.05);
        }}

        tbody td {{
            border: none;
            padding: 0;
            margin: 0 0 12px 0;
        }}

        tbody td:last-child {{
            margin-bottom: 0;
        }}

        .td-label {{
            display: block;
            font-size: 11px;
            font-weight: 700;
            text-transform: uppercase;
            letter-spacing: .4px;
            color: #64748b;
            margin-bottom: 6px;
        }}

        .monto-celda {{
            background: #f8fafc;
            border: 1px solid #e2e8f0;
            border-radius: 12px;
            padding: 10px 12px;
        }}

        .monto {{
            font-size: 18px;
            white-space: normal;
        }}

        .accion-celda {{
            border-top: 1px dashed #e2e8f0;
            padding-top: 12px;
        }}

        .foot {{
            padding: 8px 16px 20px;
        }}

        .btn-print {{
            border-radius: 14px;
            padding: 14px 16px;
        }}

        .footer-copy {{
            font-size: 11px;
        }}
    }}
    </style>
</head>
<body>
    <div class="page">
        <div class="header">
            <div class="header-top">
                <div class="brand">
                    {logo_html}
                    <div>
                        <h1>{html_lib.escape(str(titulo_reporte or 'Reporte de seguimiento'))}</h1>
                        <p class="sub">Resumen breve de cuentas que conviene revisar.</p>
                    </div>
                </div>
                <div class="meta">
                    <div><strong>Empresa:</strong> {html_lib.escape(str(empresa or ''))}</div>
                    <div><strong>Período:</strong> {html_lib.escape(str(periodo))}</div>
                    <div><strong>Emitido:</strong> {html_lib.escape(str(fecha_emision))}</div>
                </div>
            </div>
        </div>

        <div class="content">
            <div class="kpis">
                <div class="chip">Cuentas: {total_cuentas}</div>
                <div class="chip">NIT: {html_lib.escape(str(empresa_id or ''))}</div>
            </div>

            {observaciones_block}

            <div class="tabla-wrap">
                <table>
                    <thead>
                        <tr>
                            <th>Cuenta</th>
                            <th>Monto</th>
                            <th>¿Qué hacer?</th>
                        </tr>
                    </thead>
                    <tbody>
                        {filas_html}
                    </tbody>
                </table>
            </div>
        </div>

        <div class="foot">
            <div class="footer-copy">Emmanuel's, Desarrollo de Sistemas Empresariales, Correo: xarjhonnatan4@gmail.com</div>
            <button class="btn-print" onclick="window.print()">Imprimir reporte</button>
        </div>
    </div>
</body>
</html>
    """

def dataframe_partidas_fuente(apertura, partidas_mes):
    filas = []

    if apertura:
        filas.append({
            "Bloque": "SALDO INICIAL DEL MES",
            "Partida": str(apertura.get("codigo", "")),
            "Fecha": str(apertura.get("fecha", "")),
            "Cuenta": "",
            "Debe": "",
            "Haber": "",
            "Glosa": str(apertura.get("glosa", ""))
        })
        for mov in apertura.get("cuentas", []):
            filas.append({
                "Bloque": "",
                "Partida": "",
                "Fecha": "",
                "Cuenta": str(mov.get("cuenta", "")),
                "Debe": float(mov.get("debe", 0) or 0) if float(mov.get("debe", 0) or 0) else "",
                "Haber": float(mov.get("haber", 0) or 0) if float(mov.get("haber", 0) or 0) else "",
                "Glosa": ""
            })

    for p in partidas_mes or []:
        filas.append({
            "Bloque": "MOVIMIENTOS DEL MES",
            "Partida": str(p.get("codigo", "")),
            "Fecha": str(p.get("fecha", "")),
            "Cuenta": "",
            "Debe": "",
            "Haber": "",
            "Glosa": str(p.get("glosa", ""))
        })
        for mov in p.get("cuentas", []):
            filas.append({
                "Bloque": "",
                "Partida": "",
                "Fecha": "",
                "Cuenta": str(mov.get("cuenta", "")),
                "Debe": float(mov.get("debe", 0) or 0) if float(mov.get("debe", 0) or 0) else "",
                "Haber": float(mov.get("haber", 0) or 0) if float(mov.get("haber", 0) or 0) else "",
                "Glosa": ""
            })

    return pd.DataFrame(filas)




def aplicar_estilo_colores(df):
    if df is None or df.empty:
        return df

    def estilo(row):
        estilos = []
        for col in df.columns:
            if "saldo_inicial" in col:
                estilos.append("background-color: #E0F2FE")  # azul claro
            elif "movimientos" in col:
                estilos.append("background-color: #FEF3C7")  # amarillo claro
            elif "saldo_final" in col:
                estilos.append("background-color: #DCFCE7")  # verde claro
            else:
                estilos.append("")
        return estilos

    styler = df.style.apply(estilo, axis=1)

    # resaltar TOTAL
    if "cuenta" in df.columns:
        def resaltar_total(row):
            if str(row.get("cuenta", "")).strip().upper() == "TOTAL":
                return ["font-weight: bold; background-color: #E5E7EB"] * len(row)
            return [""] * len(row)

        styler = styler.apply(resaltar_total, axis=1)

    
    # formato monetario Q
    cols_monetarias = [c for c in df.columns if any(k in c for k in ["debe","haber"])]
    styler = styler.format({col: "Q {:,.2f}" for col in cols_monetarias})

    return styler

COLUMNAS_MAYOR_EXPORTAR = [
    ("partidas_movimiento", "PDA"),
    ("cuenta", "Cuenta"),
    ("saldo_inicial_debe", "Saldo inicial debe"),
    ("saldo_inicial_haber", "Saldo inicial haber"),
    ("movimientos_debe", "Movimientos debe"),
    ("movimientos_haber", "Movimientos haber"),
    ("saldo_final_debe", "Saldo final debe"),
    ("saldo_final_haber", "Saldo final haber"),
]


def numero_partida_excel(codigo_partida):
    texto = str(codigo_partida or "").strip()
    match = re.search(r"\d+", texto)
    return match.group(0) if match else texto


def obtener_partidas_por_cuenta_mes(partidas_mes):
    partidas_por_cuenta = {}

    for partida in partidas_mes or []:
        if not isinstance(partida, dict):
            continue
        codigo_partida = numero_partida_excel(partida.get("codigo", ""))
        if not codigo_partida:
            continue

        for mov in partida.get("cuentas", []) or []:
            cuenta_key = clave_cuenta(mov.get("cuenta", ""))
            debe = round(float(mov.get("debe", 0) or 0), 2)
            haber = round(float(mov.get("haber", 0) or 0), 2)
            if not cuenta_key or (debe == 0 and haber == 0):
                continue

            partidas = partidas_por_cuenta.setdefault(cuenta_key, [])
            if codigo_partida not in partidas:
                partidas.append(codigo_partida)

    return partidas_por_cuenta


def preparar_df_mayor_excel(df, partidas_mes=None):
    if df is None or df.empty:
        return pd.DataFrame(columns=[nombre for _, nombre in COLUMNAS_MAYOR_EXPORTAR])

    salida = df.copy()
    for columna, _ in COLUMNAS_MAYOR_EXPORTAR:
        if columna not in salida.columns:
            salida[columna] = ""

    partidas_por_cuenta = obtener_partidas_por_cuenta_mes(partidas_mes)
    if partidas_por_cuenta:
        salida["partidas_movimiento"] = salida["cuenta"].apply(
            lambda cuenta: ", ".join(partidas_por_cuenta.get(clave_cuenta(cuenta), []))
        )

    salida = salida[[columna for columna, _ in COLUMNAS_MAYOR_EXPORTAR]].copy()
    salida.columns = [nombre for _, nombre in COLUMNAS_MAYOR_EXPORTAR]

    columnas_monto = [c for c in salida.columns if c not in {"PDA", "Cuenta"}]
    for columna in columnas_monto:
        salida[columna] = pd.to_numeric(salida[columna], errors="coerce").fillna(0).round(2)

    return salida


def nombre_archivo_seguro(texto):
    limpio = normalizar_txt(texto).replace(" ", "_")
    limpio = "".join(c for c in limpio if c.isalnum() or c in {"_", "-"})
    return limpio or "empresa"


def aplicar_bordes_rango(ws, fila_inicio, fila_fin, col_inicio, col_fin, borde):
    for fila in range(fila_inicio, fila_fin + 1):
        for col in range(col_inicio, col_fin + 1):
            ws.cell(row=fila, column=col).border = borde


def separar_total_mayor_excel(df_excel):
    if df_excel is None or df_excel.empty:
        return df_excel, pd.DataFrame(columns=df_excel.columns if df_excel is not None else [])

    es_total = df_excel["Cuenta"].astype(str).str.strip().str.upper() == "TOTAL"
    df_detalle = df_excel[~es_total].copy()
    df_total = df_excel[es_total].copy()
    return df_detalle, df_total


def construir_orden_cuentas_mayor_excel(bloques_mensuales):
    orden = []
    vistas = set()

    for bloque in bloques_mensuales or []:
        df_excel = preparar_df_mayor_excel(bloque.get("df"), partidas_mes=bloque.get("partidas_mes"))
        df_detalle, _ = separar_total_mayor_excel(df_excel)
        for _, row in df_detalle.iterrows():
            cuenta = str(row.get("Cuenta", "") or "").strip()
            cuenta_key = clave_cuenta(cuenta)
            if cuenta_key and cuenta_key not in vistas:
                vistas.add(cuenta_key)
                orden.append((cuenta_key, cuenta))

    return orden


def construir_filas_detalle_ordenadas(df_detalle, encabezados, orden_cuentas=None):
    if not orden_cuentas:
        return [row.to_dict() for _, row in df_detalle.iterrows()]

    filas_por_cuenta = {}
    for _, row in df_detalle.iterrows():
        fila = row.to_dict()
        cuenta_key = clave_cuenta(fila.get("Cuenta", ""))
        if cuenta_key and cuenta_key not in filas_por_cuenta:
            filas_por_cuenta[cuenta_key] = fila

    filas = []
    for cuenta_key, _ in orden_cuentas:
        filas.append(
            filas_por_cuenta.get(
                cuenta_key,
                {encabezado: "" for encabezado in encabezados},
            )
        )

    return filas


def escribir_bloque_mayor_excel(ws, df, empresa, empresa_id, anio, mes, col_inicio, filas_detalle_objetivo=None, folio=None, orden_cuentas=None, partidas_mes=None):
    df_excel = preparar_df_mayor_excel(df, partidas_mes=partidas_mes)
    df_detalle, df_total = separar_total_mayor_excel(df_excel)
    filas_detalle_objetivo = max(int(filas_detalle_objetivo or 0), len(orden_cuentas or []), len(df_detalle))
    encabezados = list(df_excel.columns)
    col_fin = col_inicio + len(encabezados) - 1

    titulo_fill = PatternFill("solid", fgColor="1D4ED8")
    header_fill = PatternFill("solid", fgColor="DBEAFE")
    saldo_ini_fill = PatternFill("solid", fgColor="E0F2FE")
    mov_fill = PatternFill("solid", fgColor="FEF3C7")
    saldo_fin_fill = PatternFill("solid", fgColor="DCFCE7")
    total_fill = PatternFill("solid", fgColor="E5E7EB")
    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    ws.merge_cells(start_row=1, start_column=col_inicio, end_row=1, end_column=col_fin)
    celda_titulo = ws.cell(row=1, column=col_inicio)
    celda_titulo.value = f"Mayor - {MESES[int(mes)]} {int(anio)}"
    celda_titulo.font = Font(bold=True, color="FFFFFF", size=12)
    celda_titulo.fill = titulo_fill
    celda_titulo.alignment = center

    ws.merge_cells(start_row=2, start_column=col_inicio, end_row=2, end_column=col_fin)
    celda_empresa = ws.cell(row=2, column=col_inicio)
    celda_empresa.value = f"{empresa} | NIT: {normalizar_nit(empresa_id)}"
    celda_empresa.font = Font(bold=True, color="1F2937")
    celda_empresa.alignment = center

    ws.merge_cells(start_row=3, start_column=col_inicio, end_row=3, end_column=col_fin)
    celda_folio = ws.cell(row=3, column=col_inicio)
    celda_folio.value = f"Folio: {folio}" if folio is not None else "Folio:"
    celda_folio.font = Font(bold=True, color="1F2937")
    celda_folio.alignment = center

    fila_encabezado = 4
    for offset, encabezado in enumerate(encabezados):
        col = col_inicio + offset
        celda = ws.cell(row=fila_encabezado, column=col, value=encabezado)
        celda.font = Font(bold=True, color="111827")
        celda.fill = header_fill
        celda.alignment = center
        celda.border = border

    filas_exportar = []
    filas_exportar.extend(
        construir_filas_detalle_ordenadas(
            df_detalle,
            encabezados,
            orden_cuentas=orden_cuentas,
        )
    )

    for _ in range(filas_detalle_objetivo - len(filas_exportar)):
        filas_exportar.append({encabezado: "" for encabezado in encabezados})

    for _, row in df_total.iterrows():
        filas_exportar.append(row.to_dict())

    for idx, row in enumerate(filas_exportar):
        fila = fila_encabezado + 1 + idx
        es_total = str(row.get("Cuenta", "")).strip().upper() == "TOTAL"
        for offset, encabezado in enumerate(encabezados):
            col = col_inicio + offset
            valor = row.get(encabezado, "")
            celda = ws.cell(row=fila, column=col, value=valor)
            celda.border = border
            celda.alignment = left if encabezado == "Cuenta" else center

            if encabezado not in {"PDA", "Cuenta"}:
                celda.number_format = '"Q" #,##0.00'

            if "Saldo inicial" in encabezado:
                celda.fill = saldo_ini_fill
            elif "Movimientos" in encabezado:
                celda.fill = mov_fill
            elif "Saldo final" in encabezado:
                celda.fill = saldo_fin_fill

            if es_total:
                celda.font = Font(bold=True)
                celda.fill = total_fill

    fila_fin = fila_encabezado + max(len(filas_exportar), 1)
    aplicar_bordes_rango(ws, 1, fila_fin, col_inicio, col_fin, border)

    anchos = {
        "PDA": 16,
        "Cuenta": 30,
        "Saldo inicial debe": 16,
        "Saldo inicial haber": 16,
        "Movimientos debe": 16,
        "Movimientos haber": 16,
        "Saldo final debe": 16,
        "Saldo final haber": 16,
    }
    for offset, encabezado in enumerate(encabezados):
        col = col_inicio + offset
        letra = get_column_letter(col)
        ws.column_dimensions[letra].width = anchos.get(encabezado, 14)

    return col_fin + 1


def exportar_mayor_excel(empresa, empresa_id, anio, bloques_mensuales, folio_inicial=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Mayor"
    ws.freeze_panes = "A5"

    orden_cuentas = construir_orden_cuentas_mayor_excel(bloques_mensuales)
    filas_detalle_objetivo = len(orden_cuentas)
    for bloque in bloques_mensuales:
        df_excel = preparar_df_mayor_excel(bloque.get("df"), partidas_mes=bloque.get("partidas_mes"))
        df_detalle, _ = separar_total_mayor_excel(df_excel)
        filas_detalle_objetivo = max(filas_detalle_objetivo, len(df_detalle))

    col_inicio = 1
    for idx, bloque in enumerate(bloques_mensuales):
        folio = int(folio_inicial) + idx if folio_inicial is not None else None
        col_inicio = escribir_bloque_mayor_excel(
            ws=ws,
            df=bloque.get("df"),
            empresa=empresa,
            empresa_id=empresa_id,
            anio=anio,
            mes=bloque.get("mes"),
            col_inicio=col_inicio,
            filas_detalle_objetivo=filas_detalle_objetivo,
            folio=folio,
            orden_cuentas=orden_cuentas,
            partidas_mes=bloque.get("partidas_mes"),
        )

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def construir_bloques_mayor_hasta_mes(ventas_data, compras_data, partidas_generadas_data, empresa_id, anio, mes_fin):
    bloques = []
    for mes_iter in range(1, int(mes_fin) + 1):
        apertura_iter, partidas_mes_iter = obtener_apertura_y_movimientos_mes(
            ventas_data=ventas_data,
            compras_data=compras_data,
            partidas_generadas_data=partidas_generadas_data,
            empresa_id=empresa_id,
            anio=int(anio),
            mes=int(mes_iter),
        )
        cuentas_historicas_iter = obtener_cuentas_historicas_hasta_mes(
            ventas_data=ventas_data,
            compras_data=compras_data,
            partidas_generadas_data=partidas_generadas_data,
            empresa_id=empresa_id,
            anio=int(anio),
            mes=int(mes_iter),
        )
        df_iter = construir_resumen_mayor_mes(
            apertura_iter,
            partidas_mes_iter,
            cuentas_historicas=cuentas_historicas_iter,
        )
        bloques.append({"mes": mes_iter, "df": df_iter, "partidas_mes": partidas_mes_iter})

    return bloques


# =========================
# UI
# =========================
st.set_page_config(page_title="Mayor", page_icon="📘", layout="wide")

st.markdown("""
<style>
.block-container {
    padding-top: 2rem;
    padding-bottom: 1rem;
    max-width: 100%;
}
.panel {
    background: #ffffff;
    border: 1px solid #e5e7eb;
    border-radius: 16px;
    padding: 18px;
    margin-bottom: 14px;
    box-shadow: 0 2px 10px rgba(0,0,0,0.04);
}
.panel-title {
    font-size: 1.05rem;
    font-weight: 700;
    margin-bottom: 12px;
    color: #111827;
}
.badge-top {
    text-align: right;
    font-size: 0.95rem;
    font-weight: 700;
    color: #2563eb;
    background: #eff6ff;
    border: 1px solid #bfdbfe;
    padding: 6px 12px;
    border-radius: 999px;
    display: inline-block;
}
.cuadre-card {
    background: #ffffff;
    border: 1px solid #e5e7eb;
    border-left: 5px solid #cbd5e1;
    border-radius: 14px;
    padding: 14px 14px 12px;
    margin: 6px 0 10px;
    box-shadow: 0 2px 8px rgba(15, 23, 42, 0.04);
}
.cuadre-card.ok {
    border-left-color: #16a34a;
    background: #f8fff9;
}
.cuadre-card.warn {
    border-left-color: #dc2626;
    background: #fff9f9;
}
.cuadre-card-top {
    display: flex;
    justify-content: space-between;
    align-items: center;
    gap: 10px;
    margin-bottom: 12px;
}
.cuadre-card-title {
    font-size: 0.92rem;
    font-weight: 700;
    color: #0f172a;
}
.cuadre-card-pill {
    font-size: 0.78rem;
    font-weight: 700;
    padding: 4px 10px;
    border-radius: 999px;
    background: #e2e8f0;
    color: #334155;
}
.cuadre-card.ok .cuadre-card-pill {
    background: #dcfce7;
    color: #166534;
}
.cuadre-card.warn .cuadre-card-pill {
    background: #fee2e2;
    color: #991b1b;
}
.cuadre-card-values {
    display: grid;
    grid-template-columns: repeat(2, minmax(0, 1fr));
    gap: 10px;
}
.cuadre-card-values small {
    display: block;
    color: #64748b;
    font-size: 0.76rem;
    margin-bottom: 4px;
}
.cuadre-card-values strong {
    display: block;
    color: #0f172a;
    font-size: 1rem;
}
.cuadre-card-diff {
    margin-top: 10px;
    padding-top: 10px;
    border-top: 1px dashed #e2e8f0;
    color: #475569;
    font-size: 0.82rem;
    font-weight: 600;
}
</style>
""", unsafe_allow_html=True)

ventas_data = cargar_json_lista(VENTAS_FILE)
compras_data = cargar_json_lista(COMPRAS_FILE)
partidas_generadas_data = cargar_json_lista(PARTIDAS_GENERADAS_FILE)
empresas_data = cargar_empresas()
cierres_parciales_data = cargar_cierres_parciales()

c1, c2 = st.columns([2, 1])
with c1:
    st.markdown("## 📘 Mayor mensual")

st.markdown('</div>', unsafe_allow_html=True)

empresas_map = {}

for x in (ventas_data + compras_data):
    empresa_id = obtener_id_empresa_desde_partida(x, empresas_data)
    nombre_visible = str(x.get("empresa", "")).strip()

    if empresa_id and empresa_id not in empresas_map:
        emp = buscar_empresa_por_nit(empresa_id, empresas_data)
        empresas_map[empresa_id] = str(emp.get("nombre", nombre_visible)).strip() if emp else nombre_visible

    elif not empresa_id and nombre_visible:
        emp = buscar_empresa_por_nombre(nombre_visible, empresas_data)
        if emp:
            nit_emp = normalizar_nit(emp.get("nit"))
            if nit_emp and nit_emp not in empresas_map:
                empresas_map[nit_emp] = str(emp.get("nombre", nombre_visible)).strip()

empresas = sorted(
    [{"id": k, "nombre": v} for k, v in empresas_map.items()],
    key=lambda x: x["nombre"].lower()
)

if not empresas:
    st.warning("No se encontraron empresas válidas.")
    st.stop()

anios = sorted(set(
    int(x.get("anio", 0) or 0)
    for x in (ventas_data + compras_data)
    if int(x.get("anio", 0) or 0)
))

if not anios:
    st.warning("No se encontraron años válidos.")
    st.stop()

with st.container():
    col1, col2, col3 = st.columns(3)

    with col1:
        empresa_sel = st.selectbox("Empresa", empresas, format_func=lambda x: x["nombre"])

    with col2:
        anio = st.selectbox("Año", anios, index=len(anios) - 1 if anios else 0)

    with col3:
        mes = st.selectbox("Mes", list(MESES.keys()), format_func=lambda x: MESES[x])

regimen_empresa = obtener_regimen_empresa(empresa_sel["id"], empresas_data)
trimestre_actual = trimestre_desde_mes(mes)
es_cierre_trimestral_actual = es_mes_cierre_trimestral(mes)

apertura, partidas_mes = obtener_apertura_y_movimientos_mes(
    ventas_data=ventas_data,
    compras_data=compras_data,
    partidas_generadas_data=partidas_generadas_data,
    empresa_id=empresa_sel["id"],
    anio=int(anio),
    mes=int(mes)
)

cuentas_historicas = obtener_cuentas_historicas_hasta_mes(
    ventas_data=ventas_data,
    compras_data=compras_data,
    partidas_generadas_data=partidas_generadas_data,
    empresa_id=empresa_sel["id"],
    anio=int(anio),
    mes=int(mes)
)

df_resumen = construir_resumen_mayor_mes(apertura, partidas_mes, cuentas_historicas=cuentas_historicas)
estado_cuadre = evaluar_cuadre_resumen(df_resumen)
df_fuente = dataframe_partidas_fuente(apertura, partidas_mes)
categorias_pendientes_default = [
    "Pago pendiente",
    "Cobro a clientes",
    "Reintegro pendiente",
    "Anticipo pendiente",
    "IVA por cobrar",
    "Crédito de retención IVA",
    "Revisión / ajustar",
]
df_pendientes_base = obtener_saldos_pendientes_desde_resumen(
    df_resumen,
    incluir_saldadas=False,
    categorias=categorias_pendientes_default
)
registro_cierre_existente = None
if es_cierre_trimestral_actual:
    registro_cierre_existente = buscar_cierre_parcial(
        cierres_parciales_data,
        empresa_sel["id"],
        int(anio),
        int(trimestre_actual)
    )

tab1, tab2, tab3 = st.tabs(["Mayor mensual", "Saldos pendientes y revisión", "Partidas fuente"])

with tab1:
    render_tarjetas_cuadre(estado_cuadre)

    df_alertas_globales = obtener_dataframe_alertas_saldo_contrario(df_resumen)
    if not df_alertas_globales.empty:
        st.warning(
            f"Se detectaron {len(df_alertas_globales)} cuenta(s) con saldo final contrario a su tipo/naturaleza. "
            "Puedes seleccionarlas en la tabla para revisarlas mejor."
        )

    if df_resumen.empty:
        st.info("No hay datos para mostrar en el mayor del mes.")
    else:
        st.caption("Haz clic en una o varias filas de la tabla para resaltarlas y ubicar más fácil sus saldos finales.")

        columnas_ocultas_mayor = ["tipo", "naturaleza_catalogo", "naturaleza_real"]
        df_resumen_visual = df_resumen.drop(columns=columnas_ocultas_mayor, errors="ignore")

        tabla_mayor_key = f"tabla_mayor_{empresa_sel['id']}_{anio}_{mes}"
        filas_seleccionadas_idx = list(st.session_state.get(f"{tabla_mayor_key}_filas", []))

        try:
            evento_tabla = st.dataframe(
                aplicar_estilo_colores(df_resumen_visual),
                use_container_width=True,
                height=560,
                hide_index=True,
                on_select="rerun",
                selection_mode="multi-row",
                key=tabla_mayor_key
            )

            filas_seleccionadas = []
            if evento_tabla is not None:
                seleccion = getattr(evento_tabla, "selection", None)
                if seleccion is not None:
                    if isinstance(seleccion, dict):
                        filas_seleccionadas = list(seleccion.get("rows", []) or [])
                    else:
                        filas_seleccionadas = list(getattr(seleccion, "rows", []) or [])

            if filas_seleccionadas is not None:
                filas_seleccionadas_idx = [int(x) for x in filas_seleccionadas]
                st.session_state[f"{tabla_mayor_key}_filas"] = filas_seleccionadas_idx

        except TypeError:
            st.dataframe(
                aplicar_estilo_colores(df_resumen_visual),
                use_container_width=True,
                height=560,
                hide_index=True
            )
            st.info("Tu versión de Streamlit no soporta selección directa en la tabla; por eso aquí se muestra la vista normal.")

        filas_validas = [
            int(idx) for idx in (filas_seleccionadas_idx or [])
            if 0 <= int(idx) < len(df_resumen)
        ]

        if filas_validas:
            filas_validas = list(dict.fromkeys(filas_validas))
            st.success(f"Filas seleccionadas: {len(filas_validas)}")

            resumen_seleccion = []
            alertas_seleccion = []
            for idx in filas_validas:
                fila_sel = df_resumen.iloc[idx]
                cuenta_sel = str(fila_sel.get("cuenta", "")).strip()
                if not cuenta_sel or cuenta_sel.upper() == "TOTAL":
                    continue

                saldo_debe = round(float(fila_sel.get("saldo_final_debe", 0) or 0), 2)
                saldo_haber = round(float(fila_sel.get("saldo_final_haber", 0) or 0), 2)
                lado_saldo, monto_saldo = obtener_lado_monto_saldo_visual(saldo_debe, saldo_haber)
                alerta_saldo = obtener_alerta_saldo_contrario(fila_sel)

                resumen_seleccion.append({
                    "Cuenta": cuenta_sel,
                    "Saldo final": f"{lado_saldo} Q {monto_saldo:,.2f}",
                    "Estado": "⚠️ Revisar" if alerta_saldo else "OK"
                })

                if alerta_saldo:
                    alertas_seleccion.append({
                        "Cuenta": cuenta_sel,
                        "Alerta": alerta_saldo["mensaje"]
                    })

            if resumen_seleccion:
                st.dataframe(pd.DataFrame(resumen_seleccion), use_container_width=True, hide_index=True)

            if alertas_seleccion:
                st.error(
                    "Se detectaron saldos finales contrarios al tipo/naturaleza en las cuentas seleccionadas. "
                    "Conviene revisarlas antes de dar por bueno el mayor."
                )
                st.dataframe(pd.DataFrame(alertas_seleccion), use_container_width=True, hide_index=True)

        st.markdown("#### Exportar mayor a Excel")
        st.caption(
            "Puedes descargar solo el mes seleccionado o un archivo con los meses desde enero hasta el mes seleccionado, colocados horizontalmente."
        )

        folio_inicial_excel = st.number_input(
            "Folio inicial para el Excel",
            min_value=1,
            value=1,
            step=1,
            key=f"folio_inicial_mayor_excel_{empresa_sel['id']}_{anio}_{mes}",
        )
        nombre_base_excel = nombre_archivo_seguro(empresa_sel["nombre"])
        bloques_mes_excel = [{"mes": int(mes), "df": df_resumen, "partidas_mes": partidas_mes}]
        bloques_hasta_mes_excel = construir_bloques_mayor_hasta_mes(
            ventas_data=ventas_data,
            compras_data=compras_data,
            partidas_generadas_data=partidas_generadas_data,
            empresa_id=empresa_sel["id"],
            anio=int(anio),
            mes_fin=int(mes),
        )

        col_excel_mes, col_excel_todo = st.columns(2)
        with col_excel_mes:
            st.download_button(
                "Descargar Excel del mes",
                data=exportar_mayor_excel(
                    empresa=empresa_sel["nombre"],
                    empresa_id=empresa_sel["id"],
                    anio=int(anio),
                    bloques_mensuales=bloques_mes_excel,
                    folio_inicial=int(folio_inicial_excel),
                ),
                file_name=f"mayor_{nombre_base_excel}_{int(anio)}_{int(mes):02d}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
                key=f"descargar_mayor_mes_{empresa_sel['id']}_{anio}_{mes}",
            )
        with col_excel_todo:
            st.download_button(
                f"Descargar Excel enero a {MESES[int(mes)]}",
                data=exportar_mayor_excel(
                    empresa=empresa_sel["nombre"],
                    empresa_id=empresa_sel["id"],
                    anio=int(anio),
                    bloques_mensuales=bloques_hasta_mes_excel,
                    folio_inicial=int(folio_inicial_excel),
                ),
                file_name=f"mayor_{nombre_base_excel}_{int(anio)}_enero_a_{int(mes):02d}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
                key=f"descargar_mayor_hasta_mes_{empresa_sel['id']}_{anio}_{mes}",
            )

    st.markdown('</div>', unsafe_allow_html=True)

    st.markdown('<div class="panel">', unsafe_allow_html=True)
    st.markdown('<div class="panel-title">Cierre trimestral acumulado</div>', unsafe_allow_html=True)

    if normalizar_txt(regimen_empresa) != "trimestral":
        st.info("Esta opción solo se habilita para empresas con régimen trimestral.")
    elif not es_cierre_trimestral_actual:
        st.info("El cierre trimestral solo se puede guardar en marzo, junio, septiembre y diciembre.")
    elif df_resumen.empty:
        st.warning("No hay datos acumulados del mayor para guardar el cierre trimestral.")
    else:
        meses_txt = ", ".join(MESES[int(x)] for x in rango_meses_hasta_trimestre(mes))
        st.caption(
            f"Se guardará el cierre parcial del trimestre {int(trimestre_actual)} con los saldos acumulados de {meses_txt} hasta {MESES[int(mes)]} {int(anio)}."
        )

        if registro_cierre_existente:
            st.info(
                f"Ya existe un cierre guardado para el trimestre {int(trimestre_actual)}. "
                f"Si vuelves a cerrar, se reemplazará el JSON anterior. Último guardado: {registro_cierre_existente.get('fecha_guardado', '')}"
            )

        if st.button("Cerrar trimestre", type="primary", use_container_width=True, key=f"cerrar_trimestre_{empresa_sel['id']}_{anio}_{mes}"):
            registro_cierre = construir_registro_cierre_parcial(
                empresa=empresa_sel["nombre"],
                empresa_id=empresa_sel["id"],
                regimen=regimen_empresa,
                anio=int(anio),
                mes=int(mes),
                df_resumen=df_resumen,
                estado_cuadre=estado_cuadre,
                df_pendientes=df_pendientes_base
            )
            guardar_o_actualizar_cierre_parcial(cierres_parciales_data, registro_cierre)
            st.success(
                f"Cierre trimestral guardado en {CIERRES_PARCIALES_FILE} para {empresa_sel['nombre']} - trimestre {int(trimestre_actual)} de {int(anio)}."
            )
            st.rerun()

    st.markdown('</div>', unsafe_allow_html=True)

with tab2:
    st.markdown('<div class="panel">', unsafe_allow_html=True)
    st.markdown('<div class="panel-title">Saldos pendientes, cobros, reintegros, anticipos, IVA y cuentas por revisar</div>', unsafe_allow_html=True)
    st.caption("Aquí se toman los saldos finales del mes para detectar tanto pendientes de pago como cuentas especiales —incluyendo IVA por cobrar y crédito de retención IVA— que conviene revisar o presentar al cliente.")

    c_filtro1, c_filtro2 = st.columns([2, 1])
    with c_filtro1:
        categorias_sel = st.multiselect(
            "Categorías a mostrar",
            options=categorias_pendientes_default,
            default=categorias_pendientes_default,
            key=f"categorias_pendientes_{empresa_sel['id']}_{anio}_{mes}"
        )
    with c_filtro2:
        mostrar_cuentas_saldadas = st.checkbox(
            "Incluir saldadas en 0",
            value=False,
            key="mostrar_cuentas_saldadas_mayor"
        )

    df_pendientes = obtener_saldos_pendientes_desde_resumen(
        df_resumen,
        incluir_saldadas=mostrar_cuentas_saldadas,
        categorias=categorias_sel
    )

    if df_pendientes.empty:
        if mostrar_cuentas_saldadas:
            st.info("No se encontraron cuentas pendientes, saldadas o por revisar con esos filtros.")
        else:
            st.info("No se encontraron cuentas para seguimiento con los filtros seleccionados.")
    else:
        total_monto = round(pd.to_numeric(df_pendientes["monto"], errors="coerce").fillna(0).sum(), 2)
        c_kpi1, c_kpi2, c_kpi3 = st.columns(3)
        c_kpi1.metric("Cuentas detectadas", len(df_pendientes))
        c_kpi2.metric("Monto total", formatear_moneda(total_monto))
        c_kpi3.metric("Categorías", int(df_pendientes["categoria"].nunique()))

        st.markdown("#### Selección de cuentas para reporte")
        st.caption("Marca las cuentas y, si quieres, edita la columna '¿Qué hacer?' antes de generar el reporte.")

        df_editor = df_pendientes.copy()
        df_editor["que_hacer"] = df_editor.apply(obtener_accion_breve_cliente, axis=1)
        columnas_editor = [
            "seleccionar",
            "cuenta",
            "monto",
            "que_hacer",
            "categoria",
        ]
        df_editor = df_editor[columnas_editor]

        edited_df = st.data_editor(
            df_editor,
            use_container_width=True,
            hide_index=True,
            height=430,
            disabled=["cuenta", "monto", "categoria"],
            column_config={
                "seleccionar": st.column_config.CheckboxColumn("Seleccionar"),
                "cuenta": st.column_config.TextColumn("Cuenta", width="medium"),
                "monto": st.column_config.NumberColumn("Monto", format="Q %.2f"),
                "que_hacer": st.column_config.TextColumn("¿Qué hacer?", width="large", help="Puedes editar este texto antes de descargar el reporte."),
                "categoria": st.column_config.TextColumn("Categoría"),
            },
            key=f"editor_pendientes_{empresa_sel['id']}_{anio}_{mes}"
        )

        df_seleccionadas = edited_df[edited_df["seleccionar"] == True].copy()
        total_seleccionado = round(pd.to_numeric(df_seleccionadas.get("monto", pd.Series(dtype=float)), errors="coerce").fillna(0).sum(), 2)

        st.markdown("#### Reporte presentable para cliente")
        c_rep1, c_rep2 = st.columns([2, 1])
        with c_rep1:
            titulo_reporte = st.text_input(
                "Título del reporte",
                value="Resumen Contable",
                key=f"titulo_reporte_pendientes_{empresa_sel['id']}_{anio}_{mes}"
            )
            observaciones_reporte = st.text_area(
                "Mensaje corto (opcional)",
                value="",
                height=80,
                key=f"obs_reporte_pendientes_{empresa_sel['id']}_{anio}_{mes}"
            )
        with c_rep2:
            logo_file = st.file_uploader(
                "Logotipo (opcional)",
                type=["png", "jpg", "jpeg", "svg"],
                key=f"logo_reporte_pendientes_{empresa_sel['id']}_{anio}_{mes}"
            )
            st.caption(
                f"Seleccionadas: {len(df_seleccionadas)} cuenta(s) | Total: {formatear_moneda(total_seleccionado)}"
            )

        if df_seleccionadas.empty:
            st.warning("Marca al menos una cuenta en la tabla para generar el reporte.")
        else:
            logo_data_uri = obtener_logo_data_uri(logo_file)
            html_reporte = construir_html_reporte_pendientes(
                empresa=empresa_sel["nombre"],
                empresa_id=empresa_sel["id"],
                anio=int(anio),
                mes=int(mes),
                df_seleccionadas=df_seleccionadas,
                titulo_reporte=titulo_reporte,
                observaciones=observaciones_reporte,
                logo_data_uri=logo_data_uri,
            )

            nombre_archivo = (
                f"reporte_pendientes_{normalizar_nit(empresa_sel['id'])}_{int(anio)}_{int(mes):02d}.html"
            )

            c_btn1, c_btn2 = st.columns([1, 1])
            with c_btn1:
                st.download_button(
                    "Descargar reporte HTML",
                    data=html_reporte,
                    file_name=nombre_archivo,
                    mime="text/html",
                    use_container_width=True,
                    key=f"descargar_reporte_pendientes_{empresa_sel['id']}_{anio}_{mes}"
                )
            with c_btn2:
                st.download_button(
                    "Descargar selección CSV",
                    data=df_seleccionadas.drop(columns=["seleccionar"], errors="ignore").to_csv(index=False).encode("utf-8-sig"),
                    file_name=(
                        f"cuentas_seleccionadas_{normalizar_nit(empresa_sel['id'])}_{int(anio)}_{int(mes):02d}.csv"
                    ),
                    mime="text/csv",
                    use_container_width=True,
                    key=f"descargar_csv_pendientes_{empresa_sel['id']}_{anio}_{mes}"
                )

            with st.expander("Vista previa del reporte", expanded=False):
                components.html(html_reporte, height=760, scrolling=True)

    st.markdown('</div>', unsafe_allow_html=True)

with tab3:
    st.markdown('<div class="panel">', unsafe_allow_html=True)
    st.markdown('<div class="panel-title">Partidas utilizadas para construir el mayor del mes</div>', unsafe_allow_html=True)
    st.dataframe(df_fuente, use_container_width=True, height=560)
    st.markdown('</div>', unsafe_allow_html=True)
