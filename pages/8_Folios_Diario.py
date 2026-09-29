from __future__ import annotations

import base64
import copy
import html
import json
import re
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import streamlit as st
import streamlit.components.v1 as components
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break


st.set_page_config(page_title="Folios del Libro Diario", layout="wide")


# =========================
# Utilidades
# =========================
def leer_json_seguro_desde_bytes(raw_bytes: bytes) -> Any:
    try:
        return json.loads(raw_bytes.decode("utf-8"))
    except UnicodeDecodeError:
        return json.loads(raw_bytes.decode("latin-1"))



def leer_json_seguro_desde_ruta(ruta: str | Path) -> Any:
    ruta = Path(ruta)
    if not ruta.exists():
        return None
    with open(ruta, "rb") as f:
        return leer_json_seguro_desde_bytes(f.read())



def a_float(valor: Any) -> float:
    try:
        if valor in (None, ""):
            return 0.0
        return float(valor)
    except Exception:
        return 0.0



def fmt_monto(valor: Any) -> str:
    monto = a_float(valor)
    return f"Q {monto:,.2f}" if monto else ""



def escape(valor: Any) -> str:
    return html.escape("" if valor is None else str(valor))



def normalizar_texto(valor: Any) -> str:
    texto = str(valor or "").strip().lower()
    return re.sub(r"\s+", " ", texto)



def extraer_numero_codigo(codigo: Any) -> int:
    texto = str(codigo or "")
    m = re.search(r"(\d+)", texto)
    return int(m.group(1)) if m else 0



def parse_fecha(fecha: Any) -> datetime:
    texto = str(fecha or "").strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(texto, fmt)
        except Exception:
            pass
    return datetime(1900, 1, 1)



def encode_logo_base64(ruta_logo: str | Path) -> str:
    ruta = Path(ruta_logo)
    if not ruta.exists():
        return ""
    mime = "image/png"
    if ruta.suffix.lower() in {".jpg", ".jpeg"}:
        mime = "image/jpeg"
    elif ruta.suffix.lower() == ".webp":
        mime = "image/webp"
    with open(ruta, "rb") as f:
        contenido = base64.b64encode(f.read()).decode("utf-8")
    return f"data:{mime};base64,{contenido}"


# =========================
# Carga de datos
# =========================
def cargar_partidas_generadas(uploaded_file=None) -> List[Dict[str, Any]]:
    if uploaded_file is not None:
        try:
            data = leer_json_seguro_desde_bytes(uploaded_file.getvalue())
            return data if isinstance(data, list) else []
        except Exception as e:
            st.error(f"No se pudo leer el JSON subido: {e}")
            return []

    for candidato in [
        Path("partidas_generadas.json"),
        Path.cwd() / "partidas_generadas.json",
        Path(__file__).resolve().parent / "partidas_generadas.json",
    ]:
        data = leer_json_seguro_desde_ruta(candidato)
        if isinstance(data, list):
            return data
    return []



def cargar_empresas_json(uploaded_file=None) -> Any:
    if uploaded_file is not None:
        try:
            return leer_json_seguro_desde_bytes(uploaded_file.getvalue())
        except Exception as e:
            st.warning(f"No se pudo leer el JSON de empresas/membretes: {e}")
            return None

    for candidato in [
        Path("empresas.json"),
        Path.cwd() / "empresas.json",
        Path(__file__).resolve().parent / "empresas.json",
        Path("membretes.json"),
        Path.cwd() / "membretes.json",
        Path(__file__).resolve().parent / "membretes.json",
    ]:
        data = leer_json_seguro_desde_ruta(candidato)
        if data is not None:
            return data
    return None


# =========================
# Empresas / membretes
# =========================
def indexar_empresas_meta(data: Any) -> List[Dict[str, Any]]:
    if isinstance(data, list):
        return [x for x in data if isinstance(x, dict)]
    if isinstance(data, dict):
        resultado = []
        for clave, valor in data.items():
            if isinstance(valor, dict):
                fila = copy.deepcopy(valor)
                fila.setdefault("empresa_nit", valor.get("nit") or clave)
                resultado.append(fila)
        return resultado
    return []



def buscar_meta_empresa(empresa: str, nit: str, empresas_meta: List[Dict[str, Any]]) -> Optional[Dict[str, Any]]:
    emp_n = normalizar_texto(empresa)
    nit_n = normalizar_texto(nit)

    for item in empresas_meta:
        nit_item = normalizar_texto(item.get("empresa_nit") or item.get("nit"))
        if nit_n and nit_item and nit_n == nit_item:
            return item

    for item in empresas_meta:
        nombre_item = normalizar_texto(
            item.get("empresa")
            or item.get("nombre")
            or item.get("razon_social")
            or item.get("comercial")
        )
        if emp_n and nombre_item and emp_n == nombre_item:
            return item

    return None



def resolver_membrete(empresa: str, nit: str, empresas_meta: List[Dict[str, Any]]) -> Dict[str, str]:
    item = buscar_meta_empresa(empresa, nit, empresas_meta) or {}

    nombre = (
        item.get("razon_social")
        or item.get("empresa")
        or item.get("nombre")
        or item.get("comercial")
        or empresa
    )
    nit_final = item.get("nit") or item.get("empresa_nit") or nit
    direccion = item.get("direccion") or item.get("direccion_empresa") or ""
    telefono = item.get("telefono") or item.get("tel") or ""
    actividad = item.get("actividad") or item.get("giro") or ""
    representante = item.get("representante") or item.get("propietario") or item.get("contador") or ""
    logo = item.get("logo") or item.get("logo_path") or ""

    extras = [x for x in [telefono, actividad, representante] if str(x or "").strip()]
    linea3 = " | ".join(str(x).strip() for x in extras)

    logo_b64 = ""
    if logo:
        try:
            logo_b64 = encode_logo_base64(logo)
        except Exception:
            logo_b64 = ""

    return {
        "nombre": str(nombre or empresa),
        "nit": str(nit_final or nit),
        "linea2": str(direccion or ""),
        "linea3": str(linea3 or ""),
        "logo_b64": logo_b64,
    }


# =========================
# Preparación de partidas
# =========================
def obtener_empresas_disponibles(data: List[Dict[str, Any]]) -> List[Tuple[str, str]]:
    empresas = {}
    for bloque in data:
        empresa = str(bloque.get("empresa") or "").strip()
        nit = str(bloque.get("empresa_nit") or "").strip()
        if empresa:
            empresas[(empresa, nit)] = True
    return sorted(empresas.keys(), key=lambda x: (normalizar_texto(x[0]), normalizar_texto(x[1])))



def obtener_anios_disponibles_por_empresa(data: List[Dict[str, Any]], empresa: str, nit: str) -> List[int]:
    anios = set()
    for bloque in data:
        if str(bloque.get("empresa") or "").strip() != empresa:
            continue
        if str(bloque.get("empresa_nit") or "").strip() != nit:
            continue
        anio_bloque = bloque.get("anio")
        if isinstance(anio_bloque, int):
            anios.add(anio_bloque)
        for partida in bloque.get("partidas", []):
            anio_partida = partida.get("anio", anio_bloque)
            if isinstance(anio_partida, int):
                anios.add(anio_partida)
    return sorted(anios)



def filtrar_y_ordenar_partidas(
    data: List[Dict[str, Any]],
    empresa: str,
    nit: str,
    anios_sel: List[int],
) -> List[Dict[str, Any]]:
    partidas: List[Dict[str, Any]] = []

    for bloque in data:
        if str(bloque.get("empresa") or "").strip() != empresa:
            continue
        if str(bloque.get("empresa_nit") or "").strip() != nit:
            continue

        for partida in bloque.get("partidas", []):
            anio_partida = partida.get("anio", bloque.get("anio"))
            if anios_sel and anio_partida not in anios_sel:
                continue

            cuentas = []
            for cuenta in partida.get("cuentas", []):
                cuentas.append(
                    {
                        "cuenta": str(cuenta.get("cuenta") or "").strip(),
                        "debe": a_float(cuenta.get("debe")),
                        "haber": a_float(cuenta.get("haber")),
                    }
                )

            partidas.append(
                {
                    "empresa": str(partida.get("empresa") or empresa),
                    "empresa_nit": str(partida.get("empresa_nit") or nit),
                    "anio": anio_partida,
                    "mes": int(partida.get("mes", bloque.get("mes") or 0) or 0),
                    "fecha": str(partida.get("fecha") or ""),
                    "codigo": str(partida.get("codigo") or ""),
                    "glosa": str(partida.get("glosa") or ""),
                    "cuentas": cuentas,
                    "total_debe": a_float(partida.get("total_debe")),
                    "total_haber": a_float(partida.get("total_haber")),
                }
            )

    partidas.sort(
        key=lambda p: (
            parse_fecha(p.get("fecha")),
            int(p.get("anio") or 0),
            int(p.get("mes") or 0),
            extraer_numero_codigo(p.get("codigo")),
        )
    )
    return partidas



def construir_mapa_colores_mes(partidas: List[Dict[str, Any]]) -> Dict[Tuple[int, int], str]:
    pares = []
    vistos = set()
    for partida in partidas:
        clave = (int(partida.get("anio") or 0), int(partida.get("mes") or 0))
        if clave not in vistos:
            pares.append(clave)
            vistos.add(clave)

    colores = ["DDEBF7", "FCE4D6"]
    return {clave: colores[i % len(colores)] for i, clave in enumerate(pares)}


# =========================
# Motor de folios en 4 columnas
# =========================
def crear_folios_desde_partidas(
    partidas: List[Dict[str, Any]],
    lineas_por_folio: int,
    folio_inicial: int = 1,
    colores_mes: Optional[Dict[Tuple[int, int], str]] = None,
) -> List[Dict[str, Any]]:
    folios: List[Dict[str, Any]] = []
    folio_actual: Optional[Dict[str, Any]] = None
    numero_folio = folio_inicial - 1
    colores_mes = colores_mes or {}

    def nuevo_folio() -> None:
        nonlocal numero_folio, folio_actual
        numero_folio += 1
        folio_actual = {
            "folio": numero_folio,
            "rows": [],
            "lineas_usadas": 0,
        }

    def cerrar_folio() -> None:
        nonlocal folio_actual
        if folio_actual is None:
            return
        while folio_actual["lineas_usadas"] < lineas_por_folio:
            folio_actual["rows"].append({"tipo": "blank"})
            folio_actual["lineas_usadas"] += 1
        folios.append(folio_actual)
        folio_actual = None

    def asegurar_folio() -> None:
        if folio_actual is None:
            nuevo_folio()

    def espacio_restante() -> int:
        if folio_actual is None:
            return 0
        return lineas_por_folio - folio_actual["lineas_usadas"]

    def agregar_row(row: Dict[str, Any]) -> None:
        asegurar_folio()
        folio_actual["rows"].append(row)
        folio_actual["lineas_usadas"] += 1

    if not partidas:
        return []

    nuevo_folio()

    for partida in partidas:
        cuentas = partida.get("cuentas") or [{"cuenta": "(Sin cuentas)", "debe": 0.0, "haber": 0.0}]
        idx = 0
        acumulado_debe = 0.0
        acumulado_haber = 0.0
        primera_seccion = True
        color_mes = colores_mes.get((int(partida.get("anio") or 0), int(partida.get("mes") or 0)), "DDEBF7")

        while idx < len(cuentas):
            lineas_fijas = 2  # encabezado+cierre, o vienen+cierre
            if espacio_restante() < lineas_fijas + 1:
                cerrar_folio()
                nuevo_folio()

            if primera_seccion:
                agregar_row(
                    {
                        "tipo": "encabezado_partida",
                        "partida": partida.get("codigo", ""),
                        "fecha": partida.get("fecha", ""),
                    }
                )
            else:
                agregar_row(
                    {
                        "tipo": "vienen",
                        "concepto": "VIENEN",
                        "debe": acumulado_debe,
                        "haber": acumulado_haber,
                    }
                )

            disponibles_para_cuentas = max(1, espacio_restante() - 1)
            cuentas_restantes = len(cuentas) - idx
            cuentas_a_imprimir = min(cuentas_restantes, disponibles_para_cuentas)

            for _ in range(cuentas_a_imprimir):
                cuenta = cuentas[idx]
                debe = a_float(cuenta.get("debe"))
                haber = a_float(cuenta.get("haber"))
                agregar_row(
                    {
                        "tipo": "cuenta",
                        "concepto": cuenta.get("cuenta", ""),
                        "debe": debe,
                        "haber": haber,
                        "es_haber": str(cuenta.get("cuenta") or "").strip().startswith("A:"),
                    }
                )
                acumulado_debe += debe
                acumulado_haber += haber
                idx += 1

            if idx < len(cuentas):
                agregar_row(
                    {
                        "tipo": "van",
                        "concepto": "VAN",
                        "debe": acumulado_debe,
                        "haber": acumulado_haber,
                    }
                )
                cerrar_folio()
                nuevo_folio()
                primera_seccion = False
            else:
                agregar_row(
                    {
                        "tipo": "cierre_partida",
                        "concepto": partida.get("glosa", ""),
                        "debe": partida.get("total_debe", acumulado_debe),
                        "haber": partida.get("total_haber", acumulado_haber),
                        "fill_hex": color_mes,
                    }
                )
                break

    cerrar_folio()
    return folios


# =========================
# HTML imprimible
# =========================
def render_rows_html(rows: List[Dict[str, Any]]) -> str:
    html_rows: List[str] = []
    for row in rows:
        tipo = row.get("tipo")

        if tipo == "encabezado_partida":
            html_rows.append(
                f"""
                <tr class="row-encabezado-partida">
                    <td class="c-partida strong">{escape(row.get('partida'))}</td>
                    <td class="c-concepto fecha-top strong">{escape(row.get('fecha'))}</td>
                    <td class="c-num"></td>
                    <td class="c-num"></td>
                </tr>
                """
            )
        elif tipo == "cuenta":
            cuenta_class = "cuenta-haber" if row.get("es_haber") else "cuenta-debe"
            html_rows.append(
                f"""
                <tr class="row-cuenta">
                    <td class="c-partida"></td>
                    <td class="c-concepto {cuenta_class}">{escape(row.get('concepto'))}</td>
                    <td class="c-num">{fmt_monto(row.get('debe'))}</td>
                    <td class="c-num">{fmt_monto(row.get('haber'))}</td>
                </tr>
                """
            )
        elif tipo in {"van", "vienen"}:
            html_rows.append(
                f"""
                <tr class="row-arrastre {escape(tipo)}">
                    <td class="c-partida"></td>
                    <td class="c-concepto strong">{escape(row.get('concepto'))}</td>
                    <td class="c-num strong">{fmt_monto(row.get('debe'))}</td>
                    <td class="c-num strong">{fmt_monto(row.get('haber'))}</td>
                </tr>
                """
            )
        elif tipo == "cierre_partida":
            fill = row.get("fill_hex") or "FFFFFF"
            html_rows.append(
                f"""
                <tr class="row-cierre-partida">
                    <td class="c-partida"></td>
                    <td class="c-concepto italic">{escape(row.get('concepto'))}</td>
                    <td class="c-num strong total-cierre" style="background: #{escape(fill)};">{fmt_monto(row.get('debe'))}</td>
                    <td class="c-num strong total-cierre" style="background: #{escape(fill)};">{fmt_monto(row.get('haber'))}</td>
                </tr>
                """
            )
        else:
            html_rows.append(
                """
                <tr class="row-blank">
                    <td class="c-partida">&nbsp;</td>
                    <td class="c-concepto"></td>
                    <td class="c-num"></td>
                    <td class="c-num"></td>
                </tr>
                """
            )

    return "\n".join(html_rows)



def construir_html_empresa(
    empresa: str,
    nit: str,
    membrete: Dict[str, str],
    folios: List[Dict[str, Any]],
    titulo_libro: str,
    lineas_por_folio: int,
) -> str:
    logo_html = ""
    if membrete.get("logo_b64"):
        logo_html = f'<img class="logo" src="{membrete["logo_b64"]}" alt="Logo">'

    folios_html = []
    total_folios = len(folios)

    for folio in folios:
        rows_html = render_rows_html(folio["rows"])
        folios_html.append(
            f"""
            <section class="folio">
                <div class="folio-header">
                    <div class="membrete-left">{logo_html}</div>
                    <div class="membrete-center">
                        <div class="empresa">{escape(membrete.get('nombre') or empresa)}</div>
                        <div class="meta"><strong>NIT:</strong> {escape(membrete.get('nit') or nit)}</div>
                        <div class="meta">{escape(membrete.get('linea2'))}</div>
                        <div class="meta">{escape(membrete.get('linea3'))}</div>
                    </div>
                    <div class="membrete-right">
                        <div><strong>{escape(titulo_libro)}</strong></div>
                        <div><strong>Folio:</strong> {folio['folio']}</div>
                        <div><strong>Empresa:</strong> {escape(empresa)}</div>
                    </div>
                </div>

                <table class="libro-table">
                    <thead>
                        <tr>
                            <th class="c-partida">No. Partida</th>
                            <th class="c-concepto">Fecha / Cuenta / Descripción</th>
                            <th class="c-num">Debe</th>
                            <th class="c-num">Haber</th>
                        </tr>
                    </thead>
                    <tbody>
                        {rows_html}
                    </tbody>
                </table>

                <div class="folio-footer">
                    <span>{escape(empresa)} | NIT {escape(nit)}</span>
                    <span>Folio {folio['folio']} de {total_folios}</span>
                </div>
            </section>
            """
        )

    return "\n".join(folios_html)



def construir_documento_html(
    empresa: str,
    nit: str,
    membrete: Dict[str, str],
    folios: List[Dict[str, Any]],
    titulo_libro: str,
    lineas_por_folio: int,
) -> str:
    contenido = construir_html_empresa(
        empresa=empresa,
        nit=nit,
        membrete=membrete,
        folios=folios,
        titulo_libro=titulo_libro,
        lineas_por_folio=lineas_por_folio,
    )

    html_doc = f"""
    <!doctype html>
    <html lang="es">
    <head>
        <meta charset="utf-8">
        <title>{escape(titulo_libro)}</title>
        <style>
            @page {{
                size: letter portrait;
                margin: 8mm;
            }}

            * {{ box-sizing: border-box; }}

            body {{
                margin: 0;
                padding: 0;
                font-family: Arial, Helvetica, sans-serif;
                color: #111827;
                background: #edf2f7;
            }}

            .toolbar {{
                position: sticky;
                top: 0;
                z-index: 999;
                background: #0f172a;
                color: white;
                padding: 10px 16px;
                display: flex;
                justify-content: space-between;
                align-items: center;
                gap: 12px;
            }}

            .toolbar button {{
                background: #2563eb;
                border: none;
                color: white;
                border-radius: 8px;
                padding: 8px 14px;
                font-weight: 700;
                cursor: pointer;
            }}

            .wrap {{ padding: 12px 0 28px 0; }}

            .folio {{
                width: 195mm;
                min-height: 271mm;
                margin: 0 auto 14px auto;
                background: white;
                border: 1px solid #111827;
                padding: 7mm 7mm 6mm 7mm;
                display: flex;
                flex-direction: column;
                page-break-after: always;
            }}

            .folio-header {{
                display: grid;
                grid-template-columns: 28mm 1fr 52mm;
                gap: 8px;
                align-items: center;
                border-bottom: 1.2px solid #111827;
                padding-bottom: 5px;
                margin-bottom: 6px;
                min-height: 28mm;
            }}

            .membrete-left {{
                display: flex;
                align-items: center;
                justify-content: center;
                min-height: 22mm;
            }}

            .logo {{
                max-width: 24mm;
                max-height: 24mm;
                object-fit: contain;
            }}

            .membrete-center {{ text-align: center; }}
            .empresa {{ font-size: 16px; font-weight: 800; text-transform: uppercase; }}
            .meta {{ font-size: 11px; line-height: 1.2; margin-top: 2px; }}
            .membrete-right {{ font-size: 11px; line-height: 1.4; text-align: right; }}

            .libro-table {{
                width: 100%;
                border-collapse: collapse;
                table-layout: fixed;
            }}

            .libro-table thead th {{
                border: 1px solid #111827;
                padding: 4px 6px;
                font-size: 11px;
                font-weight: 800;
                text-transform: uppercase;
                background: #f8fafc;
                height: 8mm;
            }}

            .libro-table tbody td {{
                border: 1px solid #111827;
                padding: 2px 6px;
                font-size: 11px;
                height: 7.2mm;
                white-space: nowrap;
                overflow: hidden;
                text-overflow: ellipsis;
            }}

            .c-partida {{ width: 24mm; text-align: center; }}
            .c-concepto {{ width: auto; text-align: left; }}
            .c-num {{ width: 28mm; text-align: right; }}

            .cuenta-debe {{ padding-left: 16px !important; }}
            .cuenta-haber {{ padding-left: 48px !important; }}

            .row-arrastre td {{ font-weight: 800; background: #f8fafc; }}
            .row-arrastre .c-concepto {{
                text-align: center !important;
                letter-spacing: 0.4px;
            }}
            .row-arrastre .c-num {{
                background: #f8fafc;
            }}
            .row-cierre-partida .total-cierre {{
                border-top: 1px solid #111827 !important;
                border-bottom: 3px double #111827 !important;
            }}

            .strong {{ font-weight: 800; }}
            .italic {{ font-style: italic; }}

            .folio-footer {{
                margin-top: auto;
                padding-top: 6px;
                display: flex;
                justify-content: space-between;
                border-top: 1px solid #111827;
                font-size: 10px;
            }}

            @media print {{
                body {{ background: white; }}
                .toolbar {{ display: none !important; }}
                .wrap {{ padding: 0; }}
                .folio {{ margin: 0 auto; border: none; }}
            }}
        </style>
    </head>
    <body>
        <div class="toolbar">
            <div>{escape(titulo_libro)} | Líneas por folio: {lineas_por_folio}</div>
            <button onclick="window.print()">Imprimir</button>
        </div>
        <div class="wrap">{contenido}</div>
    </body>
    </html>
    """
    return html_doc


# =========================
# Excel
# =========================
def _aplicar_borde_rango(ws, fila: int, cols: range, border: Border) -> None:
    for col in cols:
        ws.cell(fila, col).border = border



def crear_excel_folios(
    empresa: str,
    nit: str,
    membrete: Dict[str, str],
    folios: List[Dict[str, Any]],
    titulo_libro: str,
    anio_membrete: str,
    etiqueta_anio: str,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\\/*?:\[\]]", " ", empresa)[:31] or "Folios"

    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = 0.2
    ws.page_margins.right = 0.2
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.35

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 18

    thin = Side(style="thin", color="000000")
    double_side = Side(style="double", color="000000")
    border_thin = Border(left=thin, right=thin, top=thin, bottom=thin)
    border_total = Border(left=thin, right=thin, top=thin, bottom=double_side)

    fill_header = PatternFill(fill_type="solid", fgColor="E7E6E6")
    font_title = Font(name="Arial", size=12, bold=True)
    font_sub = Font(name="Arial", size=10, bold=True)
    font_body = Font(name="Arial", size=10)
    font_italic = Font(name="Arial", size=10, italic=True)
    font_bold = Font(name="Arial", size=10, bold=True)

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    moneda_format = '"Q" #,##0.00'
    fila = 1

    for i, folio in enumerate(folios, start=1):
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=2)
        ws.merge_cells(start_row=fila, start_column=3, end_row=fila, end_column=4)
        c = ws.cell(fila, 1, membrete.get("linea3") or empresa)
        c.font = font_title
        c.alignment = align_left
        c = ws.cell(fila, 3, f"Folio: {folio['folio']}")
        c.font = font_sub
        c.alignment = align_right
        fila += 1

        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=2)
        ws.merge_cells(start_row=fila, start_column=3, end_row=fila, end_column=4)
        c = ws.cell(fila, 1, f"NIT: {membrete.get('nit') or nit}")
        c.font = font_sub
        c.alignment = align_left
        c = ws.cell(fila, 3, f"{etiqueta_anio}: {anio_membrete}")
        c.font = font_sub
        c.alignment = align_right
        fila += 1

        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=4)
        c = ws.cell(fila, 1, titulo_libro)
        c.font = font_sub
        c.alignment = align_center
        fila += 1

        headers = ["No. Partida", "Fecha / Cuenta / Descripción", "Debe", "Haber"]
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(fila, col, header)
            cell.font = font_bold
            cell.alignment = align_center
            cell.fill = fill_header
            cell.border = border_thin
        ws.row_dimensions[fila].height = 20
        fila += 1

        inicio_fila_folio = fila

        for row in folio["rows"]:
            tipo = row.get("tipo")

            for col in range(1, 5):
                ws.cell(fila, col).font = font_body
                ws.cell(fila, col).border = border_thin
                ws.cell(fila, col).alignment = align_left

            if tipo == "encabezado_partida":
                ws.cell(fila, 1, row.get("partida"))
                ws.cell(fila, 2, row.get("fecha"))
                ws.cell(fila, 1).font = font_bold
                ws.cell(fila, 2).font = font_bold
                ws.cell(fila, 1).alignment = align_center
                ws.cell(fila, 2).alignment = align_center

            elif tipo == "cuenta":
                ws.cell(fila, 2, row.get("concepto"))
                indent = 6 if row.get("es_haber") else 2
                ws.cell(fila, 2).alignment = Alignment(horizontal="left", vertical="center", indent=indent)
                if a_float(row.get("debe")):
                    ws.cell(fila, 3, a_float(row.get("debe")))
                    ws.cell(fila, 3).number_format = moneda_format
                if a_float(row.get("haber")):
                    ws.cell(fila, 4, a_float(row.get("haber")))
                    ws.cell(fila, 4).number_format = moneda_format
                ws.cell(fila, 3).alignment = align_right
                ws.cell(fila, 4).alignment = align_right

            elif tipo in {"van", "vienen"}:
                fill_arrastre = PatternFill(fill_type="solid", fgColor="F3F6FA")
                ws.cell(fila, 2, row.get("concepto"))
                ws.cell(fila, 3, a_float(row.get("debe")))
                ws.cell(fila, 4, a_float(row.get("haber")))
                ws.cell(fila, 2).font = font_bold
                ws.cell(fila, 3).font = font_bold
                ws.cell(fila, 4).font = font_bold
                ws.cell(fila, 2).alignment = align_center
                ws.cell(fila, 3).number_format = moneda_format
                ws.cell(fila, 4).number_format = moneda_format
                ws.cell(fila, 3).alignment = align_right
                ws.cell(fila, 4).alignment = align_right
                ws.cell(fila, 2).fill = fill_arrastre
                ws.cell(fila, 3).fill = fill_arrastre
                ws.cell(fila, 4).fill = fill_arrastre

            elif tipo == "cierre_partida":
                fill_mes = PatternFill(fill_type="solid", fgColor=row.get("fill_hex") or "FFFFFF")
                ws.cell(fila, 2, row.get("concepto"))
                ws.cell(fila, 2).font = font_italic
                ws.cell(fila, 2).alignment = align_left
                ws.cell(fila, 3, a_float(row.get("debe")))
                ws.cell(fila, 4, a_float(row.get("haber")))
                ws.cell(fila, 3).font = font_bold
                ws.cell(fila, 4).font = font_bold
                ws.cell(fila, 3).number_format = moneda_format
                ws.cell(fila, 4).number_format = moneda_format
                ws.cell(fila, 3).alignment = align_right
                ws.cell(fila, 4).alignment = align_right
                ws.cell(fila, 3).fill = fill_mes
                ws.cell(fila, 4).fill = fill_mes
                ws.cell(fila, 3).border = border_total
                ws.cell(fila, 4).border = border_total

            else:
                pass

            ws.row_dimensions[fila].height = 18
            fila += 1

        fin_fila_folio = fila - 1
        ws.print_area = f"A1:D{fin_fila_folio}"

        if i < len(folios):
            ws.row_breaks.append(Break(id=fin_fila_folio))
            fila += 1

        for col in range(1, 5):
            for r in range(inicio_fila_folio, fin_fila_folio + 1):
                if ws.cell(r, col).value is None:
                    ws.cell(r, col).border = border_thin

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


# =========================
# App
# =========================
def main() -> None:
    st.title("Folios del Libro Diario")
    st.caption("Selecciona una empresa y genera sus folios con VAN, VIENEN, vista imprimible y Excel.")

    with st.sidebar:
        st.subheader("Archivos")
        uploaded_partidas = st.file_uploader(
            "partidas_generadas.json",
            type=["json"],
            help="Si no subes archivo, se intentará usar el partidas_generadas.json local.",
        )
        uploaded_empresas = st.file_uploader(
            "empresas.json o membretes.json (opcional)",
            type=["json"],
            help="Opcional. Si existe localmente también se intentará cargar.",
        )

    data = cargar_partidas_generadas(uploaded_partidas)
    empresas_json = cargar_empresas_json(uploaded_empresas)
    empresas_meta = indexar_empresas_meta(empresas_json)

    if not data:
        st.error("No encontré datos válidos en partidas_generadas.json.")
        st.stop()

    empresas_disponibles = obtener_empresas_disponibles(data)
    if not empresas_disponibles:
        st.error("No hay empresas disponibles en el JSON.")
        st.stop()

    with st.sidebar:
        st.subheader("Configuración")
        titulo_libro = st.text_input("Título", value="Libro Diario")
        empresa_labels = [f"{empresa} | NIT {nit}" for empresa, nit in empresas_disponibles]
        empresa_sel_label = st.selectbox("Empresa", options=empresa_labels, index=0)
        empresa_sel, nit_sel = empresa_sel_label.split(" | NIT ", 1)
        empresa_sel = empresa_sel.strip()
        nit_sel = nit_sel.strip()

        anios_disponibles = obtener_anios_disponibles_por_empresa(data, empresa_sel, nit_sel)
        anios_sel = st.multiselect("Años", options=anios_disponibles, default=anios_disponibles)

        lineas_por_folio = st.number_input(
            "Líneas por folio",
            min_value=35,
            max_value=44,
            value=44,
            step=1,
            help="Ajusta esto si quieres más o menos líneas por hoja carta.",
        )
        folio_inicial = st.number_input("Folio inicial", min_value=1, value=1, step=1)

    if not anios_sel:
        st.warning("Selecciona al menos un año.")
        st.stop()

    anio_membrete = ", ".join(str(x) for x in sorted(anios_sel))
    etiqueta_anio = "Año" if len(anios_sel) == 1 else "Años"

    partidas = filtrar_y_ordenar_partidas(data, empresa_sel, nit_sel, anios_sel)
    if not partidas:
        st.warning("No hay partidas para la empresa y años seleccionados.")
        st.stop()

    membrete = resolver_membrete(empresa_sel, nit_sel, empresas_meta)
    colores_mes = construir_mapa_colores_mes(partidas)
    folios = crear_folios_desde_partidas(
        partidas=partidas,
        lineas_por_folio=int(lineas_por_folio),
        folio_inicial=int(folio_inicial),
        colores_mes=colores_mes,
    )

    st.subheader("Resumen")
    col1, col2, col3, col4 = st.columns(4)
    col1.metric("Empresa", empresa_sel)
    col2.metric("Partidas", len(partidas))
    col3.metric("Folios", len(folios))
    col4.metric("Años", ", ".join(str(x) for x in anios_sel))

    with st.expander("Cómo funciona", expanded=False):
        st.markdown(
            """
            - La **primera columna** es el número de partida.
            - En la **segunda columna** va primero la fecha, luego las cuentas, y al final la glosa.
            - Las cuentas que empiezan con **A:** se corren más a la derecha para distinguir el haber.
            - Si una partida no cabe en el folio, se cierra con **VAN** y continúa con **VIENEN**.
            - En Excel, los **totales de debe y haber** se alternan por mes con color y quedan cerrados con **doble línea**.
            """
        )

    html_doc = construir_documento_html(
        empresa=empresa_sel,
        nit=nit_sel,
        membrete=membrete,
        folios=folios,
        titulo_libro=titulo_libro,
        lineas_por_folio=int(lineas_por_folio),
    )

    excel_bytes = crear_excel_folios(
        empresa=empresa_sel,
        nit=nit_sel,
        membrete=membrete,
        folios=folios,
        titulo_libro=titulo_libro,
        anio_membrete=anio_membrete,
        etiqueta_anio=etiqueta_anio,
    )

    c1, c2 = st.columns(2)
    with c1:
        st.download_button(
            "Descargar vista imprimible en HTML",
            data=html_doc.encode("utf-8"),
            file_name=f"folios_{normalizar_texto(empresa_sel).replace(' ', '_')}.html",
            mime="text/html",
            use_container_width=True,
        )
    with c2:
        st.download_button(
            "Descargar en Excel",
            data=excel_bytes,
            file_name=f"folios_{normalizar_texto(empresa_sel).replace(' ', '_')}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
        )

    alto_estimado = max(900, len(folios) * 1100)
    components.html(html_doc, height=alto_estimado, scrolling=True)


if __name__ == "__main__":
    main()
