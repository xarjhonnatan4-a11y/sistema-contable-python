import streamlit as st
import pandas as pd
import json
import os
import io
import re
import unicodedata
from datetime import datetime
from copy import copy
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

CIERRES_PARCIALES_FILE = "cierres_parciales.json"
EMPRESAS_FILE = "empresas.json"
ISR_TRIMESTRAL_FILE = "isr_trimestral_por_pagar.json"

MESES = {
    1: "Enero", 2: "Febrero", 3: "Marzo", 4: "Abril",
    5: "Mayo", 6: "Junio", 7: "Julio", 8: "Agosto",
    9: "Septiembre", 10: "Octubre", 11: "Noviembre", 12: "Diciembre",
}


# =========================
# UTILIDADES GENERALES
# =========================
def cargar_json_lista(ruta):
    if os.path.exists(ruta):
        with open(ruta, "r", encoding="utf-8") as f:
            try:
                data = json.load(f)
                return data if isinstance(data, list) else []
            except json.JSONDecodeError:
                return []
    return []



def guardar_json_lista(ruta, data):
    with open(ruta, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=4)



def normalizar_nit(nit):
    return str(nit or "").strip().replace("-", "").replace(" ", "").upper()



def cargar_empresas():
    return cargar_json_lista(EMPRESAS_FILE)



def cargar_cierres_parciales():
    return cargar_json_lista(CIERRES_PARCIALES_FILE)



def cargar_isr_trimestral_guardado():
    return cargar_json_lista(ISR_TRIMESTRAL_FILE)



def trimestre_label(trimestre):
    return f"Trimestre {int(trimestre)}"



def texto_meses_acumulados(meses):
    meses = [int(x) for x in (meses or []) if int(x or 0) in MESES]
    if not meses:
        return ""
    return ", ".join(MESES[m] for m in meses)



def simplificar_txt(txt):
    base = str(txt or "").strip().lower()
    if not base:
        return ""
    base = unicodedata.normalize("NFKD", base)
    base = "".join(ch for ch in base if not unicodedata.combining(ch))
    return " ".join(base.split())



def contiene_alguno(texto, patrones):
    return any(p in texto for p in patrones)



def redondear_monto(valor):
    try:
        return round(float(valor or 0), 2)
    except Exception:
        return 0.0



def monto_neto_deudor(row):
    debe = redondear_monto(row.get("saldo_final_debe", 0))
    haber = redondear_monto(row.get("saldo_final_haber", 0))
    return round(abs(debe - haber), 2)



def monto_neto_ingreso(row):
    debe = redondear_monto(row.get("saldo_final_debe", 0))
    haber = redondear_monto(row.get("saldo_final_haber", 0))
    neto = round(haber - debe, 2)
    if abs(neto) <= 0.005:
        return 0.0
    return round(abs(neto), 2)



def clasificar_cuenta_para_cierre(row):
    tipo = simplificar_txt(row.get("tipo", ""))
    cuenta = simplificar_txt(row.get("cuenta", ""))

    if not cuenta or cuenta == "total":
        return None, None

    # Activos/pasivos/patrimonio quedan fuera del cierre parcial,
    # salvo Mercadería / Inventario inicial usado como Mercadería I.
    if contiene_alguno(cuenta, ["anticipo s/compras", "anticipo sobre compras", "anticipo de compras", "anticipo compras"]):
        return None, None

    es_mercaderia_ii = contiene_alguno(cuenta, ["inventario final", "mercaderia ii", "mercaderia ll", "mercaderia 2"])
    es_mercaderia_i = contiene_alguno(cuenta, ["inventario inicial"]) or (
        "mercaderia" in cuenta and not es_mercaderia_ii
    )

    if tipo in {"activo", "pasivo", "patrimonio"}:
        if es_mercaderia_i:
            return "costo_ventas", "mercaderia_i"
        return None, None

    # CONTRA-INGRESOS (opcionalmente reducen ingresos)
    if tipo == "ingreso" and contiene_alguno(cuenta, ["devolucion", "devoluciones", "rebaja", "rebajas", "descuento", "descuentos"]) and contiene_alguno(cuenta, ["venta", "ventas"]):
        return "contra_ingresos", "contra_ingresos"

    if tipo == "ingreso":
        return "ingresos", "ingresos"

    if es_mercaderia_i:
        return "costo_ventas", "mercaderia_i"

    if contiene_alguno(cuenta, ["flete sobre compra", "fletes sobre compra", "flete s/compras", "fletes s/compras"]):
        return "costo_ventas", "fletes_compras"

    if contiene_alguno(cuenta, ["devolucion", "devoluciones", "rebaja", "rebajas"]) and contiene_alguno(cuenta, ["compra", "compras"]):
        return "costo_ventas", "devoluciones_rebajas_compras"

    if contiene_alguno(cuenta, ["descuento", "descuentos"]) and contiene_alguno(cuenta, ["compra", "compras"]):
        return "costo_ventas", "descuentos_compras"

    if tipo == "gasto" and contiene_alguno(cuenta, ["compra", "compras"]):
        return "costo_ventas", "compras"

    if tipo == "gasto" and contiene_alguno(cuenta, ["costo de venta", "costo ventas", "costos de venta", "costos ventas"]):
        return "costo_ventas", "otros_costos"

    if tipo == "gasto":
        return "gastos", "gastos"

    return None, None


def construir_linea_detalle(row, monto, grupo):
    return {
        "codigo": str(row.get("codigo", "") or "").strip(),
        "cuenta": str(row.get("cuenta", "") or "").strip(),
        "tipo": str(row.get("tipo", "") or "").strip(),
        "grupo": grupo,
        "monto": round(float(monto or 0), 2),
    }


def normalizar_detalle_mercaderias_ii(detalle=None, total_respaldo=0.0):
    if isinstance(detalle, pd.DataFrame):
        detalle = detalle.to_dict("records")
    elif isinstance(detalle, dict):
        detalle = [detalle]
    elif detalle is None:
        detalle = []

    filas = []
    for idx, item in enumerate(detalle or []):
        cuenta = str((item or {}).get("cuenta", "") or "").strip()
        monto = redondear_monto((item or {}).get("monto", 0))
        if not cuenta and abs(monto) <= 0.005:
            continue
        if not cuenta:
            cuenta = f"Mercadería II {idx + 1}"
        if abs(monto) <= 0.005:
            continue
        filas.append({
            "codigo": "",
            "cuenta": cuenta,
            "tipo": "Ajuste manual",
            "grupo": "mercaderia_ii_manual",
            "monto": monto,
        })

    total_respaldo = redondear_monto(total_respaldo)
    if not filas and abs(total_respaldo) > 0.005:
        filas.append({
            "codigo": "",
            "cuenta": "Mercadería II",
            "tipo": "Ajuste manual",
            "grupo": "mercaderia_ii_manual",
            "monto": total_respaldo,
        })

    return sorted(filas, key=lambda x: str(x.get("cuenta", "")).lower())


def obtener_detalle_mercaderias_ii_desde_ajustes(ajustes):
    ajustes = ajustes or {}
    return normalizar_detalle_mercaderias_ii(
        ajustes.get("mercaderias_ii_detalle", []),
        total_respaldo=ajustes.get("mercaderia_ii_manual", 0),
    )


def serializar_detalle_mercaderias_ii_para_comparar(detalle):
    detalle_norm = normalizar_detalle_mercaderias_ii(detalle)
    base = [
        {"cuenta": str(x.get("cuenta", "") or "").strip(), "monto": redondear_monto(x.get("monto", 0))}
        for x in detalle_norm
    ]
    return json.dumps(base, ensure_ascii=False, sort_keys=True)


def normalizar_detalle_gastos_no_deducibles(detalle=None):
    if isinstance(detalle, pd.DataFrame):
        detalle = detalle.to_dict("records")
    elif isinstance(detalle, dict):
        detalle = [detalle]
    elif detalle is None:
        detalle = []

    filas = []
    for item in detalle or []:
        cuenta = str((item or {}).get("cuenta", "") or "").strip()
        monto = redondear_monto((item or {}).get("monto", 0))
        if not cuenta or abs(monto) <= 0.005:
            continue
        filas.append({
            "codigo": str((item or {}).get("codigo", "") or "").strip(),
            "cuenta": cuenta,
            "tipo": str((item or {}).get("tipo", "") or "").strip(),
            "grupo": "gastos_no_deducibles",
            "monto": monto,
        })

    return sorted(
        filas,
        key=lambda x: (str(x.get("codigo", "")).strip(), str(x.get("cuenta", "")).lower())
    )


def obtener_gastos_no_deducibles_desde_ajustes(ajustes):
    ajustes = ajustes or {}
    return normalizar_detalle_gastos_no_deducibles(ajustes.get("gastos_no_deducibles_detalle", []))


def serializar_detalle_gastos_no_deducibles_para_comparar(detalle):
    detalle_norm = normalizar_detalle_gastos_no_deducibles(detalle)
    base = [
        {
            "codigo": str(x.get("codigo", "") or "").strip(),
            "cuenta": str(x.get("cuenta", "") or "").strip(),
            "monto": redondear_monto(x.get("monto", 0)),
        }
        for x in detalle_norm
    ]
    return json.dumps(base, ensure_ascii=False, sort_keys=True)


def obtener_firma_ajustes_cierre(mercaderias_ii_detalle=None, gastos_no_deducibles_detalle=None):
    return json.dumps(
        {
            "mercaderias_ii": json.loads(serializar_detalle_mercaderias_ii_para_comparar(mercaderias_ii_detalle)),
            "gastos_no_deducibles": json.loads(serializar_detalle_gastos_no_deducibles_para_comparar(gastos_no_deducibles_detalle)),
        },
        ensure_ascii=False,
        sort_keys=True,
    )


def obtener_firma_ajustes_cierre_desde_registro(registro):
    ajustes = (registro or {}).get("ajustes_cierre", {}) or {}
    return obtener_firma_ajustes_cierre(
        mercaderias_ii_detalle=obtener_detalle_mercaderias_ii_desde_ajustes(ajustes),
        gastos_no_deducibles_detalle=obtener_gastos_no_deducibles_desde_ajustes(ajustes),
    )


def guardar_ajustes_cierre_si_cambian(
    registro,
    mercaderias_ii_detalle=None,
    gastos_no_deducibles_detalle=None,
    motivo="auto_guardado_ajustes",
    marcar_exportacion_excel=False,
):
    firma_actual = obtener_firma_ajustes_cierre_desde_registro(registro)
    firma_nueva = obtener_firma_ajustes_cierre(
        mercaderias_ii_detalle=mercaderias_ii_detalle,
        gastos_no_deducibles_detalle=gastos_no_deducibles_detalle,
    )
    if firma_actual == firma_nueva:
        return registro, False

    registro_actualizado = guardar_ajuste_mercaderia_en_cierre(
        registro,
        mercaderias_ii_detalle=mercaderias_ii_detalle,
        gastos_no_deducibles_detalle=gastos_no_deducibles_detalle,
        motivo=motivo,
        marcar_exportacion_excel=marcar_exportacion_excel,
    )
    return registro_actualizado, True


def es_cuenta_evaporizacion(cuenta):
    cuenta = simplificar_txt(cuenta)
    if not cuenta:
        return False
    return any(p in cuenta for p in ["evaporiz", "evaporacion", "evaporación"])


def calcular_total_evaporizacion(detalle_gastos):
    total = 0.0
    detalle = []
    for item in detalle_gastos or []:
        cuenta = str((item or {}).get("cuenta", "") or "").strip()
        if not es_cuenta_evaporizacion(cuenta):
            continue
        monto = redondear_monto((item or {}).get("monto", 0))
        if abs(monto) <= 0.005:
            continue
        total += monto
        detalle.append({
            "codigo": str((item or {}).get("codigo", "") or "").strip(),
            "cuenta": cuenta,
            "tipo": str((item or {}).get("tipo", "") or "").strip(),
            "monto": monto,
        })
    return {
        "evaporizacion_total": redondear_monto(total),
        "detalle_evaporizacion": detalle,
    }


def construir_cierre_resultados_desde_resumen(df_resumen, mercaderia_ii_manual=0.0, mercaderias_ii_detalle=None):
    detalle = {
        "ingresos": [],
        "contra_ingresos": [],
        "costo_ventas": {
            "mercaderia_i": [],
            "compras": [],
            "fletes_compras": [],
            "devoluciones_rebajas_compras": [],
            "descuentos_compras": [],
            "mercaderia_ii_manual": [],
            "otros_costos": [],
        },
        "gastos": [],
    }
    cuentas_excluidas = []

    if df_resumen is not None and not df_resumen.empty:
        for _, row in df_resumen.iterrows():
            cuenta = str(row.get("cuenta", "")).strip()
            if not cuenta or cuenta.upper() == "TOTAL":
                continue

            seccion, grupo = clasificar_cuenta_para_cierre(row)
            tipo = str(row.get("tipo", "") or "").strip()

            if seccion is None:
                cuentas_excluidas.append({
                    "codigo": str(row.get("codigo", "") or "").strip(),
                    "cuenta": cuenta,
                    "tipo": tipo,
                    "motivo": "No forma parte del estado de resultados parcial",
                })
                continue

            if seccion in {"ingresos", "contra_ingresos"}:
                monto = monto_neto_ingreso(row)
            else:
                monto = monto_neto_deudor(row)

            if abs(monto) <= 0.005:
                continue

            if seccion == "ingresos":
                detalle["ingresos"].append(construir_linea_detalle(row, monto, grupo))
            elif seccion == "contra_ingresos":
                detalle["contra_ingresos"].append(construir_linea_detalle(row, monto, grupo))
            elif seccion == "gastos":
                detalle["gastos"].append(construir_linea_detalle(row, monto, grupo))
            elif seccion == "costo_ventas":
                detalle["costo_ventas"][grupo].append(construir_linea_detalle(row, monto, grupo))

    mercaderias_ii_detalle = normalizar_detalle_mercaderias_ii(
        mercaderias_ii_detalle,
        total_respaldo=mercaderia_ii_manual,
    )
    if mercaderias_ii_detalle:
        detalle["costo_ventas"]["mercaderia_ii_manual"].extend(mercaderias_ii_detalle)

    for key in ["ingresos", "contra_ingresos", "gastos"]:
        detalle[key] = sorted(detalle[key], key=lambda x: (str(x.get("codigo", "")), str(x.get("cuenta", "")).lower()))

    for key in detalle["costo_ventas"]:
        detalle["costo_ventas"][key] = sorted(
            detalle["costo_ventas"][key],
            key=lambda x: (str(x.get("codigo", "")), str(x.get("cuenta", "")).lower())
        )

    total_ingresos_brutos = round(sum(x["monto"] for x in detalle["ingresos"]), 2)
    total_contra_ingresos = round(sum(x["monto"] for x in detalle["contra_ingresos"]), 2)
    total_ingresos = round(total_ingresos_brutos - total_contra_ingresos, 2)

    total_mercaderia_i = round(sum(x["monto"] for x in detalle["costo_ventas"]["mercaderia_i"]), 2)
    total_compras = round(sum(x["monto"] for x in detalle["costo_ventas"]["compras"]), 2)
    total_fletes = round(sum(x["monto"] for x in detalle["costo_ventas"]["fletes_compras"]), 2)
    total_dev_reb = round(sum(x["monto"] for x in detalle["costo_ventas"]["devoluciones_rebajas_compras"]), 2)
    total_desc = round(sum(x["monto"] for x in detalle["costo_ventas"]["descuentos_compras"]), 2)
    total_mercaderia_ii_manual = round(sum(x["monto"] for x in detalle["costo_ventas"]["mercaderia_ii_manual"]), 2)
    total_otros_costos = round(sum(x["monto"] for x in detalle["costo_ventas"]["otros_costos"]), 2)

    costo_ventas = round(
        total_mercaderia_i
        + total_compras
        + total_fletes
        + total_otros_costos
        - total_dev_reb
        - total_desc
        - total_mercaderia_ii_manual,
        2,
    )

    total_gastos = round(sum(x["monto"] for x in detalle["gastos"]), 2)
    utilidad_bruta = round(total_ingresos - costo_ventas, 2)
    utilidad_antes_isr = round(utilidad_bruta - total_gastos, 2)
    isr_acumulado_hasta_corte = round(max(utilidad_antes_isr, 0) * 0.25, 2)

    return {
        "tipo": "estado_resultados_parcial",
        "metodo": "basado_en_saldos_finales_acumulados_al_corte",
        "ajustes": {
            "mercaderia_ii_manual": total_mercaderia_ii_manual,
            "mercaderias_ii_detalle": detalle["costo_ventas"]["mercaderia_ii_manual"],
        },
        "detalle": detalle,
        "totales": {
            "ingresos_brutos": total_ingresos_brutos,
            "contra_ingresos": total_contra_ingresos,
            "ingresos": total_ingresos,
            "mercaderia_i": total_mercaderia_i,
            "compras": total_compras,
            "fletes_compras": total_fletes,
            "devoluciones_rebajas_compras": total_dev_reb,
            "descuentos_compras": total_desc,
            "mercaderia_ii_manual": total_mercaderia_ii_manual,
            "otros_costos": total_otros_costos,
            "costo_ventas": costo_ventas,
            "utilidad_bruta": utilidad_bruta,
            "gastos": total_gastos,
            "utilidad_antes_isr": utilidad_antes_isr,
            "isr_acumulado_hasta_corte": isr_acumulado_hasta_corte,
        },
        "cuentas_excluidas": cuentas_excluidas,
    }



def aplicar_estilo_mayor(df):
    if df is None or df.empty:
        return df

    def estilo(row):
        estilos = []
        for col in df.columns:
            if "saldo_inicial" in col:
                estilos.append("background-color: #E0F2FE")
            elif "movimientos" in col:
                estilos.append("background-color: #FEF3C7")
            elif "saldo_final" in col:
                estilos.append("background-color: #DCFCE7")
            else:
                estilos.append("")
        return estilos

    styler = df.style.apply(estilo, axis=1)

    if "cuenta" in df.columns:
        def resaltar_total(row):
            if str(row.get("cuenta", "")).strip().upper() == "TOTAL":
                return ["font-weight: bold; background-color: #E5E7EB"] * len(row)
            return [""] * len(row)

        styler = styler.apply(resaltar_total, axis=1)

    cols_monetarias = [c for c in df.columns if any(k in c for k in ["debe", "haber", "monto"])]
    if cols_monetarias:
        styler = styler.format({col: "Q {:,.2f}" for col in cols_monetarias})

    return styler



def construir_df_saldos_finales(registro):
    df = pd.DataFrame(registro.get("resumen_mayor", []))
    if df.empty:
        return df

    cols_base = [
        "codigo", "cuenta", "tipo", "naturaleza_real",
        "saldo_final_debe", "saldo_final_haber"
    ]
    disponibles = [c for c in cols_base if c in df.columns]
    df = df[disponibles].copy()

    if "cuenta" in df.columns:
        filas_normales = df[df["cuenta"].astype(str).str.upper() != "TOTAL"].copy()
        fila_total = df[df["cuenta"].astype(str).str.upper() == "TOTAL"].copy()
        df = pd.concat([filas_normales, fila_total], ignore_index=True)

    return df



def construir_df_detalle_cierre(items):
    df = pd.DataFrame(items or [])
    if df.empty:
        return df
    columnas = [c for c in ["codigo", "cuenta", "tipo", "grupo", "monto"] if c in df.columns]
    df = df[columnas].copy()
    if "monto" in df.columns:
        df["monto"] = pd.to_numeric(df["monto"], errors="coerce").fillna(0).round(2)
    return df



def construir_df_costo_ventas(cierre_resultados):
    detalle_cv = (cierre_resultados.get("detalle", {}) or {}).get("costo_ventas", {}) or {}
    filas = []
    etiquetas = {
        "mercaderia_i": "Mercadería I / Inventario inicial",
        "compras": "Compras",
        "fletes_compras": "Fletes sobre compras",
        "devoluciones_rebajas_compras": "(-) Devoluciones y rebajas s/compras",
        "descuentos_compras": "(-) Descuentos s/compras",
        "mercaderia_ii_manual": "(-) Mercadería II",
        "otros_costos": "Otros costos de venta",
    }
    for grupo, items in detalle_cv.items():
        for item in items:
            fila = dict(item)
            fila["concepto_costo"] = etiquetas.get(grupo, grupo)
            filas.append(fila)
    return construir_df_detalle_cierre(filas)



def obtener_cierre_resultados_registro(registro):
    ajustes = registro.get("ajustes_cierre", {}) or {}
    mercaderias_ii_detalle = obtener_detalle_mercaderias_ii_desde_ajustes(ajustes)
    df_resumen = pd.DataFrame(registro.get("resumen_mayor", []))
    return construir_cierre_resultados_desde_resumen(
        df_resumen,
        mercaderias_ii_detalle=mercaderias_ii_detalle,
    )



def guardar_ajuste_mercaderia_en_cierre(
    registro,
    mercaderias_ii_detalle=None,
    gastos_no_deducibles_detalle=None,
    motivo="actualizacion_manual",
    marcar_exportacion_excel=False
):
    mercaderias_ii_detalle = normalizar_detalle_mercaderias_ii(mercaderias_ii_detalle)
    total_mercaderia_ii = redondear_monto(sum(x.get("monto", 0) for x in mercaderias_ii_detalle))
    gastos_no_deducibles_detalle = normalizar_detalle_gastos_no_deducibles(gastos_no_deducibles_detalle)
    cierres = cargar_cierres_parciales()
    empresa_nit = normalizar_nit(registro.get("empresa_nit"))
    anio = int(registro.get("anio", 0) or 0)
    trimestre = int(registro.get("trimestre", 0) or 0)
    ahora = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    registro_actualizado = None
    for idx, item in enumerate(cierres):
        if (
            normalizar_nit(item.get("empresa_nit")) == empresa_nit
            and int(item.get("anio", 0) or 0) == anio
            and int(item.get("trimestre", 0) or 0) == trimestre
        ):
            item = dict(item)
            ajustes = item.get("ajustes_cierre", {}) or {}
            ajustes["mercaderias_ii_detalle"] = mercaderias_ii_detalle
            ajustes["mercaderia_ii_manual"] = total_mercaderia_ii
            ajustes["gastos_no_deducibles_detalle"] = gastos_no_deducibles_detalle
            ajustes["gastos_no_deducibles_total"] = redondear_monto(sum(x.get("monto", 0) for x in gastos_no_deducibles_detalle))
            item["ajustes_cierre"] = ajustes
            item["cierre_resultados"] = construir_cierre_resultados_desde_resumen(
                pd.DataFrame(item.get("resumen_mayor", [])),
                mercaderias_ii_detalle=mercaderias_ii_detalle,
            )
            item["fecha_actualizacion_cierre"] = ahora
            item["ultimo_movimiento_cierre"] = motivo
            if marcar_exportacion_excel:
                item["fecha_exportacion_excel"] = ahora
            cierres[idx] = item
            registro_actualizado = item
            break

    if registro_actualizado is None:
        raise ValueError("No se encontró el cierre parcial a actualizar.")

    guardar_json_lista(CIERRES_PARCIALES_FILE, cierres)
    return registro_actualizado


def nombre_archivo_seguro(texto):
    base = re.sub(r"[^A-Za-z0-9._-]+", "_", str(texto or "").strip())
    return base.strip("_") or "archivo"


def obtener_nombre_archivo_excel(registro):
    empresa = nombre_archivo_seguro(registro.get("empresa", "empresa"))
    anio = int(registro.get("anio", 0) or 0)
    trimestre = int(registro.get("trimestre", 0) or 0)
    return f"cierre_parcial_{empresa}_{anio}_T{trimestre}.xlsx"


def crear_excel_cierre_parcial(registro, cierre_resultados, resumen_isr=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Cierre Trimestral"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A8"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1
    ws.page_margins.left = 0.3
    ws.page_margins.right = 0.3
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5

    moneda_fmt = '"Q" #,##0.00;[Red]-"Q" #,##0.00'
    azul = PatternFill("solid", fgColor="DCE6F1")
    gris = PatternFill("solid", fgColor="D9D9D9")
    verde = PatternFill("solid", fgColor="DCE6D2")
    borde_top = Border(top=Side(style="thin", color="000000"))
    borde_top_double = Border(top=Side(style="double", color="000000"))
    borde_bottom_double = Border(bottom=Side(style="double", color="000000"))
    borde_bottom = Border(bottom=Side(style="thin", color="000000"))

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 48
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 18

    ws.merge_cells("A1:E1")
    ws["A1"] = "CIERRE TRIMESTRAL"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")

    empresa = str(registro.get("empresa", "") or "").strip()
    regimen = str(registro.get("regimen", "") or "").strip()
    anio = int(registro.get("anio", 0) or 0)
    trimestre = int(registro.get("trimestre", 0) or 0)
    corte = f"{registro.get('mes_corte_nombre', '')} {anio}"
    acumulado = texto_meses_acumulados(registro.get("meses_acumulados", []))

    meta = [
        ("Empresa", empresa),
        ("NIT", str(registro.get("empresa_nit", "") or "").strip()),
        ("Régimen", regimen),
        ("Trimestre", f"{trimestre}"),
        ("Corte", corte),
        ("Acum.", acumulado),
    ]
    fila = 2
    for etiqueta, valor in meta:
        ws[f"A{fila}"] = etiqueta
        ws[f"A{fila}"].font = Font(bold=True)
        ws[f"B{fila}"] = valor
        fila += 1

    ws[f"A{fila}"] = "No."
    ws[f"B{fila}"] = "Concepto"
    ws[f"D{fila}"] = "Parcial"
    ws[f"E{fila}"] = "Total"
    for col in range(1, 6):
        cell = ws.cell(fila, col)
        cell.font = Font(bold=True)
        cell.fill = azul
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(top=Side(style="thin"), bottom=Side(style="thin"))
    fila += 1

    consecutivo = 1

    def poner_monto(celda, monto, bold=False, border=None, fill=None):
        celda.value = redondear_monto(monto)
        celda.number_format = moneda_fmt
        celda.alignment = Alignment(horizontal="right", vertical="center")
        if bold:
            celda.font = Font(bold=True)
        if border is not None:
            celda.border = border
        if fill is not None:
            celda.fill = fill

    def escribir_seccion(titulo):
        nonlocal fila
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=5)
        celda = ws.cell(fila, 1)
        celda.value = titulo
        celda.font = Font(bold=True)
        celda.fill = gris
        celda.alignment = Alignment(horizontal="left", vertical="center")
        fila += 1

    def escribir_subtitulo(titulo):
        nonlocal fila
        ws.cell(fila, 2, titulo).font = Font(bold=True)
        fila += 1

    def escribir_detalles_con_subtotal_misma_linea(
        items,
        col_detalle,
        subtotal=None,
        col_subtotal=None,
        detalle_negativo=False,
        subtotal_negativo=False,
        col_total_final=None,
        total_final=None,
        borde_detalle_final=borde_bottom,
        borde_subtotal=None,
        borde_total_final=None,
    ):
        nonlocal fila, consecutivo
        if not items:
            return None
        ultima_fila = fila + len(items) - 1
        for item in items:
            ws.cell(fila, 1, consecutivo)
            ws.cell(fila, 2, str(item.get("cuenta", "") or "").strip())
            monto = redondear_monto(item.get("monto", 0))
            if detalle_negativo:
                monto *= -1
            borde_detalle = borde_detalle_final if fila == ultima_fila else None
            poner_monto(ws[f"{col_detalle}{fila}"], monto, border=borde_detalle)

            if fila == ultima_fila and subtotal is not None and col_subtotal:
                subtotal_val = redondear_monto(subtotal)
                if subtotal_negativo:
                    subtotal_val *= -1
                poner_monto(
                    ws[f"{col_subtotal}{fila}"],
                    subtotal_val,
                    bold=True,
                    border=borde_subtotal,
                )

            if fila == ultima_fila and total_final is not None and col_total_final:
                poner_monto(
                    ws[f"{col_total_final}{fila}"],
                    total_final,
                    bold=True,
                    border=borde_total_final,
                )

            fila += 1
            consecutivo += 1
        return ultima_fila

    def escribir_total_columna_total(etiqueta, monto, borde=None, fill=None):
        nonlocal fila
        ws.cell(fila, 2, etiqueta).font = Font(bold=True)
        if fill is not None:
            ws.cell(fila, 2).fill = fill
        poner_monto(
            ws.cell(fila, 5),
            monto,
            bold=True,
            border=borde,
            fill=fill,
        )
        fila += 1

    detalle = cierre_resultados.get("detalle", {}) or {}
    tot = cierre_resultados.get("totales", {}) or {}
    resumen_isr = resumen_isr or {
        "gastos_no_deducibles_total": 0.0,
        "utilidad_fiscal_antes_isr": redondear_monto(tot.get("utilidad_antes_isr", 0)),
        "isr_acumulado_hasta_corte": redondear_monto(tot.get("isr_acumulado_hasta_corte", 0)),
        "isr_trimestres_anteriores": 0.0,
        "evaporizacion_total": 0.0,
        "utilidad_sin_evaporizacion": redondear_monto(tot.get("utilidad_antes_isr", 0)),
        "ganancia_despues_impuesto": redondear_monto(tot.get("utilidad_antes_isr", 0)) - redondear_monto(tot.get("isr_acumulado_hasta_corte", 0)),
        "isr_trimestral_por_pagar": redondear_monto(tot.get("isr_acumulado_hasta_corte", 0)),
    }

    escribir_seccion("INGRESOS")
    escribir_detalles_con_subtotal_misma_linea(
        detalle.get("ingresos", []),
        col_detalle="D",
        subtotal=tot.get("ingresos_brutos", 0),
        col_subtotal="E",
    )

    if detalle.get("contra_ingresos"):
        escribir_seccion("(-) CONTRA INGRESOS")
        escribir_detalles_con_subtotal_misma_linea(
            detalle.get("contra_ingresos", []),
            col_detalle="D",
            subtotal=tot.get("contra_ingresos", 0),
            col_subtotal="E",
        )

    escribir_total_columna_total("INGRESOS NETOS", tot.get("ingresos", 0), borde=borde_top)

    escribir_seccion("COSTO DE VENTAS")
    costo_det = detalle.get("costo_ventas", {}) or {}

    grupos_costo = [
        ("Mercadería I", "mercaderia_i"),
        ("Compras", "compras"),
        ("Fletes sobre compras", "fletes_compras"),
        ("(-) Devoluciones y rebajas s/compras", "devoluciones_rebajas_compras"),
        ("(-) Descuentos s/compras", "descuentos_compras"),
        ("Otros costos de venta", "otros_costos"),
    ]

    for titulo_grupo, clave_grupo in grupos_costo:
        items = costo_det.get(clave_grupo, [])
        if not items:
            continue
        escribir_subtitulo(titulo_grupo)
        escribir_detalles_con_subtotal_misma_linea(
            items,
            col_detalle="C",
            subtotal=tot.get(clave_grupo, 0),
            col_subtotal="D",
        )

    disponible_para_venta = redondear_monto(
        tot.get("mercaderia_i", 0)
        + tot.get("compras", 0)
        + tot.get("fletes_compras", 0)
        + tot.get("otros_costos", 0)
        - tot.get("devoluciones_rebajas_compras", 0)
        - tot.get("descuentos_compras", 0)
    )
    ws.cell(fila, 2, "DISPONIBLE PARA LA VENTA").font = Font(bold=True)
    poner_monto(ws.cell(fila, 4), disponible_para_venta, bold=True, border=borde_top)
    fila += 1

    items_mercaderia_ii = costo_det.get("mercaderia_ii_manual", [])
    if items_mercaderia_ii:
        escribir_subtitulo("(-) Mercadería II")
        escribir_detalles_con_subtotal_misma_linea(
            items_mercaderia_ii,
            col_detalle="C",
            subtotal=tot.get("mercaderia_ii_manual", 0),
            col_subtotal="D",
            col_total_final="E",
            total_final=tot.get("costo_ventas", 0),
            borde_subtotal=borde_bottom,
        )
    else:
        escribir_total_columna_total("COSTO DE VENTAS", tot.get("costo_ventas", 0), borde=borde_top)

    escribir_total_columna_total("UTILIDAD BRUTA", tot.get("utilidad_bruta", 0), borde=borde_top)

    escribir_seccion("GASTOS DE OPERACIÓN")
    items_gastos = detalle.get("gastos", [])
    if items_gastos:
        escribir_detalles_con_subtotal_misma_linea(
            items_gastos,
            col_detalle="D",
            subtotal=tot.get("gastos", 0),
            col_subtotal="E",
            borde_subtotal=borde_bottom,
        )
    else:
        escribir_total_columna_total("TOTAL GASTOS", tot.get("gastos", 0), borde=borde_bottom)

    escribir_total_columna_total("UTILIDAD / PÉRDIDA ANTES DE ISR", tot.get("utilidad_antes_isr", 0), borde=borde_bottom_double)

    if redondear_monto(resumen_isr.get("gastos_no_deducibles_total", 0)) != 0:
        escribir_total_columna_total("(+) GASTOS NO DEDUCIBLES", resumen_isr.get("gastos_no_deducibles_total", 0), borde=borde_top)
        escribir_total_columna_total("UTILIDAD FISCAL ANTES DE ISR", resumen_isr.get("utilidad_fiscal_antes_isr", 0), borde=borde_top)

    escribir_total_columna_total("ISR ACUMULADO HASTA EL CORTE (25%)", resumen_isr.get("isr_acumulado_hasta_corte", 0), borde=borde_top)
    escribir_total_columna_total("(-) ISR DE TRIMESTRES ANTERIORES ACREDITADO", resumen_isr.get("isr_trimestres_anteriores", 0), borde=borde_top)

    ws.cell(fila, 2, "ISR TRIMESTRAL POR PAGAR").font = Font(bold=True)
    ws.cell(fila, 2).fill = verde
    poner_monto(
        ws.cell(fila, 5),
        resumen_isr.get("isr_trimestral_por_pagar", 0),
        bold=True,
        fill=verde,
        border=borde_top_double,
    )
    fila += 1

    escribir_total_columna_total("GANANCIA DESPUÉS DEL IMPUESTO", resumen_isr.get("ganancia_despues_impuesto", 0), borde=borde_top_double)

    for r in range(1, fila + 1):
        for c in range(1, 6):
            cell = ws.cell(r, c)
            horizontal = cell.alignment.horizontal or ("right" if c in {3, 4, 5} else "left")
            cell.alignment = Alignment(horizontal=horizontal, vertical="center")

    pie = fila + 2
    ws.merge_cells(start_row=pie, start_column=1, end_row=pie, end_column=5)
    ws.cell(pie, 1, "Formato generado desde el módulo Cierres Parciales. Este archivo refleja el cierre guardado al momento de preparar el Excel.")
    ws.cell(pie, 1).font = Font(italic=True, size=9)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def buscar_isr_guardado(data, empresa_nit, anio, trimestre):
    empresa_nit = normalizar_nit(empresa_nit)
    for item in data or []:
        if (
            normalizar_nit(item.get("empresa_nit")) == empresa_nit
            and int(item.get("anio", 0) or 0) == int(anio)
            and int(item.get("trimestre", 0) or 0) == int(trimestre)
        ):
            return item
    return None



def obtener_trimestres_acreditables(trimestre_actual):
    try:
        trimestre_actual = int(trimestre_actual or 0)
    except Exception:
        trimestre_actual = 0
    if trimestre_actual <= 1:
        return []
    return list(range(1, trimestre_actual))



def resumir_isr_para_registro(registro, cierre_resultados, data_isr=None):
    data_isr = data_isr if data_isr is not None else cargar_isr_trimestral_guardado()
    empresa_nit = normalizar_nit(registro.get("empresa_nit"))
    anio = int(registro.get("anio", 0) or 0)
    trimestre = int(registro.get("trimestre", 0) or 0)
    utilidad = redondear_monto((cierre_resultados.get("totales", {}) or {}).get("utilidad_antes_isr", 0))

    ajustes = registro.get("ajustes_cierre", {}) or {}
    gastos_no_deducibles_detalle = obtener_gastos_no_deducibles_desde_ajustes(ajustes)
    gastos_no_deducibles_total = redondear_monto(sum(x.get("monto", 0) for x in gastos_no_deducibles_detalle))
    utilidad_fiscal = redondear_monto(utilidad + gastos_no_deducibles_total)
    isr_acumulado = redondear_monto(max(utilidad_fiscal, 0) * 0.25)

    acreditables = obtener_trimestres_acreditables(trimestre)
    detalle_acreditado = []
    total_acreditado = 0.0
    for trim in acreditables:
        previo = buscar_isr_guardado(data_isr, empresa_nit, anio, trim)
        if not previo:
            continue
        monto = redondear_monto(
            previo.get("isr_trimestral_calculado", previo.get("isr_trimestral_por_pagar", 0))
        )
        if abs(monto) <= 0.005:
            continue
        total_acreditado += monto
        detalle_acreditado.append({
            "trimestre": trim,
            "mes_corte_nombre": str(previo.get("mes_corte_nombre", "") or "").strip(),
            "isr_trimestral_calculado": monto,
            "fecha_guardado": str(previo.get("fecha_guardado", "") or "").strip(),
        })

    total_acreditado = redondear_monto(total_acreditado)
    isr_trimestral_calculado = redondear_monto(max(isr_acumulado - total_acreditado, 0))
    isr_por_pagar = isr_trimestral_calculado

    detalle_gastos = (cierre_resultados.get("detalle", {}) or {}).get("gastos", []) or []
    evaporizacion_info = calcular_total_evaporizacion(detalle_gastos)
    evaporizacion_total = redondear_monto(evaporizacion_info.get("evaporizacion_total", 0))
    utilidad_sin_evaporizacion = redondear_monto(utilidad + evaporizacion_total)
    # La utilidad final después del impuesto parte de la utilidad contable antes de ISR.
    # Los gastos no deducibles solo aumentan la base fiscal del impuesto, pero no la utilidad final.
    ganancia_despues_impuesto = redondear_monto(utilidad - isr_por_pagar)

    return {
        "utilidad_antes_isr": utilidad,
        "gastos_no_deducibles_detalle": gastos_no_deducibles_detalle,
        "gastos_no_deducibles_total": gastos_no_deducibles_total,
        "utilidad_fiscal_antes_isr": utilidad_fiscal,
        "isr_acumulado_hasta_corte": isr_acumulado,
        "trimestres_acreditables": acreditables,
        "detalle_acreditado": detalle_acreditado,
        "isr_trimestres_anteriores": total_acreditado,
        "isr_trimestral_calculado": isr_trimestral_calculado,
        "isr_trimestral_por_pagar": isr_por_pagar,
        "detalle_evaporizacion": evaporizacion_info.get("detalle_evaporizacion", []),
        "evaporizacion_total": evaporizacion_total,
        "utilidad_sin_evaporizacion": utilidad_sin_evaporizacion,
        "ganancia_despues_impuesto": ganancia_despues_impuesto,
    }



def guardar_isr_trimestral_por_pagar(registro, cierre_resultados):
    data = cargar_isr_trimestral_guardado()
    empresa_nit = normalizar_nit(registro.get("empresa_nit"))
    anio = int(registro.get("anio", 0) or 0)
    trimestre = int(registro.get("trimestre", 0) or 0)
    resumen_isr = resumir_isr_para_registro(registro, cierre_resultados, data_isr=data)
    ajustes_registro = registro.get("ajustes_cierre", {}) or {}
    ajustes = cierre_resultados.get("ajustes", {}) or {}

    nuevo = {
        "empresa": str(registro.get("empresa", "") or "").strip(),
        "empresa_nit": empresa_nit,
        "regimen": str(registro.get("regimen", "") or "").strip(),
        "anio": anio,
        "trimestre": trimestre,
        "mes_corte": int(registro.get("mes_corte", 0) or 0),
        "mes_corte_nombre": str(registro.get("mes_corte_nombre", "") or "").strip(),
        "meses_acumulados": registro.get("meses_acumulados", []) or [],
        "base_calculo": "25% sobre utilidad fiscal acumulada al corte más gastos no deducibles menos ISR trimestral calculado de trimestres anteriores del mismo año",
        "porcentaje_isr": 25.0,
        "utilidad_antes_isr": resumen_isr["utilidad_antes_isr"],
        "gastos_no_deducibles_total": resumen_isr["gastos_no_deducibles_total"],
        "gastos_no_deducibles_detalle": resumen_isr["gastos_no_deducibles_detalle"],
        "utilidad_fiscal_antes_isr": resumen_isr["utilidad_fiscal_antes_isr"],
        "isr_acumulado_hasta_corte": resumen_isr["isr_acumulado_hasta_corte"],
        "trimestres_acreditables": resumen_isr["trimestres_acreditables"],
        "detalle_acreditado": resumen_isr["detalle_acreditado"],
        "isr_trimestres_anteriores": resumen_isr["isr_trimestres_anteriores"],
        "isr_trimestral_calculado": resumen_isr["isr_trimestral_calculado"],
        "mercaderias_ii_detalle": obtener_detalle_mercaderias_ii_desde_ajustes(ajustes_registro),
        "mercaderia_ii_manual": redondear_monto(ajustes.get("mercaderia_ii_manual", 0)),
        "evaporizacion_total": resumen_isr["evaporizacion_total"],
        "utilidad_sin_evaporizacion": resumen_isr["utilidad_sin_evaporizacion"],
        "ganancia_despues_impuesto": resumen_isr["ganancia_despues_impuesto"],
        "isr_trimestral_por_pagar": resumen_isr["isr_trimestral_por_pagar"],
        "fecha_guardado": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }

    actualizado = False
    for idx, item in enumerate(data):
        if (
            normalizar_nit(item.get("empresa_nit")) == empresa_nit
            and int(item.get("anio", 0) or 0) == anio
            and int(item.get("trimestre", 0) or 0) == trimestre
        ):
            data[idx] = nuevo
            actualizado = True
            break

    if not actualizado:
        data.append(nuevo)

    data.sort(key=lambda x: (
        str(x.get("empresa", "")).lower(),
        int(x.get("anio", 0) or 0),
        int(x.get("trimestre", 0) or 0)
    ))
    guardar_json_lista(ISR_TRIMESTRAL_FILE, data)
    return nuevo


# =========================
# UI
# =========================
st.set_page_config(page_title="Cierres Parciales", page_icon="📚", layout="wide")

st.markdown("""
<style>
.block-container {
    padding-top: 2rem;
    padding-bottom: 1rem;
    max-width: 100%;
}
.panel {
    background: #ffffff;
    border: 1px solid #e5e7eb;
    border-radius: 16px;
    padding: 18px;
    margin-bottom: 14px;
    box-shadow: 0 2px 10px rgba(0,0,0,0.04);
}
.panel-title {
    font-size: 1.05rem;
    font-weight: 700;
    margin-bottom: 12px;
    color: #111827;
}
.badge-top {
    text-align: right;
    font-size: 0.95rem;
    font-weight: 700;
    color: #7c3aed;
    background: #f5f3ff;
    border: 1px solid #ddd6fe;
    padding: 6px 12px;
    border-radius: 999px;
    display: inline-block;
}
.kpi {
    border: 1px solid #e5e7eb;
    border-radius: 14px;
    padding: 14px 16px;
    background: #fff;
    box-shadow: 0 2px 8px rgba(0,0,0,0.04);
}
.kpi-label {
    font-size: 12px;
    color: #6b7280;
    text-transform: uppercase;
    font-weight: 700;
    margin-bottom: 6px;
}
.kpi-value {
    font-size: 18px;
    font-weight: 700;
    color: #111827;
}
</style>
""", unsafe_allow_html=True)

cierres_data = cargar_cierres_parciales()
empresas_data = cargar_empresas()
isr_data = cargar_isr_trimestral_guardado()

st.markdown('<div class="panel">', unsafe_allow_html=True)
col_a, col_b = st.columns([2, 1])
with col_a:
    st.markdown("## 📚 Cierres parciales")
with col_b:
    st.markdown('<div class="badge-top">Lectura desde cierres_parciales.json</div>', unsafe_allow_html=True)
st.markdown('</div>', unsafe_allow_html=True)

if not cierres_data:
    st.info("Todavía no hay cierres trimestrales guardados. Primero guarda uno desde el módulo Mayor.")
    st.stop()

empresas_map = {}
for item in cierres_data:
    nit = normalizar_nit(item.get("empresa_nit"))
    nombre = str(item.get("empresa", "")).strip()
    if nit and nit not in empresas_map:
        empresas_map[nit] = nombre or nit

empresas = sorted(
    [{"id": k, "nombre": v} for k, v in empresas_map.items()],
    key=lambda x: x["nombre"].lower()
)

with st.container():
    c1, c2, c3 = st.columns(3)

    with c1:
        empresa_sel = st.selectbox("Empresa", empresas, format_func=lambda x: x["nombre"])

    cierres_empresa = [
        x for x in cierres_data
        if normalizar_nit(x.get("empresa_nit")) == normalizar_nit(empresa_sel["id"])
    ]
    anios = sorted({int(x.get("anio", 0) or 0) for x in cierres_empresa if int(x.get("anio", 0) or 0)})

    with c2:
        anio = st.selectbox("Año", anios, index=len(anios) - 1 if anios else 0)

    cierres_anio = [x for x in cierres_empresa if int(x.get("anio", 0) or 0) == int(anio)]
    trimestres = sorted({int(x.get("trimestre", 0) or 0) for x in cierres_anio if int(x.get("trimestre", 0) or 0)})

    with c3:
        trimestre = st.selectbox("Trimestre", trimestres, format_func=trimestre_label)

registro = None
for item in cierres_anio:
    if int(item.get("trimestre", 0) or 0) == int(trimestre):
        registro = item
        break

if not registro:
    st.warning("No se encontró el cierre parcial seleccionado.")
    st.stop()

estado = registro.get("estado_cuadre", {}) or {}
totales = estado.get("totales", {}) or {}
df_resumen = pd.DataFrame(registro.get("resumen_mayor", []))
df_pendientes = pd.DataFrame(registro.get("saldos_pendientes", []))
meses_txt = texto_meses_acumulados(registro.get("meses_acumulados", []))
cierre_resultados = obtener_cierre_resultados_registro(registro)
tot_cierre = cierre_resultados.get("totales", {}) or {}
resumen_isr = resumir_isr_para_registro(registro, cierre_resultados, data_isr=isr_data)
df_ingresos = construir_df_detalle_cierre((cierre_resultados.get("detalle", {}) or {}).get("ingresos", []))
df_contra_ingresos = construir_df_detalle_cierre((cierre_resultados.get("detalle", {}) or {}).get("contra_ingresos", []))
df_costo_ventas = construir_df_costo_ventas(cierre_resultados)
df_gastos = construir_df_detalle_cierre((cierre_resultados.get("detalle", {}) or {}).get("gastos", []))
df_excluidas = pd.DataFrame(cierre_resultados.get("cuentas_excluidas", []) or [])
mercaderias_ii_guardadas = obtener_detalle_mercaderias_ii_desde_ajustes(registro.get("ajustes_cierre", {}) or {})
mercaderia_ii_guardada = redondear_monto(sum(x.get("monto", 0) for x in mercaderias_ii_guardadas))
gastos_no_deducibles_guardados = obtener_gastos_no_deducibles_desde_ajustes(registro.get("ajustes_cierre", {}) or {})
isr_guardado = buscar_isr_guardado(isr_data, registro.get("empresa_nit"), registro.get("anio"), registro.get("trimestre"))

st.markdown('<div class="panel">', unsafe_allow_html=True)
st.markdown('<div class="panel-title">Resumen del cierre parcial guardado</div>', unsafe_allow_html=True)
st.caption(
    f"Empresa: {registro.get('empresa', '')} | Régimen: {registro.get('regimen', '')} | "
    f"Corte: {registro.get('mes_corte_nombre', '')} {int(registro.get('anio', 0) or 0)} | "
    f"Acumulado: {meses_txt}"
)

k1, k2, k3, k4 = st.columns(4)
with k1:
    st.markdown(
        f'<div class="kpi"><div class="kpi-label">Trimestre</div><div class="kpi-value">{int(registro.get("trimestre", 0) or 0)}</div></div>',
        unsafe_allow_html=True,
    )
with k2:
    st.markdown(
        f'<div class="kpi"><div class="kpi-label">Saldo final debe</div><div class="kpi-value">Q {float(totales.get("fin_debe", 0) or 0):,.2f}</div></div>',
        unsafe_allow_html=True,
    )
with k3:
    st.markdown(
        f'<div class="kpi"><div class="kpi-label">Saldo final haber</div><div class="kpi-value">Q {float(totales.get("fin_haber", 0) or 0):,.2f}</div></div>',
        unsafe_allow_html=True,
    )
with k4:
    st.markdown(
        f'<div class="kpi"><div class="kpi-label">Guardado</div><div class="kpi-value">{str(registro.get("fecha_guardado", ""))[:19]}</div></div>',
        unsafe_allow_html=True,
    )

if estado.get("cuadra_final", True):
    st.success("El cierre parcial guardado tiene saldo final cuadrado.")
else:
    st.error("El cierre parcial guardado no cuadra en el saldo final.")

st.markdown('</div>', unsafe_allow_html=True)

tab1, tab2, tab3, tab4, tab5 = st.tabs([
    "Cierre parcial",
    "Mayor acumulado guardado",
    "Saldos pendientes",
    "Cuentas excluidas del cierre",
    "JSON guardado"
])

with tab1:
    st.markdown('<div class="panel">', unsafe_allow_html=True)
    st.markdown('<div class="panel-title">Estado de resultados parcial acumulado al corte</div>', unsafe_allow_html=True)
    st.caption(
        "Aquí solo se toman cuentas de resultado. Activos, pasivos y patrimonio no entran al cierre parcial, "
        "salvo Mercadería I / Inventario inicial usado dentro del costo de ventas. "
        "Las Devoluciones y Rebajas s/Compras y los Descuentos s/Compras se restan del costo de ventas."
    )

    c1, c2, c3, c4, c5, c6 = st.columns(6)
    c1.metric("Ingresos netos", f"Q {float(tot_cierre.get('ingresos', 0) or 0):,.2f}")
    c2.metric("Costo de ventas", f"Q {float(tot_cierre.get('costo_ventas', 0) or 0):,.2f}")
    c3.metric("Gastos acumulados", f"Q {float(tot_cierre.get('gastos', 0) or 0):,.2f}")
    c4.metric("Utilidad/Pérdida antes ISR", f"Q {float(tot_cierre.get('utilidad_antes_isr', 0) or 0):,.2f}")
    c5.metric("ISR acumulado al corte", f"Q {float(resumen_isr.get('isr_acumulado_hasta_corte', 0) or 0):,.2f}")
    c6.metric("ISR trimestral por pagar", f"Q {float(resumen_isr.get('isr_trimestral_por_pagar', 0) or 0):,.2f}")

    if int(registro.get("trimestre", 0) or 0) == 1:
        if float(resumen_isr.get("gastos_no_deducibles_total", 0) or 0) != 0:
            st.info(
                f"Para marzo, la utilidad fiscal suma gastos no deducibles por Q {float(resumen_isr.get('gastos_no_deducibles_total', 0) or 0):,.2f}. "
                f"ISR acumulado al corte: Q {float(resumen_isr.get('isr_acumulado_hasta_corte', 0) or 0):,.2f}."
            )
        else:
            st.info("Para marzo, el ISR trimestral por pagar es igual al 25% de la utilidad acumulada al corte.")
    else:
        detalle_prev = resumen_isr.get("detalle_acreditado", []) or []
        acreditados_txt = ", ".join([
            f"T{int(x.get('trimestre', 0) or 0)}: Q {float(x.get('isr_trimestral_calculado', x.get('isr_trimestral_por_pagar', 0)) or 0):,.2f}"
            for x in detalle_prev
        ])
        extra_nd = ""
        if float(resumen_isr.get("gastos_no_deducibles_total", 0) or 0) != 0:
            extra_nd = f" | (+) Gastos no deducibles: Q {float(resumen_isr.get('gastos_no_deducibles_total', 0) or 0):,.2f}"
        if acreditados_txt:
            st.info(
                f"ISR acumulado al corte: Q {float(resumen_isr.get('isr_acumulado_hasta_corte', 0) or 0):,.2f}"
                f"{extra_nd} | (-) ISR acreditado de trimestres anteriores: Q {float(resumen_isr.get('isr_trimestres_anteriores', 0) or 0):,.2f} "
                f"({acreditados_txt}) | ISR trimestral por pagar: Q {float(resumen_isr.get('isr_trimestral_por_pagar', 0) or 0):,.2f}"
            )
        else:
            st.warning(
                f"Este trimestre descuenta ISR guardado de trimestres anteriores del mismo año. Aún no hay ISR previo guardado, "
                f"así que el ISR trimestral por pagar coincide con el acumulado: Q {float(resumen_isr.get('isr_trimestral_por_pagar', 0) or 0):,.2f}{extra_nd}"
            )

    st.markdown("### Ingresos acumulados")
    if df_ingresos.empty:
        st.info("No hay cuentas de ingresos incluidas en este cierre.")
    else:
        st.dataframe(
            df_ingresos,
            use_container_width=True,
            height=220,
            column_config={"monto": st.column_config.NumberColumn("Monto", format="Q %.2f")}
        )
    st.markdown(f"**Ingresos brutos:** Q {float(tot_cierre.get('ingresos_brutos', 0) or 0):,.2f}")
    if not df_contra_ingresos.empty:
        st.dataframe(
            df_contra_ingresos,
            use_container_width=True,
            height=160,
            column_config={"monto": st.column_config.NumberColumn("Monto", format="Q %.2f")}
        )
        st.markdown(f"**(-) Contra-ingresos / rebajas sobre ventas:** Q {float(tot_cierre.get('contra_ingresos', 0) or 0):,.2f}")
    st.markdown(f"**Total ingresos netos:** Q {float(tot_cierre.get('ingresos', 0) or 0):,.2f}")

    st.markdown("### Costo de ventas")
    col_cv_1, col_cv_2 = st.columns([1.2, 1])
    with col_cv_1:
        mercaderia_ii_manual_ui = st.number_input(
            "Mercadería II",
            min_value=0.0,
            step=0.01,
            value=float(mercaderia_ii_guardada),
            format="%.2f",
            key=f"mercaderia_ii_manual_{normalizar_nit(registro.get('empresa_nit'))}_{registro.get('anio')}_{registro.get('trimestre')}"
        )
        st.caption("Este valor se resta en el costo de ventas como inventario final manual del trimestre.")
        with st.expander("Detalle Mercaderías II (opcional)", expanded=False):
            detalle_mercaderias_ui_inicial = mercaderias_ii_guardadas or [{"cuenta": "", "monto": 0.0}]
            df_mercaderias_ii_ui = st.data_editor(
                pd.DataFrame([
                    {"cuenta": str(x.get("cuenta", "") or "").strip(), "monto": float(redondear_monto(x.get("monto", 0)))}
                    for x in detalle_mercaderias_ui_inicial
                ]),
                use_container_width=True,
                hide_index=True,
                num_rows="dynamic",
                key=f"mercaderias_ii_detalle_{normalizar_nit(registro.get('empresa_nit'))}_{registro.get('anio')}_{registro.get('trimestre')}",
                column_config={
                    "cuenta": st.column_config.TextColumn("Cuenta Mercadería II"),
                    "monto": st.column_config.NumberColumn("Monto", format="Q %.2f", min_value=0.0, step=0.01),
                },
            )
            st.caption("Si detallas varias mercaderías II, el sistema usará la suma del detalle para el cierre.")

        mercaderias_ii_ui = normalizar_detalle_mercaderias_ii(
            df_mercaderias_ii_ui,
            total_respaldo=mercaderia_ii_manual_ui,
        )
        mercaderia_ii_total_ui = redondear_monto(sum(x.get("monto", 0) for x in mercaderias_ii_ui))
        if mercaderias_ii_ui:
            st.markdown(f"**Mercadería II aplicada al cierre:** Q {mercaderia_ii_total_ui:,.2f}")
        detalle_mercaderias_cambio = (
            serializar_detalle_mercaderias_ii_para_comparar(mercaderias_ii_ui)
            != serializar_detalle_mercaderias_ii_para_comparar(
                normalizar_detalle_mercaderias_ii(mercaderias_ii_guardadas, total_respaldo=mercaderia_ii_guardada)
            )
        )

    with col_cv_2:
        st.markdown("**Fórmula aplicada**")
        st.code(
            "Costo de ventas = Mercadería I + Compras + Fletes + Otros costos"
            " - Devoluciones/Rebajas s/Compras - Descuentos s/Compras"
            " - Mercadería II",
            language="text"
        )

    gastos_disponibles_nd = normalizar_detalle_gastos_no_deducibles((cierre_resultados.get("detalle", {}) or {}).get("gastos", []))
    opciones_gastos_nd = {}
    for item in gastos_disponibles_nd:
        key = f"{str(item.get('codigo', '')).strip()}||{str(item.get('cuenta', '')).strip()}||{redondear_monto(item.get('monto', 0)):.2f}"
        opciones_gastos_nd[key] = item

    default_nd = []
    for item in gastos_no_deducibles_guardados:
        key = f"{str(item.get('codigo', '')).strip()}||{str(item.get('cuenta', '')).strip()}||{redondear_monto(item.get('monto', 0)):.2f}"
        if key in opciones_gastos_nd:
            default_nd.append(key)

    with st.expander("Ajuste fiscal para ISR: gastos no deducibles", expanded=False):
        if not opciones_gastos_nd:
            st.info("No hay gastos disponibles para marcar como no deducibles.")
            selected_nd_keys = []
        else:
            selected_nd_keys = st.multiselect(
                "Selecciona los gastos que no deben deducirse para el ISR trimestral",
                options=list(opciones_gastos_nd.keys()),
                default=default_nd,
                format_func=lambda k: f"{opciones_gastos_nd[k].get('cuenta', '')} — Q {float(opciones_gastos_nd[k].get('monto', 0) or 0):,.2f}",
                key=f"gastos_no_deducibles_{normalizar_nit(registro.get('empresa_nit'))}_{registro.get('anio')}_{registro.get('trimestre')}"
            )
            st.caption("Estos gastos se suman nuevamente para obtener la utilidad fiscal del ISR trimestral.")

    gastos_no_deducibles_ui = [opciones_gastos_nd[k] for k in selected_nd_keys] if opciones_gastos_nd else []
    gastos_no_deducibles_total_ui = redondear_monto(sum(x.get("monto", 0) for x in gastos_no_deducibles_ui))
    detalle_gastos_nd_cambio = (
        serializar_detalle_gastos_no_deducibles_para_comparar(gastos_no_deducibles_ui)
        != serializar_detalle_gastos_no_deducibles_para_comparar(gastos_no_deducibles_guardados)
    )

    registro, ajustes_auto_guardados = guardar_ajustes_cierre_si_cambian(
        registro,
        mercaderias_ii_detalle=mercaderias_ii_ui,
        gastos_no_deducibles_detalle=gastos_no_deducibles_ui,
        motivo="auto_guardado_ajustes_trimestre",
    )
    if ajustes_auto_guardados:
        cierre_resultados = obtener_cierre_resultados_registro(registro)
        tot_cierre = cierre_resultados.get("totales", {}) or {}
        resumen_isr = resumir_isr_para_registro(registro, cierre_resultados, data_isr=isr_data)
        mercaderias_ii_guardadas = obtener_detalle_mercaderias_ii_desde_ajustes(registro.get("ajustes_cierre", {}) or {})
        mercaderia_ii_guardada = redondear_monto(sum(x.get("monto", 0) for x in mercaderias_ii_guardadas))
        gastos_no_deducibles_guardados = obtener_gastos_no_deducibles_desde_ajustes(registro.get("ajustes_cierre", {}) or {})
        st.caption("Los ajustes de Mercadería II y gastos no deducibles de este trimestre ya quedan recordados automáticamente.")

    ajustes_preview = dict(registro.get("ajustes_cierre", {}) or {})
    ajustes_preview["mercaderias_ii_detalle"] = mercaderias_ii_ui
    ajustes_preview["mercaderia_ii_manual"] = mercaderia_ii_total_ui
    ajustes_preview["gastos_no_deducibles_detalle"] = gastos_no_deducibles_ui
    ajustes_preview["gastos_no_deducibles_total"] = gastos_no_deducibles_total_ui
    registro_preview = dict(registro)
    registro_preview["ajustes_cierre"] = ajustes_preview
    cierre_resultados_ui = obtener_cierre_resultados_registro(registro_preview)
    tot_cierre_ui = cierre_resultados_ui.get("totales", {}) or {}
    resumen_isr_ui = resumir_isr_para_registro(registro_preview, cierre_resultados_ui, data_isr=isr_data)
    df_costo_ventas_ui = construir_df_costo_ventas(cierre_resultados_ui)

    excel_state_key = f"excel_cierre_{normalizar_nit(registro.get('empresa_nit'))}_{registro.get('anio')}_{registro.get('trimestre')}"
    excel_name_key = f"excel_nombre_{normalizar_nit(registro.get('empresa_nit'))}_{registro.get('anio')}_{registro.get('trimestre')}"

    b1, b2, b3 = st.columns([1, 1, 1.2])
    with b1:
        if st.button("Guardar ajustes", type="secondary", use_container_width=True):
            try:
                guardar_ajuste_mercaderia_en_cierre(
                    registro,
                    mercaderias_ii_detalle=mercaderias_ii_ui,
                    gastos_no_deducibles_detalle=gastos_no_deducibles_ui,
                    motivo="ajustes_cierre"
                )
                st.success("Ajustes guardados en cierres_parciales.json")
                st.rerun()
            except Exception as e:
                st.error(f"No se pudo guardar el ajuste manual: {e}")

    with b2:
        if st.button("Guardar ISR trimestral por pagar", type="primary", use_container_width=True):
            try:
                if detalle_mercaderias_cambio or detalle_gastos_nd_cambio:
                    registro = guardar_ajuste_mercaderia_en_cierre(
                        registro,
                        mercaderias_ii_detalle=mercaderias_ii_ui,
                        gastos_no_deducibles_detalle=gastos_no_deducibles_ui,
                        motivo="guardar_isr_trimestral"
                    )
                    cierre_resultados = obtener_cierre_resultados_registro(registro)
                    tot_cierre = cierre_resultados.get("totales", {}) or {}
                    resumen_isr = resumir_isr_para_registro(registro, cierre_resultados, data_isr=isr_data)
                guardado = guardar_isr_trimestral_por_pagar(registro, cierre_resultados)
                st.success(
                    f"ISR trimestral por pagar guardado: Q {float(guardado.get('isr_trimestral_por_pagar', 0) or 0):,.2f}"
                )
                st.rerun()
            except Exception as e:
                st.error(f"No se pudo guardar el ISR trimestral: {e}")

    with b3:
        if st.button("Guardar cierre y preparar Excel", type="secondary", use_container_width=True):
            try:
                registro_excel = guardar_ajuste_mercaderia_en_cierre(
                    registro,
                    mercaderias_ii_detalle=mercaderias_ii_ui,
                    gastos_no_deducibles_detalle=gastos_no_deducibles_ui,
                    motivo="exportacion_excel",
                    marcar_exportacion_excel=True,
                )
                cierre_resultados_excel = obtener_cierre_resultados_registro(registro_excel)
                resumen_isr_excel = resumir_isr_para_registro(registro_excel, cierre_resultados_excel, data_isr=isr_data)
                excel_bytes = crear_excel_cierre_parcial(registro_excel, cierre_resultados_excel, resumen_isr=resumen_isr_excel)
                st.session_state[excel_state_key] = excel_bytes
                st.session_state[excel_name_key] = obtener_nombre_archivo_excel(registro_excel)
                st.success("Cierre parcial guardado y Excel contable preparado correctamente.")
            except Exception as e:
                st.error(f"No se pudo preparar el Excel contable: {e}")

    if st.session_state.get(excel_state_key):
        st.download_button(
            "Descargar Excel contable",
            data=st.session_state[excel_state_key],
            file_name=st.session_state.get(excel_name_key, "cierre_parcial.xlsx"),
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
        )
        st.caption("Al preparar el Excel, el sistema actualiza y guarda automáticamente el cierre parcial vigente en cierres_parciales.json.")

    if df_costo_ventas_ui.empty:
        st.info("No hay cuentas clasificadas como costo de ventas en este cierre.")
    else:
        st.dataframe(
            df_costo_ventas_ui,
            use_container_width=True,
            height=280,
            column_config={"monto": st.column_config.NumberColumn("Monto", format="Q %.2f")}
        )

    st.markdown(f"**Mercadería I / Inventario inicial:** Q {float(tot_cierre_ui.get('mercaderia_i', 0) or 0):,.2f}")
    st.markdown(f"**Compras:** Q {float(tot_cierre_ui.get('compras', 0) or 0):,.2f}")
    st.markdown(f"**Fletes sobre compras:** Q {float(tot_cierre_ui.get('fletes_compras', 0) or 0):,.2f}")
    st.markdown(f"**(-) Devoluciones y rebajas s/compras:** Q {float(tot_cierre_ui.get('devoluciones_rebajas_compras', 0) or 0):,.2f}")
    st.markdown(f"**(-) Descuentos s/compras:** Q {float(tot_cierre_ui.get('descuentos_compras', 0) or 0):,.2f}")
    st.markdown(f"**(-) Mercadería II:** Q {float(tot_cierre_ui.get('mercaderia_ii_manual', 0) or 0):,.2f}")
    if float(tot_cierre_ui.get('otros_costos', 0) or 0) != 0:
        st.markdown(f"**Otros costos de venta:** Q {float(tot_cierre_ui.get('otros_costos', 0) or 0):,.2f}")
    st.markdown(f"**Total costo de ventas:** Q {float(tot_cierre_ui.get('costo_ventas', 0) or 0):,.2f}")
    st.markdown(f"**Utilidad bruta:** Q {float(tot_cierre_ui.get('utilidad_bruta', 0) or 0):,.2f}")

    st.markdown("### Gastos acumulados")
    if df_gastos.empty:
        st.info("No hay cuentas de gastos incluidas en este cierre.")
    else:
        st.dataframe(
            df_gastos,
            use_container_width=True,
            height=260,
            column_config={"monto": st.column_config.NumberColumn("Monto", format="Q %.2f")}
        )
    st.markdown(f"**Total gastos acumulados:** Q {float(tot_cierre_ui.get('gastos', 0) or 0):,.2f}")
    st.markdown(f"**(+) Gastos no deducibles para ISR:** Q {float(resumen_isr_ui.get('gastos_no_deducibles_total', 0) or 0):,.2f}")

    resultado_final = float(tot_cierre_ui.get('utilidad_antes_isr', 0) or 0)
    if resultado_final >= 0:
        st.success(f"Utilidad acumulada antes de ISR: Q {resultado_final:,.2f}")
    else:
        st.error(f"Pérdida acumulada antes de ISR: Q {abs(resultado_final):,.2f}")

    st.markdown(f"**Utilidad fiscal para ISR:** Q {float(resumen_isr_ui.get('utilidad_fiscal_antes_isr', 0) or 0):,.2f}")
    st.markdown(f"**ISR acumulado hasta el corte (25%):** Q {float(resumen_isr_ui.get('isr_acumulado_hasta_corte', 0) or 0):,.2f}")
    if float(resumen_isr_ui.get('isr_trimestres_anteriores', 0) or 0) != 0:
        st.markdown(f"**(-) ISR acreditado de trimestres anteriores:** Q {float(resumen_isr_ui.get('isr_trimestres_anteriores', 0) or 0):,.2f}")
    st.markdown(f"**ISR trimestral calculado:** Q {float(resumen_isr_ui.get('isr_trimestral_calculado', 0) or 0):,.2f}")
    st.markdown(f"**ISR trimestral por pagar del trimestre actual:** Q {float(resumen_isr_ui.get('isr_trimestral_por_pagar', 0) or 0):,.2f}")
    if float(resumen_isr_ui.get('evaporizacion_total', 0) or 0) != 0:
        st.markdown(f"**Utilidad sin la evaporización:** Q {float(resumen_isr_ui.get('utilidad_sin_evaporizacion', 0) or 0):,.2f}")
    st.markdown(f"**Ganancia después del impuesto:** Q {float(resumen_isr_ui.get('ganancia_despues_impuesto', 0) or 0):,.2f}")

    if isr_guardado:
        st.info(
            f"ISR trimestral guardado en JSON: Q {float(isr_guardado.get('isr_trimestral_por_pagar', 0) or 0):,.2f} "
            f"(guardado el {str(isr_guardado.get('fecha_guardado', ''))[:19]})"
        )
        detalles_guardado = []
        if float(isr_guardado.get('gastos_no_deducibles_total', 0) or 0) != 0:
            detalles_guardado.append(f"gastos no deducibles por Q {float(isr_guardado.get('gastos_no_deducibles_total', 0) or 0):,.2f}")
        if float(isr_guardado.get('isr_trimestres_anteriores', 0) or 0) != 0:
            detalles_guardado.append(f"ISR previo acreditado por Q {float(isr_guardado.get('isr_trimestres_anteriores', 0) or 0):,.2f}")
        if detalles_guardado:
            st.caption(f"Ese registro ya quedó guardado considerando {' y '.join(detalles_guardado)}.")
    else:
        st.warning("Todavía no has guardado el ISR trimestral por pagar para este cierre.")

    st.markdown('</div>', unsafe_allow_html=True)

with tab2:

    st.markdown('<div class="panel">', unsafe_allow_html=True)
    st.markdown('<div class="panel-title">Mayor acumulado guardado en el JSON</div>', unsafe_allow_html=True)
    if df_resumen.empty:
        st.info("No hay resumen del mayor guardado.")
    else:
        st.dataframe(aplicar_estilo_mayor(df_resumen), use_container_width=True, height=560)
    st.markdown('</div>', unsafe_allow_html=True)

with tab3:
    st.markdown('<div class="panel">', unsafe_allow_html=True)
    st.markdown('<div class="panel-title">Cuentas pendientes guardadas con el cierre</div>', unsafe_allow_html=True)
    if df_pendientes.empty:
        st.info("No hay saldos pendientes guardados para este cierre parcial.")
    else:
        st.dataframe(aplicar_estilo_mayor(df_pendientes), use_container_width=True, height=420)
    st.markdown('</div>', unsafe_allow_html=True)

with tab4:
    st.markdown('<div class="panel">', unsafe_allow_html=True)
    st.markdown('<div class="panel-title">Cuentas excluidas del cierre parcial</div>', unsafe_allow_html=True)
    st.caption("Estas cuentas quedaron fuera porque no forman parte del estado de resultados parcial. Aquí normalmente verás activos, pasivos y patrimonio como Caja, Bancos, Clientes, Proveedores, Capital o Anticipo S/Compras.")
    if df_excluidas.empty:
        st.info("No hay cuentas excluidas registradas para este cierre.")
    else:
        st.dataframe(df_excluidas, use_container_width=True, height=420)
    st.markdown('</div>', unsafe_allow_html=True)

with tab5:
    st.markdown('<div class="panel">', unsafe_allow_html=True)
    st.markdown('<div class="panel-title">Detalle del registro JSON</div>', unsafe_allow_html=True)
    st.json(registro)
    st.markdown('</div>', unsafe_allow_html=True)
